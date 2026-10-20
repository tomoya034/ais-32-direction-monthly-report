from __future__ import annotations

import hashlib
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from datetime import date
from pathlib import Path
from unittest import mock

import ais_monthly_app as app


NS = {
    "s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "x": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}
# Independent, immutable transcription from the five accepted manual June
# workbooks' drawing1.xml; these are not calculated from producer templates.
A_ANCHORS = (((12, 9524, 1, 19050), (22, 609599, 30, 190500)),
             ((0, 9525, 43, 0), (14, 590550, 56, 142875)))
B_ANCHORS = (((11, 838200, 1, 33336), (23, 0, 31, 9525)),
             ((0, 0, 43, 14287), (15, 9524, 56, 157162)))
INTEGRATED_ANCHORS = (((12, 0, 1, 9524), (23, 9525, 31, 19049)),
                      ((0, 0, 43, 9525), (14, 590549, 56, 152400)))
EXPECTED = {
    "full_a": (A_ANCHORS, [], None),
    "summary_a": (A_ANCHORS, [], None),
    "full_b": (B_ANCHORS, [(1, 1, "14.5703125"), (2, 6, "12.7109375"),
                            (7, 7, "11.5703125"), (8, 12, "12.7109375"),
                            (24, 33, "12.7109375")], "9"),
    "summary_b": (B_ANCHORS, [(1, 1, "14.5703125")], "9"),
    "integrated": (INTEGRATED_ANCHORS, [(1, 1, "14.5703125")], "9"),
}


class HistoricalLayoutTests(unittest.TestCase):
    def setUp(self) -> None:
        temporary = tempfile.TemporaryDirectory(prefix="ais32-layout-")
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.config = app.AppConfig(self.root, self.root / "modern.xlsx", "KLNG", 2026, 6)

    def write_role(self, role: str, count: int = 30, *, port: str = "KLNG") -> Path:
        config = app.AppConfig(self.root, self.root / "modern.xlsx", port, 2026, 6)
        days = [app.DayResult(date(2026, 6, index + 1), (), {}) for index in range(count)]
        path = self.root / f"{port}-{role}-{count}.xlsx"
        if role.startswith("full_"):
            app._write_period_full_workbook(path, config, days,
                                           app.PERIOD_A if role == "full_a" else app.PERIOD_B,
                                           {}, None)
        else:
            app._write_summary_workbook(
                path, config, days, {}, scope_title="Display title is not a role",
                artifact_role=role, integrated_inputs=({}, {}) if role == "integrated" else None,
            )
        return path

    def geometry(self, path: Path):
        with zipfile.ZipFile(path) as archive:
            sheet = next(ET.fromstring(archive.read(name)) for name in archive.namelist()
                         if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
                         and b"drawing" in archive.read(name))
            drawing = ET.fromstring(archive.read("xl/drawings/drawing1.xml"))
            anchors = tuple(tuple(tuple(int(vertex.find(f"x:{field}", NS).text)
                                             for field in ("col", "colOff", "row", "rowOff"))
                                  for vertex in (item.find("x:from", NS), item.find("x:to", NS)))
                            for item in drawing.findall("x:twoCellAnchor", NS))
            columns = [(int(item.get("min")), int(item.get("max")), item.get("width"))
                       for item in sheet.findall("s:cols/s:col", NS) if item.get("hidden") != "1"]
            return anchors, columns, sheet.find("s:sheetFormatPr", NS)

    def test_all_five_roles_exact_reference_anchors_and_grid(self) -> None:
        for role, (expected, columns, default) in EXPECTED.items():
            with self.subTest(role=role):
                anchors, actual_columns, sheet_format = self.geometry(self.write_role(role))
                self.assertEqual(anchors, expected)
                self.assertEqual(actual_columns, columns)
                self.assertEqual(sheet_format.get("defaultColWidth"), default)
                self.assertEqual(sheet_format.get("defaultRowHeight"), "15.75")

    def test_formula_source_radar_missing_helpers_preserve_primary_blank(self) -> None:
        import openpyxl
        for port in ("HWLN", "KLNG"):
            for role in EXPECTED:
                with self.subTest(port=port, role=role):
                    path = self.write_role(role, count=2, port=port)
                    values = openpyxl.load_workbook(path, data_only=True)
                    formulas = openpyxl.load_workbook(path, data_only=False)
                    try:
                        sheet = values["總表"]
                        self.assertTrue(all(sheet.cell(row, col).value is None
                                            for row in (2, 3) for col in range(2, 34)))
                        formula_backed = role in {"full_a", "full_b", "integrated"}
                        for row in (2, 3):
                            for index in range(32):
                                cell = sheet.cell(row, 41 + index)
                                if formula_backed:
                                    source = f"{openpyxl.utils.get_column_letter(index + 2)}{row}"
                                    self.assertEqual(formulas["總表"].cell(row, 41 + index).value,
                                                     f'=IF(COUNT({source})=0,NA(),{source})')
                                    self.assertEqual(cell.value, "#N/A")
                                    self.assertEqual(cell.data_type, "e")
                                else:
                                    self.assertIsNone(cell.value)
                        if formula_backed:
                            self.assertTrue(formulas["總表"].column_dimensions["AO"].hidden)
                        with zipfile.ZipFile(path) as archive:
                            radar = ET.fromstring(archive.read("xl/charts/chart1.xml"))
                        source = radar.find(".//c:val/c:numRef/c:f", NS).text
                        self.assertEqual(source, "總表!$AO$2:$BT$2" if formula_backed else "總表!$B$2:$AG$2")
                        self.assertEqual(radar.find(".//c:val/c:numRef/c:numCache/c:ptCount", NS).get("val"), "32")
                        self.assertEqual(radar.findall(".//c:val/c:numRef/c:numCache/c:pt", NS), [])
                    finally:
                        values.close()
                        formulas.close()

    def test_month_lengths_shift_both_statistics_vertices_not_radar(self) -> None:
        for count in (28, 29, 31):
            for role, (expected, _, _) in EXPECTED.items():
                # June has 30 days; geometry accepts an arbitrary day sequence,
                # so use valid dates in a 31-day month for the 31-row case.
                with self.subTest(count=count, role=role):
                    days = [app.DayResult(date(2026, 7, index + 1), (), {}) for index in range(count)]
                    config = app.AppConfig(self.root, self.root / "modern.xlsx", "KLNG", 2026, 7)
                    path = self.root / f"{role}-{count}.xlsx"
                    import xlsxwriter
                    workbook = xlsxwriter.Workbook(path)
                    app._write_delivery_summary_sheet(workbook, config, days, {},
                                                      scope_title="irrelevant", artifact_role=role)
                    workbook.close()
                    app._apply_historical_geometry(path, role, count)
                    actual, _, _ = self.geometry(path)
                    self.assertEqual(actual[0], expected[0])
                    for edge in range(2):
                        col, col_off, row, row_off = expected[1][edge]
                        self.assertEqual(actual[1][edge], (col, col_off, row + count - 30, row_off))
                        self.assertNotEqual(actual[1][edge][2], row)

    def test_layout_is_artifact_role_not_port_or_display_title(self) -> None:
        for role in EXPECTED:
            self.assertEqual(self.geometry(self.write_role(role, port="HWLN"))[:2],
                             self.geometry(self.write_role(role, port="KLNG"))[:2])

    def test_unknown_or_incompatible_roles_fail_before_writing(self) -> None:
        for role in ("unknown", "", "full_a", "full_b"):
            path = self.root / "must-not-exist.xlsx"
            with self.subTest(role=role), self.assertRaises(ValueError):
                app._write_summary_workbook(path, self.config, [], {},
                                            scope_title="Historical Period A 00:00–11:59",
                                            artifact_role=role)
            self.assertFalse(path.exists())
        with self.assertRaisesRegex(ValueError, "period"):
            app._write_period_full_workbook(self.root / "must-not-exist.xlsx", self.config,
                                            [], "display Period A", {}, None)
        with self.assertRaisesRegex(ValueError, "must agree"):
            app._write_summary_workbook(self.root / "must-not-exist.xlsx", self.config,
                                        [], {}, scope_title="irrelevant", artifact_role="integrated")

    def test_exact_normal_font_and_role_specific_chart_fonts_layouts(self) -> None:
        for role in EXPECTED:
            with self.subTest(role=role), zipfile.ZipFile(self.write_role(role)) as archive:
                styles = ET.fromstring(archive.read("xl/styles.xml"))
                self.assertEqual(styles.find("s:fonts/s:font/s:name", NS).get("val"), "新細明體")
                radar = ET.fromstring(archive.read("xl/charts/chart1.xml"))
                statistics = ET.fromstring(archive.read("xl/charts/chart2.xml"))
                self.assertEqual(radar.find(".//c:legend/c:txPr//a:defRPr", NS).get("sz"), "900")
                for chart in (radar, statistics):
                    self.assertEqual(chart.find("c:lang", NS).get("val"), "zh-TW")
                    fonts = chart.findall(".//c:title//a:defRPr", NS)
                    self.assertTrue(fonts)
                    self.assertTrue(all(font.get("sz") == "1400" for font in fonts))
                    self.assertEqual(chart.find(".//c:title//a:latin", NS).get("typeface"),
                                     "Calibri" if role == "integrated" else "+mn-lt")
                    # A/B reference tokens; Integrated's literal compatibility
                    # strategy is supported by isolated native 0-pixel proof,
                    # not by rewriting reference numerical expectations.
                    for properties in chart.findall(".//a:defRPr", NS) + chart.findall(".//a:rPr", NS):
                        if properties.find("a:latin", NS) is None:
                            continue
                        self.assertEqual([properties.find(f"a:{script}", NS).get("typeface")
                                          for script in ("latin", "ea", "cs")],
                                         ["Calibri"] * 3 if role == "integrated"
                                         else ["+mn-lt", "+mn-ea", "+mn-cs"])
                    parent_default = chart.find("c:txPr/a:p/a:pPr/a:defRPr", NS)
                    self.assertIsNotNone(parent_default)
                    self.assertEqual(parent_default.attrib, {})
                    self.assertEqual(list(parent_default), [])
                    self.assertTrue(all(end.get("lang") == "zh-TW"
                                        for end in chart.findall(".//a:endParaRPr", NS)))
                    for axis_format in chart.findall(".//c:valAx/c:numFmt", NS):
                        self.assertEqual(axis_format.attrib,
                                         {"formatCode": "0.0000" if role == "integrated" else "General",
                                          "sourceLinked": "1"})
                    self.assertEqual(chart.find(".//c:title//a:srgbClr", NS).get("val"), "595959")
                    title_runs = chart.findall("c:chart/c:title/c:tx/c:rich/a:p/a:r", NS)
                    self.assertTrue(title_runs)
                    for run in title_runs:
                        text = run.find("a:t", NS).text
                        self.assertEqual(run.find("a:rPr", NS).get("lang"),
                                         "en-US" if text.isascii() else "zh-TW")
                        self.assertTrue(text.isascii() or all(ord(char) > 127 for char in text))
                self.assertIsNone(statistics.find(".//c:plotArea/c:layout/c:manualLayout", NS))
                manual = radar.find(".//c:plotArea/c:layout/c:manualLayout", NS)
                if role.endswith("_a"):
                    self.assertIsNotNone(manual)
                    self.assertAlmostEqual(float(manual.find("c:x", NS).get("val")),
                                           0.064771103328797777, places=14)
                else:
                    self.assertIsNone(manual)

    def test_script_theme_fonts_match_reference_not_chart_scalar_font_name(self) -> None:
        # Read-only extraction from the copied reference theme1.xml: A/B minor
        # Latin=Calibri, EA/CS="", Hant=新細明體; integrated primary scripts
        # Latin/EA/CS=Calibri. This must not be derived from app templates.
        for role in EXPECTED:
            with self.subTest(role=role), zipfile.ZipFile(self.write_role(role)) as archive:
                theme = ET.fromstring(archive.read("xl/theme/theme1.xml"))
                minor = theme.find("a:themeElements/a:fontScheme/a:minorFont", NS)
                self.assertIsNotNone(minor)
                self.assertEqual(minor.find("a:latin", NS).get("typeface"), "Calibri")
                self.assertEqual([minor.find(f"a:{script}", NS).get("typeface")
                                  for script in ("ea", "cs")],
                                 ["Calibri", "Calibri"] if role == "integrated" else ["", ""])
                if role == "integrated":
                    self.assertEqual([font.tag.rsplit("}", 1)[-1] for font in minor],
                                     ["latin", "ea", "cs"])
                else:
                    self.assertEqual(next(font.get("typeface") for font in minor.findall("a:font", NS)
                                          if font.get("script") == "Hant"), "新細明體")

    def test_normalization_preserves_detail_and_all_chart_series_bytes(self) -> None:
        for role in EXPECTED:
            with self.subTest(role=role):
                with mock.patch.object(app, "_apply_historical_geometry"):
                    path = self.write_role(role)
                with zipfile.ZipFile(path) as archive:
                    before = {name: archive.read(name) for name in archive.namelist()}
                app._apply_historical_geometry(path, role, 30)
                with zipfile.ZipFile(path) as archive:
                    self.assertEqual(set(archive.namelist()), set(before))
                    for name, original in before.items():
                        result = archive.read(name)
                        if name.startswith("xl/charts/chart") and name.endswith(".xml"):
                            old = ET.fromstring(original).findall(".//c:ser", NS)
                            new = ET.fromstring(result).findall(".//c:ser", NS)
                            self.assertEqual([ET.tostring(series) for series in old],
                                             [ET.tostring(series) for series in new], name)
                        elif name in ("xl/styles.xml", "xl/drawings/drawing1.xml"):
                            continue
                        elif name == "xl/theme/theme1.xml" and role == "integrated":
                            old = ET.fromstring(original)
                            new = ET.fromstring(result)
                            for script in ("latin", "ea", "cs"):
                                new.find(f"a:themeElements/a:fontScheme/a:minorFont/a:{script}", NS).set(
                                    "typeface", old.find(
                                        f"a:themeElements/a:fontScheme/a:minorFont/a:{script}", NS,
                                    ).get("typeface"),
                                )
                            old_minor = old.find("a:themeElements/a:fontScheme/a:minorFont", NS)
                            new_minor = new.find("a:themeElements/a:fontScheme/a:minorFont", NS)
                            for font in old_minor:
                                if font.tag.rsplit("}", 1)[-1] not in ("latin", "ea", "cs"):
                                    new_minor.append(ET.fromstring(ET.tostring(font)))
                            self.assertEqual(ET.tostring(new), ET.tostring(old))
                        elif name.startswith("xl/worksheets/sheet") and name.endswith(".xml") and b"drawing" in original:
                            self.assertEqual(ET.tostring(ET.fromstring(original).find("s:sheetData", NS)),
                                             ET.tostring(ET.fromstring(result).find("s:sheetData", NS)))
                        else:
                            self.assertEqual(result, original, name)

    def test_radar_spokes_and_value_axis_line_match_each_reference_role(self) -> None:
        for role in EXPECTED:
            with self.subTest(role=role), zipfile.ZipFile(self.write_role(role)) as archive:
                radar = ET.fromstring(archive.read("xl/charts/chart1.xml"))
                category = radar.find(".//c:catAx", NS)
                value = radar.find(".//c:valAx", NS)
                # Reference A has spokes. The B and integrated artifacts do
                # not; simply hiding all axes would fail the A negative case.
                spokes = category.find("c:majorGridlines", NS)
                if role.endswith("_a"):
                    self.assertIsNotNone(spokes)
                    self.assertEqual(spokes.find("c:spPr/a:ln", NS).get("w"), "9525")
                    self.assertEqual(spokes.find("c:spPr/a:ln/a:solidFill/a:srgbClr", NS).get("val"),
                                     "D9D9D9")
                else:
                    self.assertIsNone(spokes)
                self.assertIsNotNone(category.find("c:spPr/a:ln/a:solidFill", NS))
                self.assertIsNotNone(value.find("c:spPr/a:ln/a:noFill", NS))
                self.assertIsNotNone(value.find("c:majorGridlines", NS))

    def test_statistics_bar_spacing_matches_all_five_reference_artifacts(self) -> None:
        for role in EXPECTED:
            with self.subTest(role=role), zipfile.ZipFile(self.write_role(role)) as archive:
                statistics = ET.fromstring(archive.read("xl/charts/chart2.xml"))
                columns = statistics.find(".//c:barChart", NS)
                self.assertEqual(columns.find("c:grouping", NS).get("val"), "clustered")
                self.assertEqual(columns.find("c:gapWidth", NS).get("val"), "219")
                self.assertEqual(columns.find("c:overlap", NS).get("val"), "-27")
                self.assertEqual(len(columns.findall("c:ser", NS)), 2)
                for axis in statistics.findall(".//c:catAx", NS) + statistics.findall(".//c:valAx", NS):
                    self.assertEqual(axis.find("c:majorTickMark", NS).get("val"), "none")
                    self.assertEqual(axis.find("c:minorTickMark", NS).get("val"), "none")
                category_line = statistics.find(".//c:catAx/c:spPr/a:ln", NS)
                self.assertEqual(category_line.get("w"), "9525")
                self.assertEqual(category_line.find("a:solidFill/a:srgbClr", NS).get("val"), "D9D9D9")
                title_layout = statistics.find("c:chart/c:title/c:layout/c:manualLayout", NS)
                if role.endswith("_a"):
                    self.assertEqual(title_layout.find("c:xMode", NS).get("val"), "edge")
                    self.assertEqual(title_layout.find("c:yMode", NS).get("val"), "edge")
                    self.assertAlmostEqual(float(title_layout.find("c:x", NS).get("val")),
                                           0.3523162112572919, places=14)
                    self.assertAlmostEqual(float(title_layout.find("c:y", NS).get("val")),
                                           0.027777777777777776, places=14)
                else:
                    self.assertIsNone(title_layout)

    def test_all_roles_preserve_missing_statistics_and_actual_zero_in_chart_caches(self) -> None:
        config = app.AppConfig(self.root, self.root / "modern.xlsx", "HWLN", 2026, 6)
        days = [app.DayResult(date(2026, 6, 1), (), {})]
        values = {(days[0].day, app.DIRECTION_ORDER[0]): 0.0,
                  (days[0].day, app.DIRECTION_ORDER[2]): 125.25}
        for role in EXPECTED:
            path = self.root / f"partial-{role}.xlsx"
            with self.subTest(role=role):
                if role.startswith("full_"):
                    # Target the shared summary renderer with a partial-value
                    # fixture; no retained-detail record is fabricated. Full
                    # workflow coverage uses real parsed fixtures separately.
                    import xlsxwriter
                    workbook = xlsxwriter.Workbook(path)
                    app._write_delivery_summary_sheet(workbook, config, days, values,
                        scope_title="partial fixture", artifact_role=role)
                    workbook.close()
                    app._apply_historical_geometry(path, role, 1)
                else:
                    app._write_summary_workbook(path, config, days, values,
                        scope_title="fixture", artifact_role=role,
                        integrated_inputs=(values, values) if role == "integrated" else None)
                with zipfile.ZipFile(path) as archive:
                    chart = ET.fromstring(archive.read("xl/charts/chart2.xml"))
                    series = chart.findall(".//c:ser", NS)
                    self.assertEqual(len(series), 3)
                    for item in series:
                        cache = item.find("c:val/c:numRef/c:numCache", NS)
                        self.assertEqual(cache.find("c:ptCount", NS).get("val"), "32")
                        points = {int(point.get("idx")): float(point.find("c:v", NS).text)
                                  for point in cache.findall("c:pt", NS)}
                        self.assertEqual(points, {0: 0.0, 2: 125.25})
                        self.assertNotIn(28, points)  # HWLN NW is truly absent.
                import openpyxl
                workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
                try:
                    sheet = workbook["總表"]
                    number_format = "0.0000" if role == "integrated" else "General"
                    # Both radar source and statistics helper source participate
                    # in Excel's source-linked native tick-format resolution.
                    self.assertEqual(sheet["B2"].number_format, number_format)
                    self.assertEqual(sheet["AJ2"].number_format, number_format)
                    self.assertIsNone(sheet["AJ30"].value)
                    self.assertEqual(sheet["AJ2"].value, 0.0)
                    self.assertFalse(sheet["B1"].font.bold)
                    self.assertIsNone(sheet["B1"].alignment.horizontal)
                    self.assertEqual(sheet["B1"].alignment.vertical, "center")
                    self.assertIsNone(sheet["B1"].border.bottom.style)
                    self.assertEqual(sheet["A2"].alignment.vertical, "center")
                    self.assertEqual(sheet["B4"].alignment.vertical, "center")
                    self.assertFalse(sheet["A4"].font.bold)
                finally:
                    workbook.close()

    def test_failed_zip_copy_preserves_authored_file_and_cleans_only_own_temp(self) -> None:
        path = self.write_role("full_b")
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        with mock.patch.object(app.shutil, "copyfileobj", side_effect=OSError("injected copy failure")):
            with self.assertRaisesRegex(OSError, "injected copy failure"):
                app._apply_historical_geometry(path, "full_b", 30)
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), digest)
        self.assertEqual(list(self.root.glob(".*.layout-*.xlsx")), [])
        for count in (0, 32):
            with self.assertRaises(ValueError):
                app._apply_historical_geometry(path, "full_b", count)
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), digest)


if __name__ == "__main__":
    unittest.main()
