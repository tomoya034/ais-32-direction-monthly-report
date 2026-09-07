from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import math
import statistics
import sys
import time
import zipfile
from dataclasses import replace
from pathlib import Path

import openpyxl

# The historic baseline stores Chinese keys; use fixed index metadata to decode it.
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from test_direction_contracts import V150_DIRECTION_ORDER


def selector_checksum(app, day_results) -> str:
    payload = []
    for day_result in day_results:
        for direction_index, direction in enumerate(app.DIRECTION_ORDER):
            result = day_result.directions[direction]
            payload.append(
                [
                    day_result.day.isoformat(),
                    direction_index,
                    result.selected,
                    result.selected_rank,
                    result.cluster_count,
                ]
            )
    encoded = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def read_grid(path: Path) -> dict[tuple[dt.date, int], float | None]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["總表"]
        grid: dict[tuple[dt.date, int], float | None] = {}
        for row in sheet.iter_rows(
            min_row=2, max_row=31, min_col=1, max_col=33, values_only=True
        ):
            raw_day = row[0]
            day = raw_day.date() if isinstance(raw_day, dt.datetime) else raw_day
            if not isinstance(day, dt.date):
                raise AssertionError(f"invalid total-grid date in {path.name}: {raw_day!r}")
            for direction_index, raw_value in enumerate(row[1:]):
                grid[(day, direction_index)] = (
                    None if raw_value in (None, "") else float(raw_value)
                )
        return grid
    finally:
        workbook.close()


def assert_grid_equal(label: str, actual, expected, tolerance: float = 1e-9) -> int:
    if set(actual) != set(expected):
        raise AssertionError(
            f"{label}: key mismatch actual={len(actual)} expected={len(expected)}"
        )
    matched = 0
    for key in sorted(expected):
        left = actual[key]
        right = expected[key]
        if left is None or right is None:
            if left is not None or right is not None:
                raise AssertionError(f"{label}: {key}: {left!r} != {right!r}")
        elif not math.isclose(left, right, rel_tol=0.0, abs_tol=tolerance):
            raise AssertionError(f"{label}: {key}: {left!r} != {right!r}")
        matched += 1
    return matched


def expected_stats(grid, direction_index: int):
    values = [
        value
        for (day, index), value in sorted(grid.items())
        if index == direction_index and value is not None
    ]
    if not values:
        return (None, None, None, None)
    return (
        max(values),
        min(values),
        statistics.fmean(values),
        statistics.stdev(values) if len(values) > 1 else None,
    )


def validate_statistics_and_charts(app, path: Path, grid) -> None:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["總表"]
        for direction_index in range(32):
            expected = (
                expected_stats(grid, direction_index)
                if direction_index in app.OPEN_SEA_INDEXES
                else (None, None, None, None)
            )
            for offset, expected_value in enumerate(expected):
                actual = sheet.cell(row=33 + offset, column=2 + direction_index).value
                actual_value = None if actual in (None, "") else float(actual)
                if expected_value is None:
                    if actual_value is not None:
                        raise AssertionError(
                            f"{path.name} stats direction={direction_index} row={33+offset}: "
                            f"expected blank, got {actual_value}"
                        )
                elif actual_value is None or not math.isclose(
                    actual_value, expected_value, rel_tol=0.0, abs_tol=1e-9
                ):
                    raise AssertionError(
                        f"{path.name} stats direction={direction_index} row={33+offset}: "
                        f"{actual_value} != {expected_value}"
                    )
    finally:
        workbook.close()
    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
        if bad_member is not None:
            raise AssertionError(f"{path.name} corrupt zip member: {bad_member}")
        charts = [
            name
            for name in archive.namelist()
            if name.startswith("xl/charts/chart") and name.endswith(".xml")
        ]
        if len(charts) != 2:
            raise AssertionError(f"{path.name}: expected 2 charts, got {len(charts)}")


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(4 * 1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def progress(payload: dict) -> None:
    kind = payload.get("kind")
    if kind == "job":
        print(
            f"JOB logical_days={payload['source_count']} fragments={payload['fragment_count']}",
            flush=True,
        )
    elif kind in {"file_done", "cache_hit"}:
        print(
            f"{kind.upper()} {payload['position']}/{payload['total']} {payload['file']} "
            f"seconds={payload.get('seconds', 0):.2f}",
            flush=True,
        )
    elif kind == "parallel_start":
        print(f"PARALLEL workers={payload['workers']} pending={payload['total']}", flush=True)
    elif kind == "preflight":
        print(
            f"PREFLIGHT required={payload['required_bytes']} free={payload.get('free_bytes')}",
            flush=True,
        )
    elif kind == "delivery_file_done":
        print(
            f"DELIVERY {payload['position']}/{payload['total']} {payload['file']}",
            flush=True,
        )
    elif kind in {"delivery_write_start", "delivery_write_done", "write_start", "write_done"}:
        print(kind.upper(), flush=True)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--repo", type=Path, required=True)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--golden", type=Path, required=True)
    parser.add_argument("--gate", type=Path, required=True)
    parser.add_argument("--baseline", type=Path, required=True)
    parser.add_argument("--v150-cache-dir", type=Path, required=True)
    parser.add_argument("--output-root", type=Path, required=True)
    parser.add_argument("--workers", type=int, default=3)
    args = parser.parse_args()

    sys.path.insert(0, str(args.repo))
    import ais_monthly_app as app

    args.output_root.mkdir(parents=True, exist_ok=True)
    analysis = args.output_root / "KLNG_2026年06月_32方位數值_新版自動分析.xlsx"
    delivery_root = args.output_root / "正式交付" / "KLNG_202606"
    report_path = args.output_root / "full_june_regression_report.json"
    config = app.AppConfig(
        input_dir=args.input,
        output_path=analysis,
        port="KLNG",
        year=2026,
        month=6,
        delivery_dir=delivery_root,
        workers=args.workers,
        overwrite=True,
    )

    started = time.perf_counter()
    day_results, warnings = app.process_month(config, callback=progress)
    processing_seconds = time.perf_counter() - started
    nonempty_results = [item for item in day_results if item.source_fragments]
    raw_rows = sum(item.rows_scanned for item in nonempty_results)
    normalized_rows = sum(item.rows_normalized for item in nonempty_results)
    fragment_count = sum(len(item.fragment_info) for item in nonempty_results)
    if (len(nonempty_results), fragment_count, raw_rows) != (30, 60, 25_850_522):
        raise AssertionError(
            f"input cardinality mismatch: days={len(nonempty_results)} "
            f"fragments={fragment_count} rows={raw_rows}"
        )

    baseline = json.loads(args.baseline.read_text(encoding="utf-8"))
    baseline_by_day = {item["date"]: item for item in baseline["days"]}
    modern_cells = 0
    v150_modern_cells = 0
    identical_spools = 0
    for day_result in nonempty_results:
        expected_day = baseline_by_day[day_result.day.strftime("%Y%m%d")]
        expected_by_index = {
            index: expected_day["current_full"][label]
            for index, label in enumerate(V150_DIRECTION_ORDER)
        }
        v150 = json.loads((args.v150_cache_dir / f"{day_result.day:%Y%m%d}.json").read_text(encoding="utf-8"))["result"]
        v150_by_index = {index: v150["directions"][label] for index, label in enumerate(V150_DIRECTION_ORDER)}
        if day_result.spool_sha256 != v150["spool_sha256"]:
            raise AssertionError(f"normalized spool changed: {day_result.day}")
        identical_spools += 1
        for index, direction in enumerate(app.DIRECTION_ORDER):
            actual = day_result.directions[direction]
            expected = expected_by_index[index]
            if index in app.OPEN_SEA_INDEXES:
                previous = v150_by_index[index]
                if (actual.selected, actual.selected_rank, actual.cluster_count) != (
                    previous["selected"], previous["selected_rank"], previous["cluster_count"]
                ):
                    raise AssertionError(f"v1.5.0 Modern numeric mismatch: {day_result.day} index={index}")
                v150_modern_cells += 1
            actual_tuple = (actual.selected, actual.selected_rank, actual.cluster_count)
            expected_tuple = (
                expected["selected"],
                expected["rank"],
                expected["cluster_count"],
            )
            if actual_tuple != expected_tuple:
                raise AssertionError(
                    f"modern baseline mismatch {day_result.day} {direction}: "
                    f"{actual_tuple} != {expected_tuple}"
                )
            modern_cells += 1

    gate = json.loads(args.gate.read_text(encoding="utf-8"))
    evidence = gate["matched_evidence"]
    expected_evidence = {
        (
            item["date"],
            item["period"],
            int(item["direction_index"]),
            item["fragment"],
            int(item["source_row"]),
        ): item
        for item in evidence
    }
    if len(expected_evidence) != 1260:
        raise AssertionError(f"expected 1260 unique gate candidates, got {len(expected_evidence)}")

    production_candidates: dict[tuple[str, dt.date, str], str] = {}
    historical_rows = 0
    matched_rows = 0
    display_misses = []
    for day_result in nonempty_results:
        day_key = day_result.day.strftime("%Y%m%d")
        _counts, records = app.read_normalized_spool(
            app._normalized_spool_path(config, day_result.day)
        )
        for direction_index, record in records:
            if record.msg_type not in app.HISTORICAL_MESSAGE_TYPES:
                continue
            historical_rows += 1
            fragment_index = record.source_key >> 32
            source_row = record.source_key & 0xFFFFFFFF
            fragment = day_result.fragment_info[fragment_index]
            key = (
                day_key,
                record.period,
                direction_index,
                fragment.path.name,
                source_row,
            )
            expected = expected_evidence.get(key)
            if expected is None:
                continue
            timestamp = dt.datetime.combine(day_result.day, dt.time()) + dt.timedelta(
                seconds=record.second_of_day
            )
            checks = (
                math.isclose(record.distance, float(expected["distance"]), rel_tol=0.0, abs_tol=1e-9),
                math.isclose(record.bearing, float(expected["bearing"]), rel_tol=0.0, abs_tol=1e-9),
                record.msg_type == int(expected["msg_type"]),
                timestamp.isoformat() == expected["timestamp"],
                record.mmsi == expected["mmsi"],
                record.channel == expected["channel"],
                fragment.sheet == expected["sheet"],
            )
            if not all(checks):
                raise AssertionError(f"gate provenance mismatch: {key}")
            direction = app.DIRECTION_ORDER[direction_index]
            decision_key = (record.period, day_result.day, direction)
            candidate_id = app._candidate_id(fragment, record)
            if decision_key in production_candidates:
                raise AssertionError(f"duplicate Golden evidence for {decision_key}")
            production_candidates[decision_key] = candidate_id
            displayed = {
                item.candidate_id
                for item in day_result.period_directions[record.period][direction].candidates
            }
            if candidate_id not in displayed:
                display_misses.append(
                    {
                        "day": day_result.day.isoformat(),
                        "period": record.period,
                        "direction_index": direction_index,
                        "fragment": fragment.path.name,
                        "source_row": source_row,
                    }
                )
            matched_rows += 1

    if historical_rows != 23_226_151:
        raise AssertionError(f"historical eligible row mismatch: {historical_rows}")
    if matched_rows != 1260 or len(production_candidates) != 1260:
        raise AssertionError(
            f"production candidate mapping mismatch: rows={matched_rows} "
            f"unique={len(production_candidates)}"
        )
    if display_misses:
        raise AssertionError(
            f"Golden candidate-only contract is not usable from displayed candidates: "
            f"{len(display_misses)} misses; sample={display_misses[:3]}"
        )

    default_snapshot = app.build_default_decision_snapshot(config, day_results)
    reviewed_decisions = []
    for decision in default_snapshot.decisions:
        if decision.scope in (app.PERIOD_A, app.PERIOD_B):
            candidate_id = production_candidates.get(
                (decision.scope, decision.day, decision.direction)
            )
            if candidate_id is None:
                raise AssertionError(
                    f"missing reviewed period decision: "
                    f"{decision.scope} {decision.day} {decision.direction}"
                )
            reviewed_decisions.append(replace(decision, candidate_id=candidate_id))
        else:
            reviewed_decisions.append(decision)
    reviewed_snapshot = app.DecisionSnapshot(
        port=default_snapshot.port,
        year=default_snapshot.year,
        month=default_snapshot.month,
        source_manifest_hash=default_snapshot.source_manifest_hash,
        decisions=tuple(reviewed_decisions),
    )
    app.validate_decision_snapshot_contract(reviewed_snapshot, config, day_results)
    resolved = app.resolve_decision_values(reviewed_snapshot, day_results, config)
    if len(resolved) != 30 * 21 * 3:
        raise AssertionError(f"resolved decision count mismatch: {len(resolved)}")

    print("WRITE REVIEWED MODERN WORKBOOK", flush=True)
    app.write_monthly_workbook(
        config,
        day_results,
        warnings,
        reviewed_snapshot,
        callback=progress,
    )
    print("FINALIZE FIVE OFFICIAL WORKBOOKS", flush=True)
    delivery_started = time.perf_counter()
    paths = app.finalize_review_workbook(
        analysis,
        delivery_dir=delivery_root,
        overwrite=True,
        callback=progress,
    )
    delivery_seconds = time.perf_counter() - delivery_started

    golden_a = read_grid(args.golden / "KLNG 6月 32方位數值_2 總表.xlsx")
    golden_b = read_grid(args.golden / "KLNG 6月 32方位數值 總表.xlsx")
    golden_integrated = read_grid(
        args.golden / "KLNG_6月_32方位_每日較大值整合總表.xlsx"
    )
    output_grids = {
        "period_a_full": read_grid(paths.period_a_full),
        "period_b_full": read_grid(paths.period_b_full),
        "period_a_summary": read_grid(paths.period_a_summary),
        "period_b_summary": read_grid(paths.period_b_summary),
        "integrated": read_grid(paths.integrated_summary),
    }
    grid_checks = {
        "period_a_full_vs_golden": assert_grid_equal(
            "period_a_full_vs_golden", output_grids["period_a_full"], golden_a
        ),
        "period_b_full_vs_golden": assert_grid_equal(
            "period_b_full_vs_golden", output_grids["period_b_full"], golden_b
        ),
        "period_a_small_vs_big": assert_grid_equal(
            "period_a_small_vs_big",
            output_grids["period_a_summary"],
            output_grids["period_a_full"],
        ),
        "period_b_small_vs_big": assert_grid_equal(
            "period_b_small_vs_big",
            output_grids["period_b_summary"],
            output_grids["period_b_full"],
        ),
        "integrated_vs_golden": assert_grid_equal(
            "integrated_vs_golden", output_grids["integrated"], golden_integrated
        ),
    }
    for key in sorted(golden_integrated):
        left = output_grids["period_a_summary"][key]
        right = output_grids["period_b_summary"][key]
        expected = (
            max(value for value in (left, right) if value is not None)
            if any(value is not None for value in (left, right))
            else None
        )
        actual = output_grids["integrated"][key]
        if expected is None:
            if actual is not None:
                raise AssertionError(f"integrated MAX mismatch {key}")
        elif actual is None or not math.isclose(
            actual, expected, rel_tol=0.0, abs_tol=1e-9
        ):
            raise AssertionError(f"integrated MAX mismatch {key}: {actual} != {expected}")

    formula_workbook = openpyxl.load_workbook(
        paths.integrated_summary, read_only=False, data_only=False
    )
    try:
        formula_sheet = formula_workbook["總表"]
        formula_count = 0
        for row in range(2, 32):
            for column in range(2, 34):
                direction_index = column - 2
                value = formula_sheet.cell(row=row, column=column).value
                if direction_index in app.OPEN_SEA_INDEXES:
                    if not isinstance(value, str) or "MAX('Period A'!" not in value:
                        raise AssertionError(
                            f"integrated formula mismatch {formula_sheet.cell(row=row, column=column).coordinate}: {value!r}"
                        )
                    formula_count += 1
                elif value is not None:
                    raise AssertionError("integrated land cell must remain blank")
    finally:
        formula_workbook.close()
    if formula_count != 630:
        raise AssertionError(f"integrated formula count mismatch: {formula_count}")

    expected_by_path = {
        paths.period_a_full: output_grids["period_a_full"],
        paths.period_b_full: output_grids["period_b_full"],
        paths.period_a_summary: output_grids["period_a_summary"],
        paths.period_b_summary: output_grids["period_b_summary"],
        paths.integrated_summary: output_grids["integrated"],
    }
    for path, grid in expected_by_path.items():
        print(f"VALIDATE {path.name}", flush=True)
        validate_statistics_and_charts(app, path, grid)

    output_files = [analysis, *paths.all_files()]
    output_manifest = [
        {
            "name": path.name,
            "size_bytes": path.stat().st_size,
            "sha256": file_sha256(path),
        }
        for path in output_files
    ]
    report = {
        "status": "PASS",
        "version": app.APP_VERSION,
        "input": {
            "logical_days": len(nonempty_results),
            "fragments": fragment_count,
            "raw_rows": raw_rows,
            "normalized_rows": normalized_rows,
            "historical_profile_rows": historical_rows,
        },
        "modern": {
            "baseline_cells": modern_cells,
            "v150_open_sea_selected_rank_cluster_cells": v150_modern_cells,
            "v150_normalized_spools_byte_identical": identical_spools,
            "identity": "direction index (0..31); English abbreviations unchanged",
            "selector_checksum": selector_checksum(app, nonempty_results),
        },
        "candidate_contract": {
            "gate_candidates": len(expected_evidence),
            "production_matched": matched_rows,
            "display_misses": len(display_misses),
        },
        "delivery_contracts": {
            "detail_max_equals_period_final": "asserted during production renderer for every day/direction",
            "grid_checks": grid_checks,
            "integrated_formula_cells": formula_count,
            "statistics_and_two_charts": len(expected_by_path),
        },
        "timing_seconds": {
            "processing": processing_seconds,
            "delivery_and_finalizer": delivery_seconds,
            "total": time.perf_counter() - started,
        },
        "outputs": output_manifest,
    }
    if v150_modern_cells != 630 or identical_spools != 30:
        raise AssertionError("incomplete v1.5.0 comparison")
    report_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(report, ensure_ascii=False, indent=2), flush=True)
    print(f"REPORT {report_path}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
