from __future__ import annotations

import argparse
import datetime as dt
import json
import math
import sys
import time
import zipfile
from pathlib import Path

import openpyxl
from lxml import etree


NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def cached_daily_finals(path: Path, app):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        expected = [f"6月{day}日" for day in range(1, 31)]
        if workbook.sheetnames[:30] != expected:
            raise AssertionError(f"unexpected daily sheet order in {path.name}")
        result = {}
        for day, sheet_name in enumerate(expected, start=1):
            row = next(
                workbook[sheet_name].iter_rows(
                    min_row=2, max_row=2, min_col=8, max_col=39, values_only=True
                )
            )
            for direction_index, value in enumerate(row):
                result[(dt.date(2026, 6, day), direction_index)] = (
                    None if value in (None, "") else float(value)
                )
        return result
    finally:
        workbook.close()


def summary_grid(path: Path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["總表"]
        result = {}
        for row in sheet.iter_rows(
            min_row=2, max_row=31, min_col=1, max_col=33, values_only=True
        ):
            day = row[0].date() if isinstance(row[0], dt.datetime) else row[0]
            for index, value in enumerate(row[1:]):
                result[(day, index)] = None if value in (None, "") else float(value)
        return result
    finally:
        workbook.close()


def scan_sheet(handle):
    maxima = {}
    row_count = 0
    for _event, row in etree.iterparse(handle, events=("end",), tag=NS + "row"):
        row_number = int(row.get("r", "0"))
        if row_number >= 2:
            distance = None
            direction = None
            for cell in row:
                reference = cell.get("r", "")
                if reference.startswith("A"):
                    raw_value = cell.findtext(NS + "v")
                    if raw_value not in (None, ""):
                        distance = float(raw_value)
                elif reference.startswith("C"):
                    direction = "".join(cell.itertext())
                if distance is not None and direction is not None:
                    break
            if distance is not None and direction:
                maxima[direction] = max(maxima.get(direction, distance), distance)
                row_count += 1
        row.clear()
        parent = row.getparent()
        while row.getprevious() is not None:
            del parent[0]
    return maxima, row_count


def audit(path: Path, app):
    finals = cached_daily_finals(path, app)
    totals = summary_grid(path)
    comparisons = 0
    detail_rows = 0
    started = time.perf_counter()
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise AssertionError(f"corrupt zip: {path.name}")
        for day in range(1, 31):
            with archive.open(f"xl/worksheets/sheet{day}.xml") as stream:
                maxima, rows = scan_sheet(stream)
            detail_rows += rows
            current_day = dt.date(2026, 6, day)
            for direction_index in app.OPEN_SEA_INDEXES:
                direction = app.DIRECTION_ORDER[direction_index]
                final_value = finals[(current_day, direction_index)]
                total_value = totals[(current_day, direction_index)]
                detail_max = maxima.get(direction)
                for label, left, right in (
                    ("daily final vs total", final_value, total_value),
                    ("detail max vs final", detail_max, final_value),
                ):
                    if left is None or right is None:
                        if left is not None or right is not None:
                            raise AssertionError(
                                f"{path.name} {current_day} {direction} {label}: {left} != {right}"
                            )
                    elif not math.isclose(left, right, rel_tol=0.0, abs_tol=1e-9):
                        raise AssertionError(
                            f"{path.name} {current_day} {direction} {label}: {left} != {right}"
                        )
                comparisons += 1
            print(
                f"{path.name}: day {day}/30 rows={rows:,} elapsed={time.perf_counter()-started:.1f}s",
                flush=True,
            )
    return {
        "file": path.name,
        "detail_rows": detail_rows,
        "directions_compared": comparisons,
        "elapsed_seconds": time.perf_counter() - started,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--repo", type=Path, required=True)
    parser.add_argument("--delivery", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()
    sys.path.insert(0, str(args.repo))
    import ais_monthly_app as app

    reports = []
    for name in ("KLNG 6月 32方位數值_2.xlsx", "KLNG 6月 32方位數值.xlsx"):
        reports.append(audit(args.delivery / name, app))
    payload = {"status": "PASS", "files": reports}
    args.report.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2), flush=True)


if __name__ == "__main__":
    main()
