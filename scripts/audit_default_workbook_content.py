"""Connect default-run workbook cells/chart caches to independently audited selectors.

This is a serialization/decision-propagation gate, NOT a selector oracle. Run the
independent raw census and selector audit over the same completed run separately.
No production calculation or writer imports. Never modifies input workbooks.
"""
from __future__ import annotations

import argparse
import calendar
from datetime import date, datetime
import hashlib
import json
import math
from pathlib import Path
import re
import statistics
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
from openpyxl.utils.cell import get_column_letter, range_boundaries

from scripts.independent_selector_audit import LABELS
from scripts.independent_source_audit import guard_output, save_report
from scripts.calendar_source_plan import source_calendar_plan

C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
S = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
ABBREVIATIONS = ("N", "NbE", "NNE", "NEbN", "NE", "NEbE", "ENE", "EbN", "E", "EbS", "ESE",
                 "SEbE", "SE", "SEbS", "SSE", "SbE", "S", "SbW", "SSW", "SWbS", "SW", "SWbW",
                 "WSW", "WbS", "W", "WbN", "WNW", "NWbW", "NW", "NWbN", "NNW", "NbW")
CANDIDATE_HEADERS = ("Scope", "日期", "方位", "排名", "Candidate ID", "Distance (NM)",
                     "Bearing", "Timestamp", "MMSI", "Channel", "Fragment", "Sheet", "Source row", "Msg type")


def blank(value):
    return value is None or value == ""


def equal(expected, actual):
    if blank(expected) or blank(actual):
        return blank(expected) and blank(actual)
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        return (isinstance(actual, (int, float)) and not isinstance(actual, bool)
                and math.isfinite(expected) and math.isfinite(actual)
                # Installed XlsxWriter XML numbers use16G. Chart caches can
                # retain the original double; normalize both to that exact
                # serialized precision, not a broad absolute tolerance.
                and float(format(expected, ".16G")) == float(format(actual, ".16G")))
    return expected == actual


def day_value(value):
    return value.date() if isinstance(value, datetime) else value


def matrix(sheet, start_row, count):
    rows = list(sheet.iter_rows(min_row=start_row, max_row=start_row + count - 1,
                               min_col=1, max_col=33, values_only=True))
    return [day_value(row[0]) for row in rows], [list(row[1:]) for row in rows]


def stats(values):
    values = [value for value in values if not blank(value)]
    if not values:
        return (None,) * 4
    return max(values), min(values), statistics.mean(values), statistics.stdev(values) if len(values) > 1 else None


def sha256(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        while chunk := stream.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def check_modern_chart_helpers(formula_book, value_book, days, included, expected):
    """Permit NA only after independent grid/statistic/formula/value prerequisites.

    The expected matrix is reconstructed from the separately raw/selector-audited
    decision source, not read back from Excel. Every helper is checked, including
    genuine zero and present values. This is not a general error allowlist.
    """
    issues, entries, permitted = [], [], set()
    if (not days or len(expected) != len(days) or any(len(row) != 32 for row in expected)
            or len(set(included)) != len(included) or any(index not in range(32) for index in included)):
        return ["Modern: helper missing-value prerequisite matrix/coverage invalid"], [], set()
    formulas, values = formula_book["總表"], value_book["總表"]
    last_row = max(len(days) + 9, len(included) + 4)
    formula_rows = list(formulas.iter_rows(min_row=4, max_row=last_row, max_col=38))
    value_rows = list(values.iter_rows(min_row=4, max_row=last_row, max_col=38))
    # A failed or corrupted prerequisite may not be hidden by an intentional NA.
    grid_valid = True
    for position, day in enumerate(days):
        for direction in range(32):
            observed = value_rows[position + 1][direction + 1]
            if observed.data_type == "e" or not equal(expected[position][direction], observed.value):
                grid_valid = False
                issues.append(f"Modern: helper prerequisite grid {day}/{LABELS[direction]}")
    for offset, direction in enumerate(included):
        column = get_column_letter(direction + 2)
        independent_stats = stats([row[direction] for row in expected])
        for stat, function in enumerate(("MAX", "MIN", "AVERAGE")):
            source = f"{column}{len(days) + 6 + stat}"
            source_range = f"{column}5:{column}{len(days) + 4}"
            source_formula = f'=IF(COUNT({source_range})<1,"",{function}({source_range}))'
            helper_formula = f'=IF(COUNT({source})=0,NA(),{source})'
            helper_cell = formula_rows[offset + 1][35 + stat]
            helper_value = value_rows[offset + 1][35 + stat]
            source_cell = formula_rows[len(days) + 2 + stat][direction + 1]
            source_value = value_rows[len(days) + 2 + stat][direction + 1]
            expected_value = independent_stats[stat]
            missing = expected_value is None
            source_ok = (grid_valid and source_cell.data_type == "f" and source_cell.value == source_formula
                         and source_value.data_type != "e" and equal(expected_value, source_value.value))
            helper_ok = (helper_cell.data_type == "f" and helper_cell.value == helper_formula
                         and ((missing and helper_value.data_type == "e" and helper_value.value == "#N/A")
                              or (not missing and helper_value.data_type != "e"
                                  and equal(expected_value, helper_value.value))))
            if not source_ok:
                issues.append(f"Modern: helper prerequisite statistic {stat}/{LABELS[direction]}")
            if not helper_ok:
                issues.append(f"Modern: exact chart-only missing/numeric helper {helper_cell.coordinate}")
            verified = source_ok and helper_ok
            if verified and missing:
                permitted.add(("總表", helper_cell.coordinate))
            entries.append({"sheet": "總表", "cell": helper_cell.coordinate, "direction_index": direction,
                            "direction": LABELS[direction], "statistic": function, "source": source,
                            "source_range": source_range, "source_formula": source_formula,
                            "formula": helper_formula, "expected_value": expected_value,
                            "expected_missing": missing, "source_cached_value": source_value.value,
                            "source_cached_type": source_value.data_type,
                            "cached_value": helper_value.value, "cached_type": helper_value.data_type,
                            "verified": verified})
    return issues, entries, permitted


def check_saved_errors(path, permitted_missing=frozenset()):
    """Scan every saved worksheet error, never just chart/summary ranges."""
    issues, intentional = [], 0
    with zipfile.ZipFile(path) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {node.get("Id"): node.get("Target") for node in relationships}
        for sheet in workbook.findall(S + "sheets/" + S + "sheet"):
            target = targets[sheet.get(R + "id")]
            part = target.lstrip("/") if target.startswith("/") else "xl/" + target
            with archive.open(part) as stream:
                for _event, node in ET.iterparse(stream, events=("end",)):
                    if node.tag == S + "row":
                        node.clear()
                        continue
                    if node.tag != S + "c":
                        continue
                    if node.get("t") == "e":
                        key = sheet.get("name"), node.get("r")
                        value = node.findtext(S + "v")
                        if key in permitted_missing and value == "#N/A" and node.find(S + "f") is not None:
                            intentional += 1
                        else:
                            issues.append(f"{Path(path).name}: unexpected saved error {key}/{value}")
                    node.clear()
    return issues, intentional


def check_modern_radar_helpers(formula_book, value_book, days, included, expected):
    """All day/direction radar helpers, separate from the statistics NA role."""
    issues, entries, permitted = [], [], set()
    if (not days or len(expected) != len(days) or any(len(row) != 32 for row in expected)
            or len(set(included)) != len(included) or any(index not in range(32) for index in included)):
        return ["Modern: radar missing-value prerequisite matrix/coverage invalid"], [], set()
    formulas, values = formula_book["總表"], value_book["總表"]
    formula_rows = list(formulas.iter_rows(min_row=4, max_row=len(days) + 4, max_col=72))
    value_rows = list(values.iter_rows(min_row=4, max_row=len(days) + 4, max_col=72))
    if tuple(cell.value for cell in value_rows[0][40:72]) != LABELS:
        issues.append("Modern: radar helper direction header/order differs")
    for position, day in enumerate(days):
        row = position + 5
        for direction in range(32):
            source = f"{get_column_letter(direction + 2)}{row}"
            coordinate = f"{get_column_letter(direction + 41)}{row}"
            source_formula = (f'=IF(COUNT(\'{day.month}月{day.day}日\'!{get_column_letter(direction + 8)}4)=0,"",'
                              f"'{day.month}月{day.day}日'!{get_column_letter(direction + 8)}4)" if direction in included else None)
            formula = f'=IF(COUNT({source})=0,NA(),{source})'
            source_cell, source_value = formula_rows[position + 1][direction + 1], value_rows[position + 1][direction + 1]
            helper_cell, helper_value = formula_rows[position + 1][direction + 40], value_rows[position + 1][direction + 40]
            wanted = expected[position][direction]
            missing = wanted is None
            source_ok = (source_value.data_type != "e" and equal(wanted, source_value.value)
                         and ((direction in included and source_cell.data_type == "f" and source_cell.value == source_formula)
                              or (direction not in included and blank(source_cell.value))))
            helper_ok = (helper_cell.data_type == "f" and helper_cell.value == formula
                         and ((missing and helper_value.data_type == "e" and helper_value.value == "#N/A")
                              or (not missing and helper_value.data_type != "e" and equal(wanted, helper_value.value))))
            if not source_ok:
                issues.append(f"Modern: radar helper prerequisite {day}/{LABELS[direction]}")
            if not helper_ok:
                issues.append(f"Modern: exact radar missing/numeric helper {coordinate}")
            verified = source_ok and helper_ok
            if verified and missing:
                permitted.add(("總表", coordinate))
            entries.append({"sheet": "總表", "cell": coordinate, "day": day.isoformat(),
                            "direction_index": direction, "direction": LABELS[direction], "source": source,
                            "source_formula": source_formula, "formula": formula, "expected_value": wanted,
                            "expected_missing": missing, "source_cached_value": source_value.value,
                            "source_cached_type": source_value.data_type, "cached_value": helper_value.value,
                            "cached_type": helper_value.data_type, "verified": verified})
    return issues, entries, permitted


def check_statistics(sheet, start_row, values, issues, label):
    actual = list(sheet.iter_rows(min_row=start_row, max_row=start_row + 3,
                                 min_col=2, max_col=33, values_only=True))
    for index in range(32):
        expected = stats([row[index] for row in values])
        for offset in range(4):
            if not equal(expected[offset], actual[offset][index]):
                issues.append(f"{label}: statistic {offset}/{LABELS[index]}")


def check_chart_caches(path, workbook, permitted_missing=frozenset()):
    """Every numeric-reference point must match its actual saved cell, including blanks."""
    issues = []
    points = 0
    with zipfile.ZipFile(path) as archive:
        charts = [name for name in archive.namelist() if re.fullmatch(r"xl/charts/chart\d+\.xml", name)]
        if len(charts) != 2:
            issues.append(f"{path.name}: expected two charts")
        for name in charts:
            root = ET.fromstring(archive.read(name))
            series = list(root.iter(C + "ser"))
            references = root.findall(".//" + C + "val/" + C + "numRef")
            if not series or not references or len(references) != len(series):
                issues.append(f"{path.name}/{name}: every native series requires one cell-backed numeric value reference")
            bars = root.findall(".//" + C + "barChart/" + C + "ser")
            lines = root.findall(".//" + C + "lineChart/" + C + "ser")
            if bars and (len(bars) != 2 or len(lines) != 1 or len(series) != 3):
                issues.append(f"{path.name}/{name}: statistics chart must contain two bars and one mean line")
            for reference in references:
                formula = reference.findtext(C + "f", "")
                match = re.fullmatch(r"(?:'((?:[^']|'')+)'|([^!]+))!(\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?)", formula)
                if not match:
                    issues.append(f"{path.name}/{name}: unsupported numeric reference {formula}")
                    continue
                sheet_name = (match.group(1) or match.group(2)).replace("''", "'")
                bounds = range_boundaries(match.group(3))
                if bounds[0] != bounds[2] and bounds[1] != bounds[3]:
                    issues.append(f"{path.name}/{name}: non-vector numeric reference")
                    continue
                source_cells = [cell for row in workbook[sheet_name].iter_rows(
                    min_col=bounds[0], min_row=bounds[1], max_col=bounds[2], max_row=bounds[3]) for cell in row]
                source = [None if (cell.data_type == "e" and cell.value == "#N/A"
                                   and (sheet_name, cell.coordinate) in permitted_missing) else cell.value
                          for cell in source_cells]
                cache = reference.find(C + "numCache")
                if cache is None:
                    issues.append(f"{path.name}/{name}: missing numeric cache")
                    continue
                count = cache.find(C + "ptCount")
                entries = cache.findall(C + "pt")
                indexes = [int(point.get("idx", "-1")) for point in entries]
                if count is None or int(count.get("val", "-1")) != len(source) or len(indexes) != len(set(indexes)):
                    issues.append(f"{path.name}/{name}: numeric-cache shape/index mismatch")
                cached = {int(point.get("idx", "-1")): point.findtext(C + "v") for point in entries}
                if any(index < 0 or index >= len(source) for index in cached):
                    issues.append(f"{path.name}/{name}: out-of-range cache point")
                for index, value in enumerate(source):
                    raw = cached.get(index)
                    try:
                        actual = None if blank(raw) else float(raw)
                    except (TypeError, ValueError):
                        actual = raw
                    points += 1
                    if not equal(value, actual):
                        issues.append(f"{path.name}/{name}: cache/cell {formula} index {index}")
    return issues, points


def reference_key(formula):
    """Canonicalize quoting/$ only, never permit a different sheet or extent."""
    match = re.fullmatch(r"(?:'((?:[^']|'')+)'|([^!]+))!(\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?)", formula or "")
    if not match:
        return None
    return (match.group(1) or match.group(2)).replace("''", "'"), range_boundaries(match.group(3))


def check_modern_bindings(path, formula_book, days, included, expected):
    """Independently check intended Modern sources, not merely cache self-consistency.

    Coordinates follow the user-visible workbook contract: dates start at row5,
    32 compass headers at row4, daily finals at row4, and AI:AL helper vectors.
    No production chart/formula helpers are imported.
    """
    issues = []
    with zipfile.ZipFile(path) as archive:
        roots = [ET.fromstring(archive.read(name)) for name in archive.namelist()
                 if re.fullmatch(r"xl/charts/chart\d+\.xml", name)]
    radars = [root.find(".//" + C + "radarChart") for root in roots]
    radars = [node for node in radars if node is not None]
    combos = [root for root in roots if root.find(".//" + C + "barChart") is not None]
    if len(roots) != 2 or len(radars) != 1 or len(combos) != 1:
        return ["Modern: expected one radar and one bar/mean combination chart"]

    def verify_series(series, wanted, label):
        if len(series) != len(wanted):
            issues.append(f"Modern: {label} series count/order differs")
        for node, (name, categories, values) in zip(series, wanted):
            actual_name = node.findtext(C + "tx/" + C + "v")
            category = node.find(C + "cat")
            actual_cat = None if category is None else next(iter(category.findall(".//" + C + "f")), None)
            actual_val = node.findtext(C + "val/" + C + "numRef/" + C + "f")
            if actual_name != name or reference_key(None if actual_cat is None else actual_cat.text) != reference_key(categories) or reference_key(actual_val) != reference_key(values):
                issues.append(f"Modern: {label} intended name/category/value binding {name}")

    wanted_radar = [(f"{day.month}/{day.day}", "總表!B4:AG4", f"總表!AO{index + 5}:BT{index + 5}")
                    for index, day in enumerate(days) if any(not blank(value) for value in expected[index])]
    if not wanted_radar and days:
        wanted_radar = [("本時段無有效資料", "總表!B4:AG4", "總表!AO5:BT5")]
    verify_series(radars[0].findall(C + "ser"), wanted_radar, "radar")
    radar_root = next(root for root in roots if root.find(".//" + C + "radarChart") is not None)
    visibility = radar_root.find(".//" + C + "plotVisOnly")
    # Installed XlsxWriter show_hidden_data omits this element. Explicit true
    # is forbidden; a separate native gate must prove PlotVisibleOnly=False.
    if visibility is not None and visibility.get("val") != "0":
        issues.append("Modern: hidden radar helpers cannot have plotVisOnly=1/invalid")
    combo = combos[0]
    category = f"總表!AI5:AI{len(included) + 4}"
    verify_series(combo.findall(".//" + C + "barChart/" + C + "ser"),
                  [("最遠", category, f"總表!AJ5:AJ{len(included) + 4}"),
                   ("最近", category, f"總表!AK5:AK{len(included) + 4}")], "statistics bars")
    verify_series(combo.findall(".//" + C + "lineChart/" + C + "ser"),
                  [("平均", category, f"總表!AL5:AL{len(included) + 4}")], "statistics mean")
    total = formula_book["總表"]
    rows = list(total.iter_rows(min_row=4, max_row=max(len(days) + 9, len(included) + 4),
                                min_col=1, max_col=38, values_only=True))
    included_set = set(included)
    for day_index, day in enumerate(days):
        sheet_name = f"{day.month}月{day.day}日"
        daily = next(formula_book[sheet_name].iter_rows(min_row=4, max_row=4,
                                                      min_col=8, max_col=39, values_only=True))
        for index in range(32):
            if index not in included_set:
                if not blank(rows[day_index + 1][index + 1]) or not blank(daily[index]):
                    issues.append(f"Modern: excluded direction formula {day}/{LABELS[index]}")
                continue
            ledger_row = 2 + day_index * 3 * len(included) + included.index(index)
            wanted_daily = f"='決策台帳'!H{ledger_row}"
            source = f"'{sheet_name}'!{get_column_letter(index + 8)}4"
            wanted_summary = f'=IF(COUNT({source})=0,"",{source})'
            if daily[index] != wanted_daily or rows[day_index + 1][index + 1] != wanted_summary:
                issues.append(f"Modern: final/summary formula source {day}/{LABELS[index]}")
    for offset, function in enumerate(("MAX", "MIN", "AVERAGE", "_xlfn.STDEV.S")):
        for index in included:
            column = get_column_letter(index + 2)
            source = f"{column}5:{column}{len(days) + 4}"
            wanted = f'=IF(COUNT({source})<{2 if offset == 3 else 1},"",{function}({source}))'
            if rows[len(days) + 2 + offset][index + 1] != wanted:
                issues.append(f"Modern: guarded statistic formula {offset}/{LABELS[index]}")
    for offset, index in enumerate(included):
        helper = rows[offset + 1]
        if helper[34] != LABELS[index]:
            issues.append(f"Modern: helper compass order {offset}")
        for stat in range(3):
            source = f"{get_column_letter(index + 2)}{len(days) + 6 + stat}"
            wanted = f'=IF(COUNT({source})=0,NA(),{source})'
            if helper[35 + stat] != wanted:
                issues.append(f"Modern: helper statistic formula {stat}/{LABELS[index]}")
    return issues


def check_historical_bindings(path, workbook, days, included, expected, *, use_radar_helpers=False):
    """Exact retained-day radar and reference-ordered statistic source contract."""
    issues = []
    with zipfile.ZipFile(path) as archive:
        roots = [ET.fromstring(archive.read(name)) for name in archive.namelist()
                 if re.fullmatch(r"xl/charts/chart\d+\.xml", name)]
    radars = [root.find(".//" + C + "radarChart") for root in roots]
    radars = [node for node in radars if node is not None]
    combos = [root for root in roots if root.find(".//" + C + "barChart") is not None]
    if len(roots) != 2 or len(radars) != 1 or len(combos) != 1:
        return [f"{path.name}: expected one radar and one bar/mean combination"]
    first, last = ("AO", "BT") if use_radar_helpers else ("B", "AG")
    wanted = [(f"{day.month}月{day.day}日", "總表!B1:AG1", f"總表!{first}{index + 2}:{last}{index + 2}")
              for index, day in enumerate(days) if any(not blank(value) for value in expected[index])]
    if not wanted and days:
        wanted = [("本時段無有效資料", "總表!B1:AG1", f"總表!{first}2:{last}2")]

    def verify(nodes, plan, label):
        if len(nodes) != len(plan):
            issues.append(f"{path.name}: {label} series coverage/order")
        for node, (name, category, values) in zip(nodes, plan):
            actual_category = node.find(C + "cat")
            category_node = None if actual_category is None else next(iter(actual_category.findall(".//" + C + "f")), None)
            if (node.findtext(C + "tx/" + C + "v") != name
                    or reference_key(None if category_node is None else category_node.text) != reference_key(category)
                    or reference_key(node.findtext(C + "val/" + C + "numRef/" + C + "f")) != reference_key(values)):
                issues.append(f"{path.name}: {label} intended name/category/value binding {name}")
    verify(radars[0].findall(C + "ser"), wanted, "radar")
    if use_radar_helpers:
        radar_root = next(root for root in roots if root.find(".//" + C + "radarChart") is not None)
        visibility = radar_root.find(".//" + C + "plotVisOnly")
        if visibility is not None and visibility.get("val") != "0":
            issues.append(f"{path.name}: hidden radar helpers cannot have plotVisOnly=1/invalid")
    category = f"總表!AI2:AI{len(included) + 1}"
    verify(combos[0].findall(".//" + C + "barChart/" + C + "ser"),
           [("MAX", category, f"總表!AJ2:AJ{len(included) + 1}"),
            ("MIN", category, f"總表!AK2:AK{len(included) + 1}")], "statistics bars")
    verify(combos[0].findall(".//" + C + "lineChart/" + C + "ser"),
           [("平均值", category, f"總表!AL2:AL{len(included) + 1}")], "statistics mean")
    sheet = workbook["總表"]
    headers = next(sheet.iter_rows(min_row=1, max_row=1, min_col=2, max_col=33, values_only=True))
    if tuple(headers) != ABBREVIATIONS:
        issues.append(f"{path.name}: compass header/order")
    # The existing accepted KLNG statistics rotate WSW..NbW before N..ESE;
    # this order is deliberately different from the visible B:AG grid.
    order = list(range(22, 32)) + list(range(11)) if included == list(range(11)) + list(range(22, 32)) else included
    helpers = list(sheet.iter_rows(min_row=2, max_row=len(included) + 1,
                                   min_col=35, max_col=38, values_only=True))
    for offset, direction in enumerate(order):
        wanted_stats = stats([row[direction] for row in expected])[:3]
        if helpers[offset][0] != ABBREVIATIONS[direction] or any(
                not equal(wanted, actual) for wanted, actual in zip(wanted_stats, helpers[offset][1:])):
            issues.append(f"{path.name}: independent helper mapping/statistics {ABBREVIATIONS[direction]}")
    return issues


def check_candidate_serialization(workbook, scopes, days, included):
    """Every saved display row/column, with a separate selector-source prerequisite."""
    sheet = workbook["候選清單"]
    issues = []
    headers = next(sheet.iter_rows(min_row=1, max_row=1, max_col=14, values_only=True))
    if tuple(headers) != CANDIDATE_HEADERS:
        issues.append("Modern: candidate headers/order differ")
    count = sum(len(scopes[day][scope][LABELS[index]]["candidates"])
                for day in days for scope in ("logical_day", "period_a", "period_b") for index in included)
    if sheet.max_row != max(1, count + 1):
        issues.append("Modern: candidate row coverage differs")
    rows = iter(sheet.iter_rows(min_row=2, max_row=count + 1, max_col=14, values_only=True))
    position = 0
    for day in days:
        for scope in ("logical_day", "period_a", "period_b"):
            for index in included:
                for candidate in scopes[day][scope][LABELS[index]]["candidates"]:
                    position += 1
                    wanted = (scope, day, LABELS[index], candidate["rank"], candidate["candidate_id"],
                              candidate["distance"], candidate["bearing"],
                              None if not candidate["timestamp"] else datetime.fromisoformat(candidate["timestamp"]),
                              candidate["mmsi"], candidate["channel"], candidate["fragment"],
                              candidate["sheet"], candidate["source_row"], candidate["msg_type"])
                    observed = list(next(rows))
                    observed[1] = day_value(observed[1])
                    if any(not equal(left, right) for left, right in zip(wanted, observed)):
                        issues.append(f"Modern: candidate serialization row {position + 1}")
    return issues, count


def check_ledger_final_formulas(workbook, days, included, candidate_count):
    issues = []
    end = max(2, candidate_count + 1)
    rows = list(workbook["決策台帳"].iter_rows(min_row=2, max_col=8, values_only=True))
    expected_scopes = [scope for _day in days for scope in ("logical_day", "period_a", "period_b")
                       for _index in included]
    for row_number, (row, scope) in enumerate(zip(rows, expected_scopes), 2):
        lookup = (f"IFERROR(INDEX('候選清單'!$F$2:$F${end},"
                  f"MATCH(F{row_number},'候選清單'!$E$2:$E${end},0)),\"#INVALID_CANDIDATE\")")
        wanted = (f'=IF(ISNUMBER(G{row_number}),G{row_number},IF(F{row_number}="","",{lookup}))'
                  if scope == "logical_day" else
                  f'=IF(G{row_number}<>"","#INVALID_PERIOD_NUMERIC",IF(F{row_number}="","",{lookup}))')
        if row[7] != wanted:
            issues.append(f"Modern: ledger intended Final formula row {row_number}")
    return issues


def file_stamps(paths):
    result = []
    for path in sorted(paths, key=lambda value: str(value.resolve()).casefold()):
        info = path.stat()
        result.append({"path": str(path.resolve()), "size": info.st_size, "modified_ns": info.st_mtime_ns})
    return result


def build_modern_missing_contract(job, modern, days, included, expected, cache_files,
                                  bound_hashes, helper_entries, intentional_missing_count,
                                  radar_entries, radar_missing_count, issues):
    return {"schema": 1, "state": "PASS_SERIALIZATION_PREREQUISITES" if not issues else "FAIL",
            "role": "Modern/default", "token": job["token"], "port": job["port"],
            "workbook": str(modern.resolve()), "workbook_sha256": bound_hashes[str(modern.resolve())],
            "cache_sources": [{"path": str(path.resolve()), "sha256": bound_hashes[str(path.resolve())]}
                              for path in cache_files],
            "logical_days": [day.isoformat() for day in days], "direction_labels": list(LABELS),
            "included_indices": included, "expected_matrix": expected, "helpers": helper_entries,
            "verified_intentional_missing_count": intentional_missing_count,
            "radar_helpers": radar_entries, "radar_verified_intentional_missing_count": radar_missing_count,
            "raw_selector_prerequisite": "Same-root independent raw schema3 and selector schema3 PASS REQUIRED separately; NOT source truth by this gate alone"}


def build_historical_chart_contract(job, path, role, scope, days, included, expected,
                                    cache_files, bound_hashes, formula_book, value_book):
    """Independent all-point/native handoff, including every formula prerequisite.

    The matrix comes from the decision caches, never the workbook's own chart
    cache. The caller marks all five contracts FAIL if ANY six-book job check
    fails. Literal blank helper cells get no Modern NA permission.
    """
    issues = []
    values = expected[scope]
    order = list(range(22, 32)) + list(range(11)) if included == list(range(11)) + list(range(22, 32)) else list(included)
    statistic_matrix = [list(stats([row[index] for row in values])[:3]) for index in order]
    entries = []
    source_formulas = formula_book["總表"]
    source_values = value_book["總表"]
    formula_rows = list(source_formulas.iter_rows(min_row=2, max_row=len(days) + 6, max_col=38))
    value_rows = list(source_values.iter_rows(min_row=2, max_row=len(days) + 6, max_col=38))
    for position, day in enumerate(days):
        daily_formulas = daily_values = None
        if role in ("full_a", "full_b"):
            sheet_name = f"{day.month}月{day.day}日"
            daily_formulas = next(formula_book[sheet_name].iter_rows(min_row=2, max_row=2, min_col=8, max_col=39))
            daily_values = next(value_book[sheet_name].iter_rows(min_row=2, max_row=2, min_col=8, max_col=39))
        for index in range(32):
            column = get_column_letter(index + 2)
            coordinate = f"{column}{position + 2}"
            wanted = values[position][index]
            cell, cached = formula_rows[position][index + 1], value_rows[position][index + 1]
            if index not in included or role in ("summary_a", "summary_b"):
                formula = None
                valid = cell.data_type != "f" and cached.data_type != "e" and equal(wanted, cached.value)
            elif role == "integrated":
                sources = f"'Period A'!{coordinate},'Period B'!{coordinate}"
                formula = f'=IF(COUNT({sources})=0,"",MAX({sources}))'
                valid = cell.data_type == "f" and cell.value == formula and cached.data_type != "e" and equal(wanted, cached.value)
            else:
                source = f"'{day.month}月{day.day}日'!{get_column_letter(index + 8)}2"
                formula = f'=IF(COUNT({source})=0,"",{source})'
                valid = cell.data_type == "f" and cell.value == formula and cached.data_type != "e" and equal(wanted, cached.value)
            if daily_formulas is not None:
                daily = daily_formulas[index]
                daily_value = daily_values[index]
                daily_column = get_column_letter(index + 8)
                daily_formula = (f'=IF(COUNTIF($C:$C,{daily_column}$1)=0,"",_xlfn.MAXIFS($A:$A,$C:$C,{daily_column}$1))'
                                 if index in included else None)
                valid = valid and daily_value.data_type != "e" and equal(wanted, daily_value.value) and (
                    daily.data_type == "f" and daily.value == daily_formula if daily_formula is not None else blank(daily.value))
                entries.append({"sheet": f"{day.month}月{day.day}日", "cell": daily_column + "2",
                                "formula": daily_formula, "expected_value": wanted, "verified": valid})
            if not valid:
                issues.append(f"{path.name}: Historical primary/formula prerequisite {day}/{ABBREVIATIONS[index]}")
            entries.append({"sheet": "總表", "cell": coordinate, "formula": formula,
                            "expected_value": wanted, "verified": valid})
    if role == "integrated":
        for sheet_name, period in (("Period A", "period_a"), ("Period B", "period_b")):
            formula_sources = list(formula_book[sheet_name].iter_rows(min_row=2, max_row=len(days) + 1, min_col=2, max_col=33))
            value_sources = list(value_book[sheet_name].iter_rows(min_row=2, max_row=len(days) + 1, min_col=2, max_col=33))
            for position, day in enumerate(days):
                for index in range(32):
                    cell, cached = formula_sources[position][index], value_sources[position][index]
                    valid = cell.data_type != "f" and cached.data_type != "e" and equal(expected[period][position][index], cached.value)
                    if not valid:
                        issues.append(f"{path.name}: integrated {sheet_name} source prerequisite {day}/{ABBREVIATIONS[index]}")
                    entries.append({"sheet": sheet_name, "cell": f"{get_column_letter(index + 2)}{position + 2}", "formula": None,
                                    "expected_value": expected[period][position][index], "verified": valid})
    for stat, function in enumerate(("MAX", "MIN", "AVERAGE", "_xlfn.STDEV.S")):
        for index in range(32):
            column = get_column_letter(index + 2)
            source = f"{column}2:{column}{len(days) + 1}"
            coordinate = f"{column}{len(days) + 3 + stat}"
            formula = f'=IF(COUNT({source})<{2 if stat == 3 else 1},"",{function}({source}))' if index in included else None
            cell, cached = formula_rows[len(days) + 1 + stat][index + 1], value_rows[len(days) + 1 + stat][index + 1]
            wanted = stats([row[index] for row in values])[stat]
            valid = cached.data_type != "e" and equal(wanted, cached.value) and (
                cell.data_type == "f" and cell.value == formula if formula is not None else blank(cell.value))
            if not valid:
                issues.append(f"{path.name}: Historical guarded statistic prerequisite {function}/{ABBREVIATIONS[index]}")
            entries.append({"sheet": "總表", "cell": coordinate, "formula": formula,
                            "expected_value": wanted, "verified": valid})
    first, last = ("AO", "BT") if role in ("full_a", "full_b", "integrated") else ("B", "AG")
    radar = [{"name": f"{day.month}月{day.day}日", "categories": "總表!B1:AG1",
              "values": f"總表!{first}{position + 2}:{last}{position + 2}", "expected_values": values[position]}
             for position, day in enumerate(days) if any(not blank(value) for value in values[position])]
    if not radar and days:
        radar = [{"name": "本時段無有效資料", "categories": "總表!B1:AG1",
                  "values": f"總表!{first}2:{last}2", "expected_values": values[0]}]
    statistic_series = [{"name": name, "categories": f"總表!AI2:AI{len(included) + 1}",
                         "values": f"總表!{column}2:{column}{len(included) + 1}",
                         "expected_values": [row[position] for row in statistic_matrix]}
                        for position, (name, column) in enumerate((("MAX", "AJ"), ("MIN", "AK"), ("平均值", "AL")))]
    contract = {"schema": 1, "state": "PENDING_WHOLE_JOB", "role": role, "scope": scope,
                "token": job["token"], "port": job["port"], "workbook": str(path.resolve()),
                "workbook_sha256": bound_hashes[str(path.resolve())],
                "cache_sources": [{"path": str(item.resolve()), "sha256": bound_hashes[str(item.resolve())]} for item in cache_files],
                "logical_days": [day.isoformat() for day in days], "direction_labels": list(ABBREVIATIONS),
                "included_indices": list(included), "expected_matrix": values,
                "statistic_order": order, "statistic_matrix": statistic_matrix,
                "source_cells": entries, "radar_series": radar, "statistic_series": statistic_series,
                "verified_source_cell_count": sum(entry["verified"] for entry in entries),
                "raw_selector_prerequisite": "Same-root independent raw schema3 and selector schema3 PASS REQUIRED separately; NOT source truth by this gate alone"}
    return contract, issues


def check_historical_radar_helpers(formula_book, value_book, contract):
    """Only formula-backed primary roles may use independently proven NA helpers."""
    if contract["role"] not in ("full_a", "full_b", "integrated"):
        return [], [], set()
    issues, entries, permitted = [], [], set()
    days, expected = contract["logical_days"], contract["expected_matrix"]
    if not days or len(expected) != len(days) or any(len(row) != 32 for row in expected):
        return ["Historical: radar helper prerequisite matrix/coverage invalid"], [], set()
    formulas, values = formula_book["總表"], value_book["總表"]
    formula_rows = list(formulas.iter_rows(min_row=1, max_row=len(days) + 1, max_col=72))
    value_rows = list(values.iter_rows(min_row=1, max_row=len(days) + 1, max_col=72))
    if tuple(cell.value for cell in value_rows[0][40:72]) != ABBREVIATIONS:
        issues.append("Historical: radar helper direction header/order differs")
    prerequisite_valid = all(entry["verified"] for entry in contract["source_cells"])
    for position, day in enumerate(days):
        row = position + 2
        for direction in range(32):
            source = f"{get_column_letter(direction + 2)}{row}"
            coordinate = f"{get_column_letter(direction + 41)}{row}"
            formula = f'=IF(COUNT({source})=0,NA(),{source})'
            source_cell, source_value = formula_rows[position + 1][direction + 1], value_rows[position + 1][direction + 1]
            helper_cell, helper_value = formula_rows[position + 1][direction + 40], value_rows[position + 1][direction + 40]
            wanted = expected[position][direction]
            missing = wanted is None
            source_ok = prerequisite_valid and source_value.data_type != "e" and equal(wanted, source_value.value)
            helper_ok = (helper_cell.data_type == "f" and helper_cell.value == formula
                         and ((missing and helper_value.data_type == "e" and helper_value.value == "#N/A")
                              or (not missing and helper_value.data_type != "e" and equal(wanted, helper_value.value))))
            if not source_ok:
                issues.append(f"Historical: radar helper prerequisite {day}/{ABBREVIATIONS[direction]}")
            if not helper_ok:
                issues.append(f"Historical: exact chart-only missing/numeric helper {coordinate}")
            verified = source_ok and helper_ok
            if verified and missing:
                permitted.add(("總表", coordinate))
            entries.append({"sheet": "總表", "cell": coordinate, "day": day,
                            "direction_index": direction, "direction": ABBREVIATIONS[direction], "source": source,
                            "source_formula": source_cell.value if source_cell.data_type == "f" else None,
                            "formula": formula, "expected_value": wanted, "expected_missing": missing,
                            "source_cached_value": source_value.value, "source_cached_type": source_value.data_type,
                            "cached_value": helper_value.value, "cached_type": helper_value.data_type,
                            "verified": verified})
    return issues, entries, permitted


def audit_job(job):
    modern = Path(job["modern"])
    day_strings = job["logical_days"]
    source_days,days = source_calendar_plan(int(job['year']),int(job['month']),day_strings)
    included = list(range(11)) + list(range(22, 32)) if job["port"] == "KLNG" else list(range(32))
    caches = list(modern.parent.rglob("*_cache"))
    if len(caches) != 1:
        raise ValueError("exactly one cache directory required")
    cache_files = sorted(caches[0].glob("*.json"))
    if {path.stem for path in cache_files} != {day.strftime("%Y%m%d") for day in source_days}:
        raise ValueError("cache day coverage differs from complete plan")
    files = sorted(Path(job["delivery"]).glob("*.xlsx"))
    input_files = [modern, *files, *cache_files]
    cache_before = file_stamps(cache_files)
    scopes = {}
    for day in source_days:
        payload = json.loads((caches[0] / (day.strftime("%Y%m%d") + ".json")).read_text(encoding="utf-8"))
        result = payload["result"]
        if result["day"] != day.isoformat() or payload["signature"]["port"] != job["port"]:
            raise ValueError("cache day/port mismatch")
        scopes[day] = {"logical_day": result["directions"], **result["period_directions"]}
    # Missing-source calendar rows have no invented cache, candidates or identity.
    # All workbook/ledger/helper cells for those days remain independently checked.
    for day in set(days)-set(source_days):
        scopes[day] = {scope:{label:{"selected":None,"selected_candidate_id":None,"candidates":[]}
                             for label in LABELS}
                       for scope in ("logical_day","period_a","period_b")}
    if file_stamps(cache_files) != cache_before:
        raise ValueError("cache changed while expected serialization values were read")
    before = file_stamps(input_files)
    bound_hashes = {str(path.resolve()): sha256(path) for path in input_files}
    expected = {scope: [[scopes[day][scope][direction]["selected"] if index in included else None
                         for index, direction in enumerate(LABELS)] for day in days]
                for scope in ("logical_day", "period_a", "period_b")}
    expected["integrated"] = [[max(value for value in (a, b) if not blank(value))
                               if not (blank(a) and blank(b)) else None
                               for a, b in zip(left, right)]
                              for left, right in zip(expected["period_a"], expected["period_b"])]
    issues = []
    points = 0
    wb = openpyxl.load_workbook(modern, read_only=True, data_only=True)
    formula_book = openpyxl.load_workbook(modern, read_only=True, data_only=False)
    try:
        headers = next(wb["總表"].iter_rows(min_row=4, max_row=4, min_col=2, max_col=33, values_only=True))
        if tuple(headers) != LABELS:
            issues.append("Modern: 32 direction header/order mismatch")
        actual_days, actual = matrix(wb["總表"], 5, len(days))
        if actual_days != days:
            issues.append("Modern: dates differ from plan")
        for row, (wanted, observed) in enumerate(zip(expected["logical_day"], actual)):
            for index in range(32):
                if not equal(wanted[index], observed[index]):
                    issues.append(f"Modern: daily summary {days[row]}/{LABELS[index]}")
        check_statistics(wb["總表"], len(days) + 6, expected["logical_day"], issues, "Modern")
        ledger = list(wb["決策台帳"].iter_rows(min_row=2, max_col=10, values_only=True))
        wanted_keys = [(scope, day, LABELS[index]) for day in days for scope in
                       ("logical_day", "period_a", "period_b") for index in included]
        wanted_key_set = set(wanted_keys)
        if [(row[0], day_value(row[1]), row[2]) for row in ledger] != wanted_keys:
            issues.append("Modern: ledger coverage/order differs")
        for row in ledger:
            key = (row[0], day_value(row[1]), row[2])
            if key not in wanted_key_set:
                continue
            selected = scopes[key[1]][key[0]][key[2]]
            candidate_id = selected["selected_candidate_id"]
            if not (equal(candidate_id, row[3]) and equal(candidate_id, row[5])
                    and blank(row[6]) and equal(selected["selected"], row[4]) and equal(selected["selected"], row[7])):
                issues.append(f"Modern: default decision propagation {key}")
        for day in days:
            sheet = wb[f"{day.month}月{day.day}日"]
            daily = list(sheet.iter_rows(min_row=2, max_row=4, min_col=8, max_col=39, values_only=True))
            for index, direction in enumerate(LABELS):
                value = scopes[day]["logical_day"][direction]["selected"] if index in included else None
                if not equal(value, daily[0][index]) or not equal(value, daily[2][index]):
                    issues.append(f"Modern: day sheet {day}/{direction}")
        helper_issues, helper_entries, permitted_missing = check_modern_chart_helpers(
            formula_book, wb, days, included, expected["logical_day"])
        issues.extend(helper_issues)
        radar_issues, radar_entries, radar_missing = check_modern_radar_helpers(
            formula_book, wb, days, included, expected["logical_day"])
        issues.extend(radar_issues)
        error_issues, total_intentional_count = check_saved_errors(modern, permitted_missing | radar_missing)
        issues.extend(error_issues)
        if total_intentional_count != len(permitted_missing | radar_missing):
            issues.append("Modern: saved intentional error coverage differs from exact verified helper roles")
        chart_issues, points = check_chart_caches(modern, wb, permitted_missing | radar_missing)
        issues.extend(chart_issues)
        candidate_issues, candidate_count = check_candidate_serialization(wb, scopes, days, included)
        issues.extend(candidate_issues)
        issues.extend(check_modern_bindings(modern, formula_book, days, included, expected["logical_day"]))
        issues.extend(check_ledger_final_formulas(formula_book, days, included, candidate_count))
    finally:
        wb.close()
        formula_book.close()
    if len(files) != 5:
        issues.append("Historical: expected five files")
    stem = f"{job['port']} {days[0].month}月 32方位數值"
    roles = {stem + suffix + ".xlsx": role for suffix, role in (("_2", "full_a"), ("", "full_b"),
                                                               ("_2 總表", "summary_a"), (" 總表", "summary_b"))}
    roles[f"{job['port']}_{days[0].month}月_32方位_每日較大值整合總表.xlsx"] = "integrated"
    wanted_names = set(roles)
    if {path.name for path in files} != wanted_names:
        issues.append("Historical: exact five role filenames/coverage differ")
    historical_contracts = []
    for path in files:
        scope = "integrated" if path.stem.endswith("每日較大值整合總表") else (
            "period_a" if re.search(r"_2(?: 總表)?$", path.stem) else "period_b")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        historical_formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            actual_days, actual = matrix(wb["總表"], 2, len(days))
            if actual_days != days:
                issues.append(f"{path.name}: dates differ from plan")
            for row, (wanted, observed) in enumerate(zip(expected[scope], actual)):
                for index in range(32):
                    if not equal(wanted[index], observed[index]):
                        issues.append(f"{path.name}: selector-to-final {days[row]}/{LABELS[index]}")
            check_statistics(wb["總表"], len(days) + 3, expected[scope], issues, path.name)
            historical_missing = set()
            role = roles.get(path.name)
            if role is not None:
                historical_contract, historical_issues = build_historical_chart_contract(
                    job, path, role, scope, days, included, expected, cache_files,
                    bound_hashes, historical_formulas, wb)
                issues.extend(historical_issues)
                helper_issues, historical_radar_entries, historical_missing = check_historical_radar_helpers(
                    historical_formulas, wb, historical_contract)
                issues.extend(helper_issues)
                historical_contract.update(radar_helpers=historical_radar_entries,
                                           radar_verified_intentional_missing_count=len(historical_missing))
                historical_contracts.append(historical_contract)
            error_issues, intentional_count = check_saved_errors(path, historical_missing)
            issues.extend(error_issues)
            if intentional_count != len(historical_missing):
                issues.append(f"{path.name}: exact Historical intentional missing error census differs")
            issues.extend(check_historical_bindings(path, wb, days, included, expected[scope],
                                                   use_radar_helpers=role in ("full_a", "full_b", "integrated")))
            chart_issues, count = check_chart_caches(path, wb, historical_missing)
            points += count
            issues.extend(chart_issues)
        finally:
            wb.close()
            historical_formulas.close()
    if file_stamps(input_files) != before or sorted(Path(job["delivery"]).glob("*.xlsx")) != files:
        issues.append("inputs/outputs changed during workbook audit")
    if any(sha256(Path(path)) != digest for path, digest in bound_hashes.items()):
        issues.append("workbook/cache hashes changed during chart-contract audit")
    for historical_contract in historical_contracts:
        historical_contract["state"] = "PASS_SERIALIZATION_PREREQUISITES" if not issues else "FAIL"
    contract = build_modern_missing_contract(job, modern, days, included, expected["logical_day"], cache_files,
                                             bound_hashes, helper_entries, len(permitted_missing),
                                             radar_entries, len(radar_missing), issues)
    return {"token": job["token"], "days": len(days), "workbooks": len(files) + 1,
            "chart_numeric_points": points, "candidate_rows_checked": candidate_count,
            "modern_chart_missing_contract": contract,
            "historical_chart_contracts": historical_contracts,
            "input_stamps": before, "issues": issues, "pass": not issues}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("root", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args(argv)
    from scripts.validate_full_regression import audit_run_scope
    run = json.loads((args.root / "run_report.json").read_text(encoding="utf-8"))
    protected_roots = [args.root, Path(run["source_root"])]
    protected_files = [args.root / "run_report.json"]
    guard_output(args.output, protected_roots, protected_files)
    scope = audit_run_scope(run, Path(run["source_root"]))
    report = {"root": str(args.root.resolve()), "scope": scope, "jobs": [],
              "boundary": "Default-run serialization gate; same-root independent raw and selector PASS required separately"}
    if scope["complete"]:
        for job in run["jobs"]:
            try:
                result = audit_job(job)
            except Exception as error:
                result = {"token": job["token"], "pass": False, "issues": [f"{type(error).__name__}: {error}"]}
            report["jobs"].append(result)
            save_report(args.output, report, protected_roots, protected_files)
    report["pass"] = scope["complete"] and len(report["jobs"]) == scope["expected_job_count"] and all(job["pass"] for job in report["jobs"])
    save_report(args.output, report, protected_roots, protected_files)
    print(json.dumps({"pass": report["pass"], "jobs": len(report["jobs"]), "issues": sum(len(job["issues"]) for job in report["jobs"])}, ensure_ascii=False))
    return 0 if report["pass"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
