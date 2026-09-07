from __future__ import annotations

import datetime as dt
import json
import math
import os
import unittest
from pathlib import Path

import openpyxl

from ais_monthly_app import (
    AppConfig,
    DayInput,
    DIRECTION_ORDER,
    LEGACY_DIRECTION_ABBREVIATIONS,
    OPEN_SEA_INDEXES,
    SourceFragment,
    build_processing_job,
    parse_source_filename,
    process_day_input,
    scan_source_files,
    select_cluster_from_sorted_rows,
)


def _read_legacy_grid(path: Path) -> dict[tuple[dt.date, int], float | None]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows)
        abbreviation_to_index = {
            abbreviation: index
            for index, abbreviation in enumerate(LEGACY_DIRECTION_ABBREVIATIONS)
        }
        grid: dict[tuple[dt.date, int], float | None] = {}
        for row in rows:
            raw_date = row[0] if row else None
            if isinstance(raw_date, dt.datetime):
                day = raw_date.date()
            elif isinstance(raw_date, dt.date):
                day = raw_date
            else:
                continue
            for column, header in enumerate(headers[1:], start=1):
                direction_index = abbreviation_to_index.get(str(header))
                if direction_index is None:
                    continue
                value = row[column] if column < len(row) else None
                grid[(day, direction_index)] = None if value is None else float(value)
        return grid
    finally:
        workbook.close()


class SelectorBaselineTests(unittest.TestCase):
    def test_full_search_result_does_not_depend_on_display_limit(self) -> None:
        rows = [
            (100.0, 2, 0.0),
            (80.0, 3, 0.0),
            (60.0, 4, 0.0),
            (40.0, 5, 0.0),
            (20.0, 6, 0.0),
            (19.0, 7, 0.0),
            (18.0, 8, 0.0),
        ]
        selected = []
        for display_limit in (3, 5, 50):
            config = AppConfig(
                input_dir=Path("."),
                output_path=Path("result.xlsx"),
                port="KLNG",
                year=2026,
                month=6,
                top_candidates=display_limit,
            )
            result = select_cluster_from_sorted_rows(
                DIRECTION_ORDER[0], rows, config, over_cap_count=0, over_cap_max=None
            )
            selected.append((result.selected, result.selected_rank, result.cluster_count))
        self.assertEqual(selected, [(20.0, 5, 3)] * 3)


class FragmentCatalogTests(unittest.TestCase):
    @staticmethod
    def _write_fragment(path: Path, distances: list[float]) -> None:
        workbook = openpyxl.Workbook(write_only=True)
        sheet = workbook.create_sheet("AIS")
        sheet.append(["msg_type", "LONGITUDE_DESC", "bearing", "distance in nautical miles"])
        for distance in distances:
            sheet.append([1, "East", 0.0, distance])
        workbook.save(path)

    def test_suffix_is_opaque_traceability_not_source_identity(self) -> None:
        first = parse_source_filename("D&TMOK KLNG_20260601_11.xlsx")
        second = parse_source_filename("D&TMOK KLNG_20260601_part-anything.xlsx")
        self.assertIsNotNone(first)
        self.assertIsNotNone(second)
        self.assertEqual((first.port, first.day), (second.port, second.day))
        self.assertEqual(first.suffix, "11")
        self.assertEqual(second.suffix, "PART-ANYTHING")

    def test_same_day_fragments_are_all_kept_and_union_selected(self) -> None:
        import tempfile

        with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as temporary:
            folder = Path(temporary)
            first_path = folder / "D&TMOK KLNG_20260601_11.xlsx"
            second_path = folder / "D&TMOK KLNG_20260601_23.xlsx"
            self._write_fragment(first_path, [100.0, 99.0])
            self._write_fragment(second_path, [98.0])

            catalog, warnings = scan_source_files(folder)
            day_input = catalog["KLNG"][dt.date(2026, 6, 1)]
            self.assertEqual(
                [fragment.path.name for fragment in day_input.fragments],
                [first_path.name, second_path.name],
            )
            self.assertTrue(any("合併 2 個來源分片" in warning for warning in warnings["KLNG"]))

            config = AppConfig(
                input_dir=folder,
                output_path=folder / "result.xlsx",
                port="KLNG",
                year=2026,
                month=6,
                top_candidates=3,
            )
            result = process_day_input(day_input, config)
            self.assertEqual(result.source_fragments, (first_path, second_path))
            self.assertIsNone(result.source_file)
            self.assertEqual(result.rows_scanned, 3)
            self.assertEqual(result.directions["北"].selected, 100.0)

            job = build_processing_job(config)
            self.assertEqual(len(job.days), 1)
            self.assertEqual(len(job.days[0].fragments), 2)
            self.assertEqual(len(job.files), 2)


class ExternalJuneContractTests(unittest.TestCase):
    def test_period_candidate_contract_gate_artifact(self) -> None:
        raw_path = os.environ.get("AIS_JUNE_CONTRACT_GATE_JSON")
        if not raw_path:
            self.skipTest("AIS_JUNE_CONTRACT_GATE_JSON is not configured")
        payload = json.loads(Path(raw_path).read_text(encoding="utf-8"))
        result = payload["results"]
        self.assertEqual(result["nonempty_period_finals"], 1260)
        self.assertEqual(result["matched"], 1260)
        self.assertEqual(result["unmatched"], 0)
        self.assertEqual(result["review_scope_blank_finals"], 0)
        self.assertEqual(result["blank_with_retained_legacy_detail"], 0)
        self.assertEqual(result["structural_land_blank_cells"], 660)
        self.assertEqual(result["structural_land_nonblank_cells"], 0)

    def test_five_file_golden_grid_relationships(self) -> None:
        raw_directory = os.environ.get("AIS_JUNE_GOLDEN_DIR")
        if not raw_directory:
            self.skipTest("AIS_JUNE_GOLDEN_DIR is not configured")
        directory = Path(raw_directory)
        period_a = _read_legacy_grid(directory / "KLNG 6月 32方位數值_2 總表.xlsx")
        period_b = _read_legacy_grid(directory / "KLNG 6月 32方位數值 總表.xlsx")
        integrated = _read_legacy_grid(directory / "KLNG_6月_32方位_每日較大值整合總表.xlsx")

        self.assertEqual(len(period_a), 30 * 32)
        self.assertEqual(len(period_b), 30 * 32)
        self.assertEqual(len(integrated), 30 * 32)
        for key in sorted(integrated):
            left = period_a[key]
            right = period_b[key]
            expected = max(value for value in (left, right) if value is not None) if any(
                value is not None for value in (left, right)
            ) else None
            actual = integrated[key]
            if expected is None:
                self.assertIsNone(actual, key)
            else:
                self.assertTrue(
                    actual is not None
                    and math.isclose(actual, expected, rel_tol=0.0, abs_tol=1e-9),
                    (key, left, right, actual),
                )
        for day in range(1, 31):
            current = dt.date(2026, 6, day)
            for index in range(32):
                if index in OPEN_SEA_INDEXES:
                    self.assertIsNotNone(period_a[(current, index)])
                    self.assertIsNotNone(period_b[(current, index)])
                else:
                    self.assertIsNone(period_a[(current, index)])
                    self.assertIsNone(period_b[(current, index)])


if __name__ == "__main__":
    unittest.main()
