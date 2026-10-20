from __future__ import annotations

import argparse
import json
import math
import zipfile
from datetime import date, datetime
from pathlib import Path

import openpyxl


REFERENCE_CASES = (
    (
        "HWLN",
        1,
        Path(r"B:\AIS\AIS_32方位月報工具\output\AIS月報\HWLN_202601_修復驗證_20260915\正式交付"),
        "v1.5.3 HWLN one-source-day repaired acceptance set",
    ),
    (
        "KLNG",
        4,
        Path(r"B:\AIS\目前可上傳\D&TMOK KLNG_2026-04 已整理資料"),
        "prior accepted automated set",
    ),
    ("KLNG", 5, Path(r"B:\AIS\AIS檔案\五月最後需求檔案"), "partial legacy final-request set"),
    ("KLNG", 6, Path(r"B:\AIS\AIS檔案\六月最後需求檔案"), "manual/senior final-request reference"),
    (
        "KLNG",
        7,
        Path(r"B:\AIS\目前可上傳\D&TMOK KLNG_2026-07 已整理資料"),
        "prior accepted automated set",
    ),
)


def read_grid(path: Path, year: int, month: int) -> list[list[float | None]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["總表"]
        rows: list[list[float | None]] = []
        row = 2
        while True:
            raw_day = sheet.cell(row, 1).value
            if isinstance(raw_day, datetime):
                raw_day = raw_day.date()
            if not isinstance(raw_day, date) or raw_day.year != year or raw_day.month != month:
                break
            rows.append(
                [
                    None if sheet.cell(row, column).value in (None, "")
                    else float(sheet.cell(row, column).value)
                    for column in range(2, 34)
                ]
            )
            row += 1
        return rows
    finally:
        workbook.close()


def compare(left: list[list[float | None]], right: list[list[float | None]]) -> dict[str, object]:
    result: dict[str, object] = {
        "reference_days": len(left),
        "generated_days": len(right),
        "cells": 0,
        "equal": 0,
        "different": 0,
        "reference_blank_generated_value": 0,
        "reference_value_generated_blank": 0,
        "both_values": 0,
        "both_values_equal": 0,
        "both_values_different": 0,
        "examples": [],
    }
    for day_index in range(max(len(left), len(right))):
        left_row = left[day_index] if day_index < len(left) else [None] * 32
        right_row = right[day_index] if day_index < len(right) else [None] * 32
        for direction_index, (expected, actual) in enumerate(zip(left_row, right_row)):
            result["cells"] += 1
            same = expected is actual or (
                expected is not None
                and actual is not None
                and math.isclose(expected, actual, rel_tol=0.0, abs_tol=1e-9)
            )
            if same:
                result["equal"] += 1
                if expected is not None and actual is not None:
                    result["both_values"] += 1
                    result["both_values_equal"] += 1
                continue
            result["different"] += 1
            if expected is not None and actual is not None:
                result["both_values"] += 1
                result["both_values_different"] += 1
            if expected is None and actual is not None:
                result["reference_blank_generated_value"] += 1
            if expected is not None and actual is None:
                result["reference_value_generated_blank"] += 1
            if len(result["examples"]) < 12:
                result["examples"].append(
                    {
                        "day": day_index + 1,
                        "direction_index": direction_index,
                        "reference": expected,
                        "generated": actual,
                    }
                )
    result["match_rate"] = result["equal"] / result["cells"] if result["cells"] else 0.0
    return result


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("generated_root", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    arguments = parser.parse_args()
    report = {"generated_root": str(arguments.generated_root), "comparisons": []}
    for port, month, reference_folder, reference_role in REFERENCE_CASES:
        generated_folder = (
            arguments.generated_root
            / f"{port}_2026{month:02d}"
            / "正式交付"
            / f"{port}_2026{month:02d}"
        )
        for reference in sorted(reference_folder.glob("*.xlsx")):
            generated = generated_folder / reference.name
            if not generated.is_file():
                continue
            try:
                comparison = compare(
                    read_grid(reference, 2026, month),
                    read_grid(generated, 2026, month),
                )
                report["comparisons"].append(
                    {
                        "port": port,
                        "month": month,
                        "reference_role": reference_role,
                        "reference": str(reference),
                        "generated": str(generated),
                        **comparison,
                    }
                )
            except (KeyError, ValueError, OSError, zipfile.BadZipFile) as error:
                report["comparisons"].append(
                    {
                        "port": port,
                        "month": month,
                        "reference_role": reference_role,
                        "reference": str(reference),
                        "generated": str(generated),
                        "error": f"{type(error).__name__}: {error}",
                    }
                )
    arguments.output.write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps([
        {
            "port": item["port"],
            "month": item["month"],
            "file": Path(item["generated"]).name,
            "cells": item.get("cells"),
            "different": item.get("different"),
            "match_rate": item.get("match_rate"),
        }
        for item in report["comparisons"]
    ], ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
