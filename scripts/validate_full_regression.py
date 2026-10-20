from __future__ import annotations

import argparse
import calendar
import hashlib
import json
import math
import posixpath
import re
import struct
import zipfile
import xml.etree.ElementTree as ET
from datetime import date, datetime
from pathlib import Path
from typing import Iterable

import openpyxl

from ais_monthly_app import (
    AppConfig,
    derive_delivery_paths,
)
from scripts.run_full_regression import discover_jobs, job_token, serializable_job
from scripts.verify_source_copy_manifest import atomic_json
from scripts.independent_source_audit import guard_output
from scripts.calendar_source_plan import source_calendar_plan


# Read-only transcription from all five accepted June reference chart1.xml
# parts (2026-10-08). Do not import the producer's palette or abbreviations:
# a shared producer regression must not silently redefine the visual oracle.
REFERENCE_DIRECTION_ABBREVIATIONS = (
    "N", "NbE", "NNE", "NEbN", "NE", "NEbE", "ENE", "EbN", "E", "EbS", "ESE",
    "SEbE", "SE", "SEbS", "SSE", "SbE", "S", "SbW", "SSW", "SWbS", "SW", "SWbW",
    "WSW", "WbS", "W", "WbN", "WNW", "NWbW", "NW", "NWbN", "NNW", "NbW",
)
AUDIT_DIRECTION_NAMES = (
    "北", "北微東", "東北偏北", "東北微北", "東北", "東北微東", "東北偏東", "東微北",
    "東", "東微南", "東南偏東", "東南微東", "東南", "東南微南", "東南偏南", "南微東",
    "南", "南微西", "西南偏南", "西南微南", "西南", "西南微西", "西南偏西", "西微南",
    "西", "西微北", "西北偏西", "西北微西", "西北", "西北微北", "西北偏北", "北微西",
)
DETAIL_ROW = struct.Struct("<ddQIIBq16s")
DETAIL_COUNTS = struct.Struct("<32Q")
REFERENCE_RADAR_PALETTE = (
    "5B9BD5", "ED7D31", "A5A5A5", "FFC000", "4472C4", "70AD47",
    "255E91", "9E480E", "636363", "997300", "264478", "43682B",
    "7CAFDD", "F1975A", "B7B7B7", "FFCD33", "698ED0", "8CC168",
    "327DC2", "D26012", "848484", "CC9A00", "335AA1", "5A8A39",
    "9DC3E6", "F4B183", "C9C9C9", "FFD966", "8FAADC", "A9D18E",
)


C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
XDR_NS = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
NS = {"c": C_NS, "a": A_NS, "xdr": XDR_NS}
S_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS.update(s=S_NS, r=R_NS)

# Independently transcribed from the five copied, accepted June reference ZIPs
# (reference_role_geometry_audit_20261007.json), not from producer layouts.
# Tuples use OOXML (col, colOff, row, rowOff); B full/summary share anchors,
# but their visible column widths differ and thus their physical chart widths.
REFERENCE_DELIVERY_LAYOUTS = {
    "full_a": {"columns": (), "default_col": None,
        "anchors": (((12, 9524, 1, 19050), (22, 609599, 30, 190500)),
                    ((0, 9525, 43, 0), (14, 590550, 56, 142875))),
        "legend": (0.76830972615675164, 0.10649836959552665, 0.22035882908404159, 0.83653926912995602),
        "plot": (0.064771103328797777, 0.17572538384888339, 0.65954183489103524, 0.74501858184039171)},
    "summary_a": {"columns": (), "default_col": None,
        "anchors": (((12, 9524, 1, 19050), (22, 609599, 30, 190500)),
                    ((0, 9525, 43, 0), (14, 590550, 56, 142875))),
        "legend": (0.76830972615675164, 0.10649836959552665, 0.22035882908404159, 0.83653926912995602),
        "plot": (0.064771103328797777, 0.17572538384888339, 0.65954183489103524, 0.74501858184039171)},
    "full_b": {"columns": ((1, 1, 14.5703125), (2, 6, 12.7109375), (7, 7, 11.5703125),
                            (8, 12, 12.7109375), (24, 33, 12.7109375)), "default_col": 9.0,
        "anchors": (((11, 838200, 1, 33336), (23, 0, 31, 9525)),
                    ((0, 0, 43, 14287), (15, 9524, 56, 157162))),
        "legend": (0.8682679434523134, 0.11656551957448454, 0.12020467902607275, 0.82470676043372704), "plot": None},
    "summary_b": {"columns": ((1, 1, 14.5703125),), "default_col": 9.0,
        "anchors": (((11, 838200, 1, 33336), (23, 0, 31, 9525)),
                    ((0, 0, 43, 14287), (15, 9524, 56, 157162))),
        "legend": (0.8682679434523134, 0.11656551957448454, 0.12020467902607275, 0.82470676043372704), "plot": None},
    "integrated": {"columns": ((1, 1, 14.5703125),), "default_col": 9.0,
        "anchors": (((12, 0, 1, 9524), (23, 9525, 31, 19049)),
                    ((0, 0, 43, 9525), (14, 590549, 56, 152400))),
        "legend": (0.87979532097392721, 0.12502090170582877, 0.10867730150445892, 0.82013252305426954), "plot": None},
}


def delivery_artifact_role(path: Path) -> str:
    if path.stem.endswith("每日較大值整合總表"):
        return "integrated"
    is_a = re.search(r"_2(?: 總表)?$", path.stem) is not None
    return ("summary_" if path.stem.endswith(" 總表") else "full_") + ("a" if is_a else "b")


def validate_reference_layout(archive: zipfile.ZipFile, path: Path, day_count: int,
                              drawing: ET.Element, radar: ET.Element | None,
                              combo: ET.Element | None,
                              drawing_part: str = "xl/drawings/drawing1.xml") -> list[str]:
    issues: list[str] = []
    role = delivery_artifact_role(path)
    layout = REFERENCE_DELIVERY_LAYOUTS[role]
    anchors = drawing.findall("xdr:twoCellAnchor", NS)
    if len(anchors) != 2 or len(list(drawing)) != 2:
        return [f"{role}: drawing must contain exactly two twoCellAnchor charts"]
    if not 1 <= day_count <= 31:
        issues.append("day count outside 1..31")
    for index, (anchor, vertices) in enumerate(zip(anchors, layout["anchors"])):
        if anchor.get("editAs") not in (None, "twoCell"):
            issues.append(f"{role}: chart {index + 1} anchor behavior")
        for edge, expected in zip(("from", "to"), vertices):
            vertex = anchor.find(f"xdr:{edge}", NS)
            try:
                actual = tuple(int(vertex.find(f"xdr:{tag}", NS).text)
                               for tag in ("col", "colOff", "row", "rowOff"))
                shifted = tuple(value + (day_count - 30 if index == 1 and pos == 2 else 0)
                                for pos, value in enumerate(expected))
                if actual != shifted:
                    issues.append(f"{role}: chart {index + 1} {edge} anchor {actual} != {shifted}")
            except (AttributeError, TypeError, ValueError):
                issues.append(f"{role}: chart {index + 1} malformed {edge} anchor")
    drawing_rels = posixpath.join(posixpath.dirname(drawing_part), "_rels",
                                 posixpath.basename(drawing_part) + ".rels")
    if drawing_rels not in archive.namelist():
        issues.append(f"{role}: missing chart drawing relationships")
    else:
        targets = {item.get("Id"): item.get("Target") for item in ET.fromstring(archive.read(drawing_rels))}
        for index, anchor in enumerate(anchors):
            node = anchor.find(".//c:chart", NS)
            target = None if node is None else targets.get(node.get(f"{{{R_NS}}}id"))
            part = "" if target is None else (target.lstrip("/") if target.startswith("/")
                else posixpath.normpath(posixpath.join(posixpath.dirname(drawing_part), target)))
            if part not in archive.namelist():
                issues.append(f"{role}: chart {index + 1} missing drawing source")
            else:
                root = ET.fromstring(archive.read(part))
                kind = "radarChart" if index == 0 else "lineChart"
                if root.find(f".//c:{kind}", NS) is None:
                    issues.append(f"{role}: chart {index + 1} anchored to wrong chart role")
    # Read only the sheet owning a drawing relationship, never decompress all
    # retained detail sheets just to find the chart summary.
    summary_parts = []
    for name in archive.namelist():
        if re.fullmatch(r"xl/worksheets/_rels/sheet\d+\.xml\.rels", name):
            rels = ET.fromstring(archive.read(name))
            sheet_part = "xl/worksheets/" + Path(name).name.removesuffix(".rels")
            if any(item.get("Type", "").endswith("/drawing") and
                (item.get("Target", "").lstrip("/") if item.get("Target", "").startswith("/") else
                 posixpath.normpath(posixpath.join("xl/worksheets", item.get("Target", "")))) == drawing_part for item in rels):
                summary_parts.append(sheet_part)
    if len(summary_parts) != 1:
        issues.append(f"{role}: drawing-owner sheet count {len(summary_parts)} != 1")
    else:
        sheet = ET.fromstring(archive.read(summary_parts[0]))
        formatting = sheet.find("s:sheetFormatPr", NS)
        if formatting is None or float(formatting.get("defaultRowHeight", "0")) != 15.75:
            issues.append(f"{role}: default row height must be 15.75")
        actual_default = None if formatting is None else formatting.get("defaultColWidth")
        actual_default = None if actual_default is None else float(actual_default)
        if actual_default != layout["default_col"]:
            issues.append(f"{role}: default column width {actual_default}")
        expected_columns = {column: width for first, last, width in layout["columns"]
                            for column in range(first, last + 1)}
        actual_columns = {}
        for column in sheet.findall("s:cols/s:col", NS):
            for index in range(int(column.get("min")), min(int(column.get("max")), 33) + 1):
                actual_columns[index] = float(column.get("width", "0"))
                if column.get("hidden") == "1":
                    issues.append(f"{role}: visible column {index} hidden")
        if actual_columns != expected_columns:
            issues.append(f"{role}: visible column geometry differs from manual reference")
        for row in sheet.findall("s:sheetData/s:row", NS):
            if row.get("hidden") == "1" or (row.get("ht") is not None and float(row.get("ht")) != 15.75):
                issues.append(f"{role}: row {row.get('r')} overrides reference height/visibility")
    def manual_layout(node: ET.Element | None, expected: tuple | None, label: str) -> None:
        manual = None if node is None else node.find("c:layout/c:manualLayout", NS)
        if expected is None:
            if manual is not None:
                issues.append(f"{role}: {label} must use automatic layout")
            return
        try:
            values = tuple(float(manual.find(f"c:{tag}", NS).get("val")) for tag in ("x", "y", "w", "h"))
            modes = tuple(manual.find(f"c:{tag}", NS).get("val") for tag in ("xMode", "yMode"))
            if modes != ("edge", "edge") or not all(math.isclose(a, b, abs_tol=1e-12, rel_tol=0) for a, b in zip(values, expected)):
                issues.append(f"{role}: {label} manual layout differs from reference")
            if label == "radar plot":
                target = manual.find("c:layoutTarget", NS)
                if target is None or target.get("val") != "inner":
                    issues.append(f"{role}: radar plot layout target must be inner")
        except (AttributeError, TypeError, ValueError):
            issues.append(f"{role}: {label} missing/malformed manual layout")
    manual_layout(None if radar is None else radar.find("c:chart/c:legend", NS), layout["legend"], "radar legend")
    manual_layout(None if radar is None else radar.find("c:chart/c:plotArea", NS), layout["plot"], "radar plot")
    manual_layout(None if combo is None else combo.find("c:chart/c:legend", NS), None, "statistics legend")
    manual_layout(None if combo is None else combo.find("c:chart/c:plotArea", NS), None, "statistics plot")
    theme_colors = {"tx1": "000000", "dk1": "000000"}
    if "xl/theme/theme1.xml" in archive.namelist():
        theme = ET.fromstring(archive.read("xl/theme/theme1.xml"))
        for entry in theme.findall("a:themeElements/a:clrScheme/*", NS):
            if len(entry):
                theme_colors[entry.tag.rsplit("}", 1)[-1]] = entry[0].get("val") if entry[0].tag.endswith("srgbClr") else entry[0].get("lastClr")
        theme_colors["tx1"] = theme_colors.get("dk1", "000000")
        # Independent reference transcription: all chart font owners use
        # minor-theme tokens. A/B use Calibri Latin plus Hant 新細明體;
        # integrated defines the three primary script faces as Calibri.
        minor = theme.find("a:themeElements/a:fontScheme/a:minorFont", NS)
        expected_faces = ("Calibri", "Calibri", "Calibri") if role == "integrated" else ("Calibri", "", "")
        actual_faces = tuple(None if minor is None or minor.find(f"a:{script}", NS) is None
                             else minor.find(f"a:{script}", NS).get("typeface")
                             for script in ("latin", "ea", "cs"))
        if actual_faces != expected_faces:
            issues.append(f"{role}: minor theme script font faces differ from reference")
        if role == "integrated" and (minor is None or
                [font.tag.rsplit("}", 1)[-1] for font in minor] != ["latin", "ea", "cs"]):
            issues.append(f"{role}: integrated minor theme must not retain supplemental script fonts")
        if role != "integrated" and (minor is None or not any(
                font.get("script") == "Hant" and font.get("typeface") == "新細明體"
                for font in minor.findall("a:font", NS))):
            issues.append(f"{role}: Hant minor theme font must be 新細明體")
    else:
        issues.append(f"{role}: missing chart font theme")
    def text_color(fill: ET.Element) -> str | None:
        if not len(fill):
            return None
        color = fill[0]
        kind = color.tag.rsplit("}", 1)[-1]
        value = color.get("val") if kind == "srgbClr" else color.get("lastClr") if kind == "sysClr" else theme_colors.get(color.get("val"))
        try:
            channels = [int(value[offset:offset + 2], 16) / 255 for offset in (0, 2, 4)]
            # Accepted references use black plus lumMod/lumOff (gray595959).
            # Unsupported transforms fail closed instead of guessing a color.
            for transform in color:
                tag = transform.tag.rsplit("}", 1)[-1]
                amount = int(transform.get("val")) / 100000
                if tag not in ("lumMod", "lumOff") or len(set(channels)) != 1:
                    return None
                channels = [channel * amount if tag == "lumMod" else channel + amount for channel in channels]
            return "".join(f"{round(max(0, min(1, channel)) * 255):02X}" for channel in channels)
        except (TypeError, ValueError):
            return None
    for label, chart, legend_position in (("radar", radar, "r"), ("statistics", combo, "b")):
        if chart is None:
            continue
        language = chart.find("c:lang", NS)
        if language is None or language.get("val") != "zh-TW":
            issues.append(f"{role}: {label} chart language must be zh-TW")
        for scope, expected_size in (("title", "1400"), ("legend", "900")):
            node = chart.find(f"c:chart/c:{scope}", NS)
            properties = [] if node is None else node.findall(".//a:defRPr", NS) + node.findall(".//a:rPr", NS)
            sizes = [item.get("sz") for item in properties if item.get("sz") is not None]
            if not sizes or any(size != expected_size for size in sizes):
                issues.append(f"{role}: {label} {scope} font must be {int(expected_size) / 100:g}pt")
            for properties_node in properties:
                if properties_node.get("sz") is None:
                    continue
                # Native reference rich runs may inherit their fonts from the
                # paragraph defRPr; explicit partial/incorrect tokens fail.
                if properties_node.tag.endswith("}rPr") and not any(
                        properties_node.find(f"a:{script}", NS) is not None
                        for script in ("latin", "ea", "cs")):
                    continue
                triplet = tuple(None if properties_node.find(f"a:{script}", NS) is None
                                else properties_node.find(f"a:{script}", NS).get("typeface")
                                for script in ("latin", "ea", "cs"))
                # Integrated reference text renders Calibri. Native isolated
                # tests prove XlsxWriter requires explicit child-owner faces;
                # identical theme tokens alone rendered the wrong glyphs.
                expected_triplet = (("Calibri",) * 3 if role == "integrated"
                                    else ("+mn-lt", "+mn-ea", "+mn-cs"))
                if triplet != expected_triplet:
                    issues.append(f"{role}: {label} {scope} script font tokens differ from reference")
            if scope == "title" and node is not None:
                for run in node.findall("c:tx/c:rich/a:p/a:r", NS):
                    text = run.find("a:t", NS)
                    properties_node = run.find("a:rPr", NS)
                    if text is None or text.text is None or properties_node is None:
                        issues.append(f"{role}: {label} title malformed rich text run")
                    elif (properties_node.get("lang") != ("en-US" if text.text.isascii() else "zh-TW") or
                          (not text.text.isascii() and any(ord(char) < 128 for char in text.text))):
                        issues.append(f"{role}: {label} title mixed-script language differs from reference")
            fills = [] if node is None else node.findall(".//a:solidFill", NS)
            if not fills or any(text_color(fill) != "595959" for fill in fills):
                issues.append(f"{role}: {label} {scope} text color must be 595959")
        position = chart.find("c:chart/c:legend/c:legendPos", NS)
        if position is None or position.get("val") != legend_position:
            issues.append(f"{role}: {label} legend position")
        default = chart.find("c:txPr/a:p/a:pPr/a:defRPr", NS)
        language = chart.find("c:txPr/a:p/a:endParaRPr", NS)
        if default is None or default.attrib or len(default) or language is None or language.get("lang") != "zh-TW":
            issues.append(f"{role}: {label} chart-level text inheritance differs from reference")
    columns = None if combo is None else combo.find(".//c:barChart", NS)
    if columns is None or any(columns.find(f"c:{tag}", NS) is None or
                             columns.find(f"c:{tag}", NS).get("val") != expected
                             for tag, expected in (("grouping", "clustered"), ("gapWidth", "219"), ("overlap", "-27"))):
        issues.append(f"{role}: statistics grouped bar gap/overlap differs from reference")
    if combo is not None:
        for axis in combo.findall(".//c:catAx", NS) + combo.findall(".//c:valAx", NS):
            if any(axis.find(f"c:{mark}", NS) is None or axis.find(f"c:{mark}", NS).get("val") != "none"
                   for mark in ("majorTickMark", "minorTickMark")):
                issues.append(f"{role}: statistics axis tick marks differ from reference")
        line = combo.find(".//c:catAx/c:spPr/a:ln", NS)
        color = None if line is None else line.find("a:solidFill/a:srgbClr", NS)
        if line is None or line.get("w") != "9525" or color is None or color.get("val") != "D9D9D9":
            issues.append(f"{role}: statistics category baseline differs from reference")
        title_position = combo.find("c:chart/c:title/c:layout/c:manualLayout", NS)
        if role.endswith("_a"):
            try:
                if (title_position.find("c:xMode", NS).get("val"), title_position.find("c:yMode", NS).get("val")) != ("edge", "edge") or not all(
                        math.isclose(float(title_position.find(f"c:{tag}", NS).get("val")), expected, rel_tol=0, abs_tol=1e-12)
                        for tag, expected in (("x", 0.3523162112572919), ("y", 0.027777777777777776))):
                    issues.append(f"{role}: statistics title position differs from reference")
            except (AttributeError, TypeError, ValueError):
                issues.append(f"{role}: missing/malformed reference statistics title position")
        elif title_position is not None:
            issues.append(f"{role}: statistics title must use automatic position")
    return issues


def semantic_package_fingerprint(path: Path, *, application_version: str | None = None) -> dict[str, object]:
    """Hash all parts, normalizing only validated creation/self-path provenance.

    This is one repeatability gate, not a replacement for numeric/native QA.
    Unrecognized fields, paths, formulas, cell formats and chart content stay
    significant. Invalid expected provenance raises rather than being ignored.
    """
    changed: dict[str, bytes] = {}
    provenance: list[dict[str, str]] = []
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        if len(names) != len(set(names)):
            raise ValueError(f"duplicate ZIP parts in {path}")
        if application_version is not None:
            # Only an explicitly declared, exact application-version property
            # may differ across a hardening baseline. All other custom fields,
            # data, formula/style/chart parts remain significant.
            if not re.fullmatch(r"\d+\.\d+\.\d+",application_version):
                raise ValueError("Invalid declared application version")
            custom = ET.fromstring(archive.read('docProps/custom.xml'))
            selected = [p for p in custom if p.get('name')=='Application version']
            if (len(selected)!=1 or len(selected[0])!=1 or
                    selected[0][0].tag!='{http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes}lpwstr' or
                    selected[0][0].text!=application_version):
                raise ValueError('Application version differs from explicit comparison declaration')
            provenance.append({'part':'docProps/custom.xml','field':'Application version','value':application_version})
            selected[0][0].text='<DECLARED_APPLICATION_VERSION>'
            changed['docProps/custom.xml']=ET.tostring(custom,encoding='utf-8')
        if "docProps/core.xml" in names:
            core = ET.fromstring(archive.read("docProps/core.xml"))
            dcterms = "http://purl.org/dc/terms/"
            for field in ("created", "modified"):
                nodes = core.findall(f"{{{dcterms}}}{field}")
                if len(nodes) != 1 or not re.fullmatch(
                        r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}Z", nodes[0].text or ""):
                    raise ValueError(f"invalid core {field} provenance in {path}")
                datetime.fromisoformat(nodes[0].text.replace("Z", "+00:00"))
                provenance.append({"part": "docProps/core.xml", "field": field,
                                   "value": nodes[0].text})
                nodes[0].text = "2000-01-01T00:00:00Z"
            # Retain title/subject/creator/description/lastModifiedBy, rather
            # than dropping the whole core.xml as the previous helper did.
            changed["docProps/core.xml"] = ET.tostring(core, encoding="utf-8")
        if "xl/workbook.xml" in names:
            workbook = ET.fromstring(archive.read("xl/workbook.xml"))
            sheets = workbook.findall("s:sheets/s:sheet", NS)
            by_name = {item.get("name"): item for item in sheets}
            modern_names = {"操作說明", "系統資料"}
            if modern_names & by_name.keys():
                if (not modern_names <= by_name.keys() or
                        any(sum(item.get("name") == name for item in sheets) != 1 for name in modern_names)):
                    raise ValueError(f"incomplete/duplicate Modern provenance sheets in {path}")
                rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
                targets = {item.get("Id"): item.get("Target", "") for item in rels}
                strings = ET.fromstring(archive.read("xl/sharedStrings.xml"))
                string_nodes = strings.findall("s:si", NS)
                permitted: set[tuple[str, str]] = set()
                indexes: set[int] = set()
                expected_path = str(path.resolve())
                for name, row, label in (("操作說明", 7, "Modern analysis"),
                                         ("系統資料", 11, "Analysis workbook path")):
                    target = targets[by_name[name].get(f"{{{R_NS}}}id")]
                    part = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
                    sheet = ET.fromstring(archive.read(part))
                    cells = sheet.findall(f"s:sheetData/s:row[@r='{row}']/s:c", NS)
                    selected = {cell.get("r"): cell for cell in cells}
                    if len(selected) != len(cells):
                        raise ValueError(f"duplicate provenance cell in {name} of {path}")
                    def shared_value(coordinate: str) -> tuple[int, str]:
                        cell = selected.get(coordinate)
                        if (cell is None or cell.get("t") != "s" or cell.find("s:f", NS) is not None or
                                cell.find("s:v", NS) is None):
                            raise ValueError(f"invalid provenance cell {name}!{coordinate} in {path}")
                        index = int(cell.find("s:v", NS).text)
                        if not 0 <= index < len(string_nodes):
                            raise ValueError(f"invalid provenance string index in {path}")
                        si = string_nodes[index]
                        if len(si) != 1 or si[0].tag != f"{{{S_NS}}}t":
                            raise ValueError(f"nonplain provenance string in {path}")
                        return index, si[0].text or ""
                    if shared_value(f"A{row}")[1] != label:
                        raise ValueError(f"wrong provenance label {name}!A{row} in {path}")
                    index, value = shared_value(f"B{row}")
                    if value != expected_path:
                        raise ValueError(f"unexpected self workbook path {name}!B{row}: {value!r} != {expected_path!r}")
                    permitted.add((part, f"B{row}"))
                    indexes.add(index)
                    provenance.append({"part": part, "field": f"{name}!B{row}", "value": value})
                # A shared string may be reused elsewhere. Replacing it without
                # checking every consumer could hide an unintended path edit.
                observed: list[tuple[str, str]] = []
                for part in names:
                    if not re.fullmatch(r"xl/worksheets/sheet\d+\.xml", part):
                        continue
                    with archive.open(part) as handle:
                        for _, node in ET.iterparse(handle, events=("end",)):
                            if node.tag == f"{{{S_NS}}}c":
                                value = node.find("s:v", NS)
                                if node.get("t") == "s" and value is not None and int(value.text) in indexes:
                                    observed.append((part, node.get("r")))
                                node.clear()
                            elif node.tag == f"{{{S_NS}}}row":
                                node.clear()
                if len(observed) != len(permitted) or set(observed) != permitted:
                    raise ValueError(f"self workbook path has unexpected consumers in {path}: {observed}")
                for index in indexes:
                    string_nodes[index][0].text = "<VERIFIED_SELF_ANALYSIS_WORKBOOK_PATH>"
                changed["xl/sharedStrings.xml"] = ET.tostring(strings, encoding="utf-8")
                timing_names = {"總表", "處理紀錄"}
                if timing_names & by_name.keys():
                    if not timing_names <= by_name.keys():
                        raise ValueError(f"incomplete Modern timing provenance sheets in {path}")
                    def worksheet(name: str) -> tuple[str, ET.Element]:
                        target = targets[by_name[name].get(f"{{{R_NS}}}id")]
                        part = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
                        return part, ET.fromstring(archive.read(part))
                    def plain_number(cell: ET.Element | None, field: str) -> float:
                        if (cell is None or cell.get("t") not in (None, "n") or
                                len(cell) != 1 or cell[0].tag != f"{{{S_NS}}}v"):
                            raise ValueError(f"invalid plain numeric provenance {field} in {path}")
                        value = float(cell[0].text)
                        if not math.isfinite(value):
                            raise ValueError(f"nonfinite provenance {field} in {path}")
                        return value
                    def label_text(cell: ET.Element | None) -> str | None:
                        if cell is None or cell.get("t") != "s" or cell.find("s:f", NS) is not None:
                            return None
                        value = cell.find("s:v", NS)
                        return None if value is None else "".join(string_nodes[int(value.text)].itertext())
                    _, total = worksheet("總表")
                    if label_text(total.find("s:sheetData/s:row/s:c[@r='A4']", NS)) != "日期":
                        raise ValueError(f"wrong Modern summary date header in {path}")
                    date_serials: list[float] = []
                    for row in total.findall("s:sheetData/s:row", NS):
                        index = int(row.get("r"))
                        cell = row.find(f"s:c[@r='A{index}']", NS)
                        if index >= 5 and cell is not None and cell.get("t") in (None, "n"):
                            if index != len(date_serials) + 5:
                                raise ValueError(f"nonconsecutive Modern logical day rows in {path}")
                            serial = plain_number(cell, f"總表!A{index}")
                            if serial != int(serial) or serial < 1:
                                raise ValueError(f"invalid logical day date serial in {path}")
                            date_serials.append(serial)
                    if not date_serials or date_serials != sorted(set(date_serials)):
                        raise ValueError(f"invalid Modern logical day sequence in {path}")
                    log_part, log = worksheet("處理紀錄")
                    headers = ("日期", "來源檔", "掃描列數", "合格列數", "錯誤列數", "訊息類型略過", "非 East 略過", "秒數")
                    for column, expected in zip("ABCDEFGH", headers):
                        if label_text(log.find(f"s:sheetData/s:row/s:c[@r='{column}5']", NS)) != expected:
                            raise ValueError(f"wrong timing log header {column}5 in {path}")
                    timed_rows: list[int] = []
                    for row in log.findall("s:sheetData/s:row", NS):
                        index = int(row.get("r"))
                        if index < 6:
                            continue
                        if index != len(timed_rows) + 6 or index > len(date_serials) + 5:
                            raise ValueError(f"unexpected timing log row {index} in {path}")
                        serial = plain_number(row.find(f"s:c[@r='A{index}']", NS), f"處理紀錄!A{index}")
                        if serial != date_serials[index - 6]:
                            raise ValueError(f"timing logical day mapping mismatch {index} in {path}")
                        timings = row.findall(f"s:c[@r='H{index}']", NS)
                        if len(timings) != 1 or plain_number(timings[0], f"處理紀錄!H{index}") < 0:
                            raise ValueError(f"invalid elapsed seconds H{index} in {path}")
                        provenance.append({"part": log_part, "field": f"處理紀錄!H{index}",
                                           "value": timings[0][0].text})
                        timings[0][0].text = "0"
                        timed_rows.append(index)
                    if len(timed_rows) != len(date_serials):
                        raise ValueError(f"missing timing log rows in {path}")
                    # Only elapsed values change. Dates, row counts, source names,
                    # cell formats, formulas and all other log cells remain hashed.
                    changed[log_part] = ET.tostring(log, encoding="utf-8")
        digest = hashlib.sha256()
        for name in sorted(names):
            digest.update(name.encode("utf-8") + b"\0")
            payload = changed.get(name)
            digest.update(str(len(payload) if payload is not None else archive.getinfo(name).file_size).encode("ascii") + b"\0")
            if payload is not None:
                digest.update(payload)
            else:
                with archive.open(name) as handle:
                    while chunk := handle.read(1024 * 1024):
                        digest.update(chunk)
    return {"sha256": digest.hexdigest(), "provenance_normalizations": provenance}


def semantic_package_hash(path: Path) -> str:
    return str(semantic_package_fingerprint(path)["sha256"])


def as_number(value: object) -> float | None:
    if value is None or value == "":
        return None
    return float(value)


def close_enough(left: float | None, right: float | None) -> bool:
    if left is None or right is None:
        return left is right
    return math.isclose(left, right, rel_tol=0.0, abs_tol=1e-9)


def read_summary(path: Path, day_count: int) -> tuple[list[date], list[list[float | None]]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["總表"]
        headers = tuple(sheet.cell(1, column).value for column in range(2, 34))
        if headers != REFERENCE_DIRECTION_ABBREVIATIONS:
            raise AssertionError(f"direction headers mismatch: {path}")
        days: list[date] = []
        matrix: list[list[float | None]] = []
        for row in range(2, day_count + 2):
            raw_day = sheet.cell(row, 1).value
            if isinstance(raw_day, datetime):
                raw_day = raw_day.date()
            if not isinstance(raw_day, date):
                raise AssertionError(f"invalid day at {path} row {row}: {raw_day!r}")
            days.append(raw_day)
            matrix.append([as_number(sheet.cell(row, column).value) for column in range(2, 34)])
        return days, matrix
    finally:
        workbook.close()


def independent_detail_bin(bearing: object) -> int | None:
    if isinstance(bearing, bool) or not isinstance(bearing, (int, float)) or not math.isfinite(bearing):
        return None
    numerator, denominator = float(bearing).as_integer_ratio()
    return ((numerator % (360 * denominator)) * 4) // (45 * denominator)


def expected_retained_detail(spool: Path, period_code: int, included: set[int],
                             finals: list[float | None]) -> Iterable[tuple[float, float, str]]:
    """Independent full row projection; no production reader/filter imports.

    Values are rounded only to the actual installed XlsxWriter 16G XML numeric
    serialization, not to the displayed three-decimal format or a broad epsilon.
    Selector/raw provenance gates independently establish finals/spool inputs.
    """
    with spool.open("rb") as handle:
        if handle.read(8) != b"AISNRM1\0":
            raise ValueError(f"invalid detail spool magic: {spool}")
        header = handle.read(DETAIL_COUNTS.size)
        if len(header) != DETAIL_COUNTS.size:
            raise ValueError("truncated detail spool header")
        counts = DETAIL_COUNTS.unpack(header)
        if spool.stat().st_size != 8 + DETAIL_COUNTS.size + sum(counts) * DETAIL_ROW.size:
            raise ValueError("detail spool length/count mismatch")
        for index, count in enumerate(counts):
            for _ in range(count):
                raw = handle.read(DETAIL_ROW.size)
                if len(raw) != DETAIL_ROW.size:
                    raise ValueError("truncated detail spool record")
                distance, bearing, _key, second, msg, stored_period, _mmsi, _channel = DETAIL_ROW.unpack(raw)
                if (not math.isfinite(distance) or distance < 0 or not math.isfinite(bearing)
                        or not 0 <= bearing < 360 or independent_detail_bin(bearing) != index
                        or second >= 86400 or stored_period != int(second >= 43200)):
                    raise ValueError("invalid detail normalized record")
                if msg not in {1, 3, 4, 18, 19} or stored_period != period_code:
                    continue
                if index in included:
                    if finals[index] is None or distance > finals[index] + 1e-12:
                        continue
                elif distance > 500:
                    continue
                yield float(format(distance, ".16G")), float(format(bearing, ".16G")), AUDIT_DIRECTION_NAMES[index]
        if handle.read(1):
            raise ValueError("unexpected detail spool tail")


def inspect_detail_mapping(
    path: Path, expected_days: list[date], included: set[int], spool_dir: Path | None = None,
    source_days: set[date] | None = None,
) -> dict[str, int]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = 0
    mismatches = 0
    summary_mismatches = 0
    retained_mismatches = 0
    expected_rows = 0
    role = delivery_artifact_role(path)
    if spool_dir is not None and role not in {"full_a", "full_b"}:
        raise ValueError("Retained detail audit requires a full A/B artifact")
    try:
        summary = workbook["總表"]
        summary_values = {
            day: [as_number(summary.cell(index + 2, column).value) for column in range(2, 34)]
            for index, day in enumerate(expected_days)
        }
        for day in expected_days:
            sheet = workbook[f"{day.month}月{day.day}日"]
            if tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1, max_col=3))) != ("DISTANCE", "DEGREE", "方位"):
                raise ValueError("wrong retained detail column headers")
            maxima: dict[int, float] = {}
            if spool_dir is None:
                expected_iterator = None
            elif source_days is not None and day not in source_days:
                expected_iterator = iter(())
            else:
                expected_iterator = iter(expected_retained_detail(
                    spool_dir / (day.strftime("%Y%m%d") + ".bin"), 0 if role == "full_a" else 1,
                    included, summary_values[day]))
            for distance, bearing, direction in sheet.iter_rows(
                min_row=2, min_col=1, max_col=3, values_only=True
            ):
                if distance is None and bearing is None and direction is None:
                    continue
                rows += 1
                if expected_iterator is not None:
                    expected = next(expected_iterator, None)
                    if expected is not None:
                        expected_rows += 1
                    if (expected is None or isinstance(distance, bool) or isinstance(bearing, bool)
                            or (distance, bearing, direction) != expected):
                        retained_mismatches += 1
                parsed = independent_detail_bin(bearing)
                if parsed is None or AUDIT_DIRECTION_NAMES[parsed] != direction:
                    mismatches += 1
                    continue
                if isinstance(distance, bool) or not isinstance(distance, (int, float)) or not math.isfinite(distance) or distance < 0:
                    mismatches += 1
                    continue
                maxima[parsed] = max(maxima.get(parsed, float("-inf")), float(distance))
            if expected_iterator is not None:
                for _record in expected_iterator:
                    expected_rows += 1
                    retained_mismatches += 1
            for index, expected in enumerate(summary_values[day]):
                if index not in included:
                    continue
                actual = maxima.get(index)
                if not close_enough(actual, expected):
                    summary_mismatches += 1
        return {
            "detail_rows": rows,
            "direction_mapping_mismatches": mismatches,
            "detail_max_summary_mismatches": summary_mismatches,
            "retained_expected_rows": expected_rows,
            "retained_row_mismatches": retained_mismatches,
        }
    finally:
        workbook.close()


def _color(series: ET.Element, *, line: bool = False) -> str | None:
    path = ".//c:spPr/a:ln/a:solidFill/*" if line else ".//c:spPr/a:solidFill/*"
    node = series.find(path, NS)
    if node is None:
        return None
    return node.get("val")


def validate_delivery_charts(path: Path, day_count: int) -> dict[str, object]:
    issues: list[str] = []
    formula_backed_radar = delivery_artifact_role(path) in {"full_a", "full_b", "integrated"}
    radar_start, radar_end = ("AO", "BT") if formula_backed_radar else ("B", "AG")
    with zipfile.ZipFile(path) as archive:
        chart_parts = sorted(
            name for name in archive.namelist()
            if re.fullmatch(r"xl/charts/chart\d+\.xml", name)
        )
        if len(chart_parts) != 2:
            issues.append(f"chart count {len(chart_parts)} != 2")
        roots = [ET.fromstring(archive.read(name)) for name in chart_parts]
        radar = next((root for root in roots if root.find(".//c:radarChart", NS) is not None), None)
        combo = next((root for root in roots if root.find(".//c:lineChart", NS) is not None), None)
        if radar is None:
            issues.append("missing radar")
        else:
            style = radar.find(".//c:radarStyle", NS)
            if style is None or style.get("val") != "filled":
                issues.append("radar is not filled")
            series = radar.findall(".//c:radarChart/c:ser", NS)
            for position, item in enumerate(series):
                category = item.find("c:cat/*/c:f", NS)
                value = item.find("c:val/c:numRef/c:f", NS)
                if category is None or "$B$1:$AG$1" not in (category.text or ""):
                    issues.append(f"radar series {position + 1} category range")
                if value is None or not re.search(rf"\${radar_start}\$\d+:\${radar_end}\$\d+$", value.text or ""):
                    issues.append(f"radar series {position + 1} value range")
                fill = item.find(".//c:spPr/a:solidFill/a:srgbClr", NS)
                if fill is None:
                    issues.append(f"radar series {position + 1} missing RGB fill")
                    continue
                # A filled radar may omit a wholly empty logical day. Its
                # worksheet row is the day-plan offset; series position is not.
                row_match = None if value is None else re.search(
                    rf"\${radar_start}\$(\d+):\${radar_end}\$(\d+)$", value.text or "")
                if (row_match is None or row_match[1] != row_match[2] or
                        not 2 <= int(row_match[1]) <= day_count + 1):
                    issues.append(f"radar series {position + 1} invalid logical day row")
                elif fill.get("val") != REFERENCE_RADAR_PALETTE[(int(row_match[1]) - 2) % 30]:
                    issues.append(f"radar series {position + 1} color")
                alpha = fill.find("a:alpha", NS)
                if alpha is None or alpha.get("val") != "40000":
                    issues.append(f"radar series {position + 1} transparency")
            if formula_backed_radar:
                visibility = radar.find("c:chart/c:plotVisOnly", NS)
                if visibility is not None and visibility.get("val") != "0":
                    issues.append("formula-backed radar hidden helpers are not plotted")
        if combo is None:
            issues.append("missing combined statistics chart")
        else:
            bars = combo.findall(".//c:barChart/c:ser", NS)
            lines = combo.findall(".//c:lineChart/c:ser", NS)
            if len(bars) != 2 or len(lines) != 1:
                issues.append(f"combo series {len(bars)}+{len(lines)}")
            expected_colors = ("00B0F0", "ED7D31")
            expected_titles = ("MAX", "MIN")
            for index, item in enumerate(bars[:2]):
                fill = item.find(".//c:spPr/a:solidFill/a:srgbClr", NS)
                if fill is None or fill.get("val") != expected_colors[index]:
                    issues.append(f"bar series {index + 1} color")
                title = item.find("c:tx/c:v", NS)
                if title is None or title.text != expected_titles[index]:
                    issues.append(f"bar series {index + 1} title")
            if lines:
                line_color = lines[0].find(".//c:spPr/a:ln/a:solidFill/a:srgbClr", NS)
                width = lines[0].find(".//c:spPr/a:ln", NS)
                marker = lines[0].find("c:marker/c:symbol", NS)
                if line_color is None or line_color.get("val") != "FF0000":
                    issues.append("average line color")
                if width is None or width.get("w") != "28575":
                    issues.append("average line width")
                if marker is None or marker.get("val") != "none":
                    issues.append("average marker")
                title = lines[0].find("c:tx/c:v", NS)
                if title is None or title.text != "平均值":
                    issues.append("average title")
            expected_indexes = (
                tuple(range(22, 32)) + tuple(range(0, 11))
                if path.name.upper().startswith("KLNG")
                else tuple(range(32))
            )
            expected_categories = [
                REFERENCE_DIRECTION_ABBREVIATIONS[index] for index in expected_indexes
            ]
            for position, item in enumerate([*bars, *lines]):
                category_values = [
                    point.find("c:v", NS).text
                    for point in item.findall("c:cat/c:strRef/c:strCache/c:pt", NS)
                ]
                if category_values != expected_categories:
                    issues.append(f"combo series {position + 1} category order")
        drawing_parts = sorted(
            name for name in archive.namelist()
            if re.fullmatch(r"xl/drawings/drawing\d+\.xml", name)
        )
        if not drawing_parts:
            issues.append("missing drawing")
        else:
            drawing = ET.fromstring(archive.read(drawing_parts[0]))
            if len(drawing_parts) != 1:
                issues.append(f"drawing part count {len(drawing_parts)} != 1")
            issues.extend(validate_reference_layout(archive, path, day_count, drawing, radar, combo, drawing_parts[0]))
    return {"issues": issues}


def expected_delivery(job: dict[str, object]) -> tuple[Path, list[Path]]:
    delivery_root = Path(str(job["delivery"]))
    config = AppConfig(
        input_dir=Path(str(job["folder"])),
        output_path=Path(str(job["modern"])),
        port=str(job["port"]),
        year=int(job["year"]),
        month=int(job["month"]),
        delivery_dir=delivery_root,
        overwrite=True,
    )
    paths = derive_delivery_paths(config, root=delivery_root)
    return Path(str(job["modern"])), list(paths.all_files())


def audit_run_scope(run_report: dict, source_root: Path | None = None) -> dict:
    """A complete source plan is required before any output can be called full PASS."""
    issues: list[str] = []
    records = run_report.get("jobs")
    if not isinstance(records, list):
        issues.append("run jobs must be a list")
        records = []
    if not records:
        issues.append("run contains no completed jobs")
    declared_plan = run_report.get("planned_jobs")
    if declared_plan is not None and not isinstance(declared_plan, list):
        issues.append("planned_jobs must be a list")
        declared_plan = []
    discovered = None
    if source_root is not None:
        try:
            discovered = [serializable_job(job) for job in discover_jobs(source_root)]
        except Exception as error:
            issues.append(f"source scope discovery failed: {type(error).__name__}: {error}")
    plan = discovered if discovered is not None else declared_plan
    if not plan:
        issues.append("no nonempty source plan; supply --source-root for a legacy report")
        plan = []
    if run_report.get("scope") == "SUBSET":
        issues.append("run scope is SUBSET, not a complete source regression")
    status = run_report.get("status")
    if status is not None and status != "COMPLETE":
        issues.append(f"run status {status} is not COMPLETE")
    if declared_plan is not None:
        if status is None:
            issues.append("planned run is missing completion status")
        if run_report.get("expected_job_count") != len(declared_plan):
            issues.append("expected_job_count does not match planned_jobs")

    def indexed(items: list, label: str) -> dict[str, dict]:
        result = {}
        for position, job in enumerate(items):
            try:
                if not isinstance(job, dict):
                    raise ValueError("job is not an object")
                token = job_token(job)
                date(int(job["year"]), int(job["month"]), 1)
                if job.get("token", token) != token:
                    issues.append(f"{label} token does not match port/year/month: {token}")
                if token in result:
                    issues.append(f"duplicate {label} token: {token}")
                result[token] = job
            except (KeyError, TypeError, ValueError) as error:
                issues.append(f"invalid {label} job {position}: {error}")
        return result

    planned = indexed(plan, "planned")
    actual = indexed(records, "recorded")
    if not planned:
        issues.append("no valid planned jobs")
    missing = sorted(set(planned) - set(actual))
    extra = sorted(set(actual) - set(planned))
    if missing:
        issues.append(f"missing planned jobs: {missing}")
    if extra:
        issues.append(f"unexpected recorded jobs: {extra}")
    if len(records) != len(plan):
        issues.append(f"completed job count {len(records)} != planned {len(plan)}")

    def compare_spec(expected: dict, observed: dict, token: str, label: str,
                     require_lists: bool) -> None:
        for field in ("physical_files", "tolerant_logical_days"):
            try:
                left, right = int(expected[field]), int(observed[field])
                if left < 1 or right != left:
                    issues.append(f"{label} {token} {field} {right} != planned {left}")
            except (KeyError, TypeError, ValueError):
                issues.append(f"{label} {token} missing/invalid {field}")
        try:
            if Path(str(expected["folder"])).resolve() != Path(str(observed["folder"])).resolve():
                issues.append(f"{label} {token} source folder differs from plan")
        except (KeyError, OSError):
            issues.append(f"{label} {token} missing/invalid source folder")
        for field in ("logical_days", "source_files"):
            values = expected.get(field)
            if values is not None:
                if not isinstance(values, list) or len(values) != len(set(values)):
                    issues.append(f"planned {token} invalid/duplicate {field}")
                elif observed.get(field) != values and (require_lists or field in observed):
                    issues.append(f"{label} {token} {field} differs from plan")
                count_field = "tolerant_logical_days" if field == "logical_days" else "physical_files"
                if isinstance(values, list) and len(values) != expected.get(count_field):
                    issues.append(f"planned {token} {field} count differs from {count_field}")

    if discovered is not None and declared_plan is not None:
        declared = indexed(declared_plan, "declared planned")
        if set(declared) != set(planned):
            issues.append("declared plan tokens differ from current source scope")
        for token in set(declared) & set(planned):
            compare_spec(planned[token], declared[token], token, "declared plan", True)
    for token in set(planned) & set(actual):
        compare_spec(planned[token], actual[token], token, "recorded", declared_plan is not None)
        if actual[token].get("status") != "PASS":
            issues.append(f"recorded {token} is not PASS: {actual[token].get('status')}")
        if actual[token].get("current_parser_logical_days") != planned[token].get("tolerant_logical_days"):
            issues.append(f"recorded {token} parser day count differs from plan")
    return {"complete": not issues, "issues": issues, "planned_jobs": plan,
            "expected_job_count": len(plan), "completed_job_count": len(records),
            "missing_jobs": missing, "extra_jobs": extra,
            "source_root": str(source_root) if source_root is not None else run_report.get("source_root")}


def validate_run(root: Path, compare_root: Path | None = None,
                 source_root: Path | None = None,
                 application_versions: tuple[str,str] | None = None) -> dict[str, object]:
    if application_versions is not None and compare_root is None:
        raise ValueError('Declared application versions require a comparison root')
    run_report = json.loads((root / "run_report.json").read_text(encoding="utf-8"))
    scope = audit_run_scope(run_report, source_root)
    report: dict[str, object] = {
        "root": str(root),
        "jobs": [],
        "issues": [],
        "workbook_count": 0,
        "detail_rows_checked": 0,
        "semantic_compare_mismatches": [],
        "allowed_provenance_differences": [],
        "scope": scope,
        "declared_application_versions": list(application_versions) if application_versions else None,
    }
    global_issues: list[str] = report["issues"]  # type: ignore[assignment]
    global_issues.extend(scope["issues"])
    if global_issues:
        report["pass"] = False
        return report
    planned_by_token = {job_token(job): job for job in scope["planned_jobs"]}
    compare_fingerprints: dict[str, dict[str, object]] = {}
    if compare_root is not None:
        for path in compare_root.rglob("*.xlsx"):
            try:
                compare_fingerprints[str(path.relative_to(compare_root))] = semantic_package_fingerprint(
                    path,application_version=application_versions[0] if application_versions else None)
            except Exception as error:
                global_issues.append(f"compare provenance validation {path}: {type(error).__name__}: {error}")
    for job in run_report["jobs"]:
        token = f"{job['port']}_{int(job['year'])}{int(job['month']):02d}"
        job_issues: list[str] = []
        if job.get("status") != "PASS":
            job_issues.append(f"run status {job.get('status')}: {job.get('error')}")
        if job.get("current_parser_logical_days") != job.get("tolerant_logical_days"):
            job_issues.append(
                f"detected days {job.get('current_parser_logical_days')} != catalog {job.get('tolerant_logical_days')}"
            )
        modern, delivery_files = expected_delivery(job)
        expected_files = [modern, *delivery_files]
        existing = sorted((root / token).rglob("*.xlsx"))
        if len(existing) != 6:
            job_issues.append(f"output count {len(existing)} != 6")
        for path in expected_files:
            if not path.is_file():
                job_issues.append(f"missing output {path}")
                continue
            if not zipfile.is_zipfile(path):
                job_issues.append(f"invalid OOXML zip {path}")
        source_day_list,calendar_days = source_calendar_plan(int(job['year']),int(job['month']),
                                                          planned_by_token[token].get('logical_days'))
        source_days = set(source_day_list)
        day_count = len(calendar_days)
        period_a_full, period_b_full, period_a_summary, period_b_summary, integrated = delivery_files
        matrices: dict[str, list[list[float | None]]] = {}
        days: list[date] = []
        if all(path.is_file() for path in delivery_files):
            try:
                days, matrices["period_a_full"] = read_summary(period_a_full, day_count)
                days_b, matrices["period_b_full"] = read_summary(period_b_full, day_count)
                days_as, matrices["period_a_summary"] = read_summary(period_a_summary, day_count)
                days_bs, matrices["period_b_summary"] = read_summary(period_b_summary, day_count)
                days_i, matrices["integrated"] = read_summary(integrated, day_count)
                if not (days == days_b == days_as == days_bs == days_i):
                    job_issues.append("summary day sequences differ")
                if days != calendar_days:
                    job_issues.append("summary day sequence differs from full calendar")
                for position,day in enumerate(days):
                    if day not in source_days and any(value is not None for matrix in matrices.values() for value in matrix[position]):
                        job_issues.append(f"missing-source calendar day contains data: {day}")
                if matrices["period_a_full"] != matrices["period_a_summary"]:
                    job_issues.append("Period A full/summary mismatch")
                if matrices["period_b_full"] != matrices["period_b_summary"]:
                    job_issues.append("Period B full/summary mismatch")
                for row_index, (left_row, right_row, integrated_row) in enumerate(
                    zip(matrices["period_a_full"], matrices["period_b_full"], matrices["integrated"])
                ):
                    for direction_index, (left, right, actual) in enumerate(
                        zip(left_row, right_row, integrated_row)
                    ):
                        expected = max(value for value in (left, right) if value is not None) if any(
                            value is not None for value in (left, right)
                        ) else None
                        if not close_enough(expected, actual):
                            job_issues.append(
                                f"integrated mismatch row {row_index + 2} direction {direction_index}"
                            )
                            break
                included = set(range(32)) if job["port"] != "KLNG" else set(range(11)) | set(range(22, 32))
                for matrix_name, matrix in matrices.items():
                    for row_index, row in enumerate(matrix):
                        for direction_index, value in enumerate(row):
                            if direction_index not in included and value is not None:
                                job_issues.append(
                                    f"{matrix_name} excluded direction {direction_index} not blank at row {row_index + 2}"
                                )
                for path in delivery_files:
                    chart_result = validate_delivery_charts(path, day_count)
                    job_issues.extend(f"{path.name}: {issue}" for issue in chart_result["issues"])
                for path in (period_a_full, period_b_full):
                    spool_dir = modern.parent / f".{modern.stem}_{job['port']}_{job['year']}_{int(job['month']):02d}_rows"
                    detail = inspect_detail_mapping(path, days, included, spool_dir, source_days)
                    report["detail_rows_checked"] += detail["detail_rows"]  # type: ignore[index,operator]
                    if detail["direction_mapping_mismatches"]:
                        job_issues.append(f"{path.name}: direction mapping mismatches {detail['direction_mapping_mismatches']}")
                    if detail["detail_max_summary_mismatches"]:
                        job_issues.append(f"{path.name}: detail/summary mismatches {detail['detail_max_summary_mismatches']}")
                    if detail["retained_row_mismatches"] or detail["retained_expected_rows"] != detail["detail_rows"]:
                        job_issues.append(f"{path.name}: full retained row mismatches {detail['retained_row_mismatches']}; "
                                          f"expected {detail['retained_expected_rows']} / observed {detail['detail_rows']}")
            except Exception as error:
                job_issues.append(f"validation exception {type(error).__name__}: {error}")

        hashes: dict[str, str] = {}
        normalized_provenance: dict[str, object] = {}
        for path in expected_files:
            if not path.is_file():
                continue
            relative = str(path.relative_to(root))
            try:
                details = semantic_package_fingerprint(path,application_version=application_versions[1] if application_versions else None)
            except Exception as error:
                job_issues.append(f"provenance validation {path.name}: {type(error).__name__}: {error}")
                continue
            fingerprint = str(details["sha256"])
            hashes[relative] = fingerprint
            normalized_provenance[relative] = details["provenance_normalizations"]
            previous = compare_fingerprints.get(relative)
            if compare_root is not None and (previous is None or previous["sha256"] != fingerprint):
                report["semantic_compare_mismatches"].append(relative)  # type: ignore[union-attr]
            if previous is not None:
                before = {(item["part"], item["field"]): item["value"]
                          for item in previous["provenance_normalizations"]}
                for item in details["provenance_normalizations"]:
                    prior = before.get((item["part"], item["field"]))
                    if prior is not None and prior != item["value"]:
                        report["allowed_provenance_differences"].append({  # type: ignore[union-attr]
                            "workbook": relative, "part": item["part"], "field": item["field"],
                            "current": item["value"], "compare": prior})
        report["workbook_count"] += len(hashes)  # type: ignore[operator]
        report["jobs"].append(  # type: ignore[union-attr]
            {
                "token": token,
                "physical_files": job["physical_files"],
                "logical_days": job["tolerant_logical_days"],
                "output_count": len(existing),
                "issues": job_issues,
                "semantic_hashes": hashes,
                "semantic_provenance_normalizations": normalized_provenance,
            }
        )
        global_issues.extend(f"{token}: {issue}" for issue in job_issues)
    report["pass"] = not global_issues and not report["semantic_compare_mismatches"]
    return report


def guard_validation_output(output: Path, root: Path, compare_root: Path | None,
                            source_root: Path | None) -> None:
    run_report_path = root / "run_report.json"
    run_report = json.loads(run_report_path.read_text(encoding="utf-8"))
    roots = [root]
    if compare_root is not None:
        roots.append(compare_root)
    if source_root is not None:
        roots.append(source_root)
    elif run_report.get("source_root"):
        roots.append(Path(str(run_report["source_root"])))
    for job in run_report.get("jobs", []):
        if isinstance(job, dict) and job.get("folder"):
            roots.append(Path(str(job["folder"])))
    protected = [run_report_path]
    for directory in roots:
        protected.extend(directory.rglob("*.xlsx"))
    guard_output(output, roots, protected)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("root", type=Path)
    parser.add_argument("--compare-root", type=Path)
    parser.add_argument('--compare-application-versions',nargs=2,metavar=('BASELINE','CURRENT'),
                        help='Allow only the exact declared Application version custom property to differ; all data/formulas/styles/charts remain exact')
    parser.add_argument("--source-root", type=Path,
                        help="Re-discover the full source scope, required for reports without a saved plan")
    parser.add_argument("--output", type=Path, required=True)
    arguments = parser.parse_args(argv)
    guard_validation_output(arguments.output, arguments.root, arguments.compare_root, arguments.source_root)
    report = validate_run(arguments.root, arguments.compare_root, arguments.source_root,
                          tuple(arguments.compare_application_versions) if arguments.compare_application_versions else None)
    atomic_json(arguments.output, report)
    print(json.dumps({
        "pass": report["pass"],
        "workbook_count": report["workbook_count"],
        "detail_rows_checked": report["detail_rows_checked"],
        "issues": len(report["issues"]),
        "semantic_compare_mismatches": len(report["semantic_compare_mismatches"]),
    }, ensure_ascii=False, indent=2))
    return 0 if report["pass"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
