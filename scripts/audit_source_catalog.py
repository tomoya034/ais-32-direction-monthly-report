from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl

from ais_monthly_app import locate_data_worksheet, parse_source_filename


TOLERANT_FILENAME_RE = re.compile(
    r"^D&TMOK[ \t]+(?P<port>[A-Z][A-Z0-9]{1,15})_"
    r"(?P<date>\d{7,8})(?:_(?P<suffix>[A-Z0-9_-]+))?\.xlsx$",
    re.IGNORECASE,
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(4 * 1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def tolerant_identity(name: str) -> tuple[str, str, str | None] | None:
    match = TOLERANT_FILENAME_RE.fullmatch(name)
    if match is None:
        return None
    raw = match.group("date")
    format_code = "%Y%m%d" if len(raw) == 8 else "%Y%m%d"
    try:
        parsed = datetime.strptime(raw, format_code).date()
    except ValueError:
        return None
    return match.group("port").upper(), parsed.isoformat(), match.group("suffix")


def inspect_source(path: Path, root: Path, with_hashes: bool) -> dict[str, object]:
    current = parse_source_filename(path.name)
    tolerant = tolerant_identity(path.name)
    record: dict[str, object] = {
        "relative_path": str(path.relative_to(root)),
        "name": path.name,
        "size": path.stat().st_size,
        "current_parser": None
        if current is None
        else {"port": current.port, "date": current.day.isoformat(), "suffix": current.suffix},
        "tolerant_parser": None
        if tolerant is None
        else {"port": tolerant[0], "date": tolerant[1], "suffix": tolerant[2]},
    }
    if with_hashes:
        record["sha256"] = sha256(path)
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            sheet, indexes = locate_data_worksheet(workbook)
            header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
            record.update(
                {
                    "status": "OK",
                    "data_sheet": sheet.title,
                    "sheet_names": workbook.sheetnames,
                    "max_row": sheet.max_row,
                    "max_column": sheet.max_column,
                    "headers": [None if value is None else str(value) for value in header],
                    "header_indexes": indexes,
                }
            )
        finally:
            workbook.close()
    except Exception as error:
        record.update({"status": "ERROR", "error": f"{type(error).__name__}: {error}"})
    return record


def summarize(records: list[dict[str, object]]) -> dict[str, object]:
    by_port_month: Counter[tuple[str, str]] = Counter()
    parsed_days: defaultdict[tuple[str, str], set[str]] = defaultdict(set)
    schema_counts: Counter[str] = Counter()
    for record in records:
        identity = record.get("tolerant_parser")
        if isinstance(identity, dict):
            port = str(identity["port"])
            day = str(identity["date"])
            month = day[:7]
            by_port_month[(port, month)] += 1
            parsed_days[(port, month)].add(day)
        headers = record.get("headers")
        if isinstance(headers, list):
            schema_counts[json.dumps(headers, ensure_ascii=False)] += 1
    return {
        "file_count": len(records),
        "total_bytes": sum(int(record["size"]) for record in records),
        "current_parser_rejected": [
            record["relative_path"] for record in records if record["current_parser"] is None
        ],
        "workbook_errors": [
            {"relative_path": record["relative_path"], "error": record.get("error")}
            for record in records
            if record.get("status") != "OK"
        ],
        "port_months": [
            {
                "port": port,
                "month": month,
                "physical_files": count,
                "logical_days": len(parsed_days[(port, month)]),
                "first_day": min(parsed_days[(port, month)]),
                "last_day": max(parsed_days[(port, month)]),
            }
            for (port, month), count in sorted(by_port_month.items())
        ],
        "header_schemas": [
            {"count": count, "headers": json.loads(headers)}
            for headers, count in schema_counts.most_common()
        ],
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("source_root", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--hashes", action="store_true")
    arguments = parser.parse_args()
    files = sorted(
        (
            path
            for path in arguments.source_root.rglob("*.xlsx")
            if not path.name.startswith("~$")
        ),
        key=lambda path: str(path.relative_to(arguments.source_root)).casefold(),
    )
    records = [
        inspect_source(path, arguments.source_root, arguments.hashes) for path in files
    ]
    payload = {"summary": summarize(records), "files": records}
    arguments.output.parent.mkdir(parents=True, exist_ok=True)
    arguments.output.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(payload["summary"], ensure_ascii=False, indent=2))
    return 1 if payload["summary"]["workbook_errors"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
