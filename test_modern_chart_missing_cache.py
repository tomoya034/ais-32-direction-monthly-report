"""Actual Modern writer chart caches: blanks/zero and both port policies."""
from datetime import date, datetime
from pathlib import Path
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
from openpyxl.utils.cell import get_column_letter
import ais_monthly_app as app

C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"


class ModernMissingCacheTests(unittest.TestCase):
    def make_result(self, blank_indexes):
        day = date(2026, 3, 1)
        result = app.empty_day_result(day)
        for index, label in enumerate(app.DIRECTION_ORDER):
            if index in blank_indexes:
                continue
            value = 0. if index == 0 else float(index + 1)
            candidate = app.Candidate(distance=value, bearing=index * 11.25 + 1,
                                      source_row=index + 2, rank=1, candidate_id=f"id-{index}",
                                      timestamp=datetime(2026, 3, 1), mmsi=123456789, channel="A",
                                      fragment="synthetic.xlsx", sheet="AIS", msg_type=1)
            selected = app.DirectionResult(direction=label, selected=value, selected_bearing=candidate.bearing,
                                           selected_rank=1, selected_candidate_id=candidate.candidate_id,
                                           cluster_count=3, candidates=[candidate], status="自動採用")
            result.directions[label] = selected
            result.period_directions[app.PERIOD_A][label] = selected
            result.period_directions[app.PERIOD_B][label] = selected
        return result

    def assert_writer(self, port, blank_indexes):
        with tempfile.TemporaryDirectory(prefix="ais-modern-missing-cache-") as temporary:
            root = Path(temporary)
            source = root / "input"
            source.mkdir()
            config = app.AppConfig(input_dir=source, output_path=root / "modern.xlsx",
                                   port=port, year=2026, month=3)
            result = self.make_result(blank_indexes)
            app.write_monthly_workbook(config, [result], [], app.build_default_decision_snapshot(config, [result]))
            with zipfile.ZipFile(config.output_path) as archive:
                charts = [ET.fromstring(archive.read(name)) for name in archive.namelist()
                          if name.startswith("xl/charts/chart") and name.endswith(".xml")]
            radar = next(chart for chart in charts if chart.find(".//" + C + "radarChart") is not None)
            # Installed XlsxWriter.show_hidden_data() deliberately omits this
            # element. Native acceptance separately asserts PlotVisibleOnly=False.
            self.assertIsNone(radar.find(".//" + C + "plotVisOnly"))
            radar_values = radar.find(".//" + C + "val/" + C + "numRef")
            self.assertEqual(radar_values.findtext(C + "f"), "總表!$AO$5:$BT$5")
            radar_cache = radar_values.find(C + "numCache")
            self.assertEqual(int(radar_cache.find(C + "ptCount").get("val")), 32)
            radar_points = {int(p.get("idx")): p.findtext(C + "v") for p in radar_cache.findall(C + "pt")}
            combo = next(chart for chart in charts if chart.find(".//" + C + "barChart") is not None)
            references = combo.findall(".//" + C + "val/" + C + "numRef")
            self.assertEqual(len(references), 3)
            included = list(range(32)) if port == "HWLN" else list(range(11)) + list(range(22, 32))
            wb = openpyxl.load_workbook(config.output_path, read_only=True, data_only=True)
            formulas = openpyxl.load_workbook(config.output_path, read_only=True, data_only=False)
            try:
                for index in range(32):
                    expected = None if index in blank_indexes or index not in included else (0. if index == 0 else float(index + 1))
                    raw = radar_points.get(index)
                    self.assertEqual(None if raw is None or raw == "" else float(raw), expected)
                    helper = wb["總表"].cell(5, 41 + index)
                    self.assertEqual(helper.value, "#N/A" if expected is None else expected)
                    self.assertEqual(helper.data_type, "e" if expected is None else "n")
                    source = f"{get_column_letter(index + 2)}5"
                    self.assertEqual(formulas["總表"].cell(5, 41 + index).value,
                                     f'=IF(COUNT({source})=0,NA(),{source})')
                    self.assertEqual(formulas["總表"].cell(4, 41 + index).value, app.DIRECTION_ORDER[index])
                for reference, column in zip(references, (36, 37, 38)):
                    cache = reference.find(C + "numCache")
                    self.assertEqual(int(cache.find(C + "ptCount").get("val")), len(included))
                    points = {int(point.get("idx")): point.findtext(C + "v") for point in cache.findall(C + "pt")}
                    for offset, index in enumerate(included):
                        expected = None if index in blank_indexes else (0. if index == 0 else float(index + 1))
                        raw = points.get(offset)
                        observed = None if raw is None or raw == "" else float(raw)
                        self.assertEqual(observed, expected, (port, index, column))
                        # Native Excel plots formula "" as zero. Chart-only
                        # helpers deliberately use NA(), but the primary data
                        # and actual statistical cells must stay blank.
                        helper = wb["總表"].cell(offset + 5, column)
                        self.assertEqual(helper.value, "#N/A" if expected is None else expected)
                        self.assertEqual(helper.data_type, "e" if expected is None else "n")
                        source = f"{get_column_letter(index + 2)}{column - 36 + 7}"
                        self.assertEqual(formulas["總表"].cell(offset + 5, column).value,
                                         f'=IF(COUNT({source})=0,NA(),{source})')
                        self.assertEqual(wb["總表"].cell(5, index + 2).value, expected)
                        for stat_row in (7, 8, 9):
                            self.assertEqual(wb["總表"].cell(stat_row, index + 2).value, expected)
            finally:
                wb.close()
                formulas.close()

    def test_trailing_blank_and_genuine_zero_hwln32_klng21(self):
        for port in ("HWLN", "KLNG"):
            with self.subTest(port=port):
                self.assert_writer(port, {31})

    def test_middle_blank_preserves_series_alignment_and_zero(self):
        for port in ("HWLN", "KLNG"):
            with self.subTest(port=port):
                self.assert_writer(port, {1})

    def test_all_empty_statistics_are_blank_not_fabricated_zero(self):
        for port in ("HWLN", "KLNG"):
            with self.subTest(port=port):
                self.assert_writer(port, set(range(32)))


if __name__ == "__main__":
    unittest.main()
