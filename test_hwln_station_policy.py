"""Regression coverage for the HWLN southern-sector data-loss diagnosis.

Fixtures deliberately contain accepted records in all eleven sectors previously
blanked by the KLNG mask. Expected numbers come from the fixture, not a regenerated
baseline. All outputs and caches are isolated in TemporaryDirectory.
"""
from __future__ import annotations

import datetime as dt
from dataclasses import replace
from pathlib import Path
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile

import openpyxl

import ais_monthly_app as app


DAY = dt.date(2026, 6, 1)
CHART_NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
DRAWING_NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
SHEET_NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
KLNG_INDEXES = tuple(range(11)) + tuple(range(22, 32))


def write_station_fixture(path: Path, *, morning_only: bool = False) -> None:
    """31 occupied bins; south index16 has morning only; index31 has no AIS."""
    workbook = openpyxl.Workbook(write_only=True)
    sheet = workbook.create_sheet("AIS")
    sheet.append(["Year", "Month", "Day", "Hour", "Minute", "Second", "channel",
                  "msg_type", "mmsi", "LONGITUDE_DESC", "bearing", "distance in nautical miles"])
    for hour, base in ((1, 50.0), (13, 70.0)):
        if morning_only and hour == 13:
            continue
        for index in range(31):
            if hour == 13 and index == 16:
                continue
            for offset in range(3):
                second = index * 3 + offset
                sheet.append([2026, 6, 1, hour, second // 60, second % 60,
                              "A", 1, 800000000 + index, "East", index * 11.25 + 1,
                              base + index - offset])
    workbook.save(path)


def radar_xml(path: Path) -> ET.Element:
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name.startswith("xl/charts/chart") and name.endswith(".xml"):
                root = ET.fromstring(archive.read(name))
                if root.find(".//c:radarChart", CHART_NS) is not None:
                    return root
    raise AssertionError(f"No radar chart: {path}")


def statistics_chart_xml(path: Path) -> ET.Element:
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name.startswith("xl/charts/chart") and name.endswith(".xml"):
                root = ET.fromstring(archive.read(name))
                if (
                    root.find(".//c:barChart", CHART_NS) is not None
                    and root.find(".//c:lineChart", CHART_NS) is not None
                ):
                    return root
    raise AssertionError(f"No combined statistics chart: {path}")


def chart_title(root: ET.Element) -> str:
    return "".join(node.text or "" for node in root.findall(".//c:title//a:t", {**CHART_NS, **DRAWING_NS}))


class StationPolicyTests(unittest.TestCase):
    def config(self, folder: Path, *, port: str = "HWLN", **kwargs) -> app.AppConfig:
        return app.AppConfig(folder, folder / "analysis.xlsx", port, 2026, 6, **kwargs)

    def prepare(self, folder: Path, *, port: str = "HWLN", **kwargs):
        source = folder / f"D&TMOK {port}_20260601_any.xlsx"
        write_station_fixture(source)
        config = self.config(folder, port=port, **kwargs)
        day_input = app.build_processing_job(config).days[0]
        result = app.process_day_input(day_input, config)
        return config, day_input, result

    def expected(self, index: int, scope: str) -> float | None:
        if index == 31 or (scope == app.PERIOD_B and index == 16):
            return None
        base = 50.0 if scope == app.PERIOD_A or index == 16 else 70.0
        return base + index

    def assert_summary(self, path: Path, indexes: tuple[int, ...], scope: str,
                       *, row: int = 2) -> None:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            values = tuple(workbook["總表"].iter_rows(min_row=row, max_row=row,
                                                    min_col=2, max_col=33, values_only=True))[0]
            for index, actual in enumerate(values):
                expected = self.expected(index, scope) if index in indexes else None
                with self.subTest(file=path.name, direction=app.LEGACY_DIRECTION_ABBREVIATIONS[index]):
                    self.assertEqual(actual, expected)
        finally:
            workbook.close()

    def test_station_resolution_and_explicit_overrides(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            for port, policy, indexes, style in (
                ("KLNG", "klng21", KLNG_INDEXES, "filled"),
                ("hwln", "all32", tuple(range(32)), "filled"),
                ("TCHG", "all32", tuple(range(32)), "filled"),
            ):
                with self.subTest(port=port):
                    config = self.config(folder, port=port)
                    config.validate()
                    self.assertEqual(config.resolved_direction_policy, policy)
                    self.assertEqual(config.direction_indexes, indexes)
                    self.assertEqual(tuple(config.directions), tuple(app.DIRECTION_ORDER[i] for i in indexes))
                    self.assertEqual(config.resolved_radar_style, style)
            manual = self.config(folder, direction_policy="klng21", radar_style="line", radar_max=150)
            self.assertEqual(manual.direction_indexes, KLNG_INDEXES)
            self.assertEqual(manual.resolved_radar_style, "line")
            klng_override = self.config(folder, port="KLNG", direction_policy="all32", radar_style="filled")
            self.assertEqual(klng_override.direction_indexes, tuple(range(32)))
            self.assertEqual(klng_override.resolved_radar_style, "filled")

    def test_invalid_policy_style_and_axis_are_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            for kwargs in ({"direction_policy": "guess"}, {"radar_style": "surface"},
                           *({"radar_max": value} for value in (0, -1, float("nan"), float("inf")))):
                with self.subTest(kwargs=kwargs), self.assertRaises(ValueError):
                    self.config(folder, **kwargs).validate()

    def test_hwln_values_survive_every_workbook_and_decision_scope(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            config, _, result = self.prepare(folder, legacy_output_path=folder / "legacy.xlsx")
            snapshot = app.build_default_decision_snapshot(config, [result])
            self.assertEqual(len(snapshot.decisions), 3 * 32)
            resolved = app.resolve_decision_values(snapshot, [result], config)
            for scope in (app.LOGICAL_DAY, app.PERIOD_A, app.PERIOD_B):
                for index, direction in enumerate(app.DIRECTION_ORDER):
                    self.assertEqual(resolved[(scope, DAY, direction)], self.expected(index, scope))
            app.write_monthly_workbook(config, [result], [], snapshot)
            app.write_legacy_workbook(config, [result])
            paths = app.write_delivery_workbooks(config, [result], snapshot, root=folder / "delivery")
            self.assert_summary(config.output_path, tuple(range(32)), app.LOGICAL_DAY, row=5)
            self.assert_summary(config.legacy_output_path, tuple(range(32)), app.LOGICAL_DAY)
            for path, scope in ((paths.period_a_full, app.PERIOD_A),
                                (paths.period_a_summary, app.PERIOD_A),
                                (paths.period_b_full, app.PERIOD_B),
                                (paths.period_b_summary, app.PERIOD_B),
                                (paths.integrated_summary, app.LOGICAL_DAY)):
                self.assert_summary(path, tuple(range(32)), scope)
            workbook = openpyxl.load_workbook(config.output_path, read_only=True, data_only=True)
            try:
                ledger = list(workbook["決策台帳"].iter_rows(min_row=2, values_only=True))
                self.assertEqual(len(ledger), 96)
                self.assertEqual({row[2] for row in ledger}, set(app.DIRECTION_ORDER))
                for index in range(11, 22):
                    self.assertTrue(any(row[2] == app.DIRECTION_ORDER[index] for row in ledger))
            finally:
                workbook.close()
            for path, scope in ((paths.period_a_full, app.PERIOD_A), (paths.period_b_full, app.PERIOD_B)):
                workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
                try:
                    details = list(workbook["6月1日"].iter_rows(min_row=2, max_col=3, values_only=True))
                    for index in range(11, 22):
                        retained = [row[0] for row in details if row[2] == app.DIRECTION_ORDER[index]]
                        expected = self.expected(index, scope)
                        if expected is None:
                            self.assertEqual(retained, [])
                        else:
                            self.assertEqual(max(retained), expected)
                            self.assertEqual(len(retained), 3)
                finally:
                    workbook.close()

    def test_klng_unchanged_21_direction_delivery_with_southern_raw_records(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            config, _, result = self.prepare(folder, port="KLNG", legacy_output_path=folder / "legacy.xlsx")
            snapshot = app.build_default_decision_snapshot(config, [result])
            self.assertEqual(len(snapshot.decisions), 63)
            self.assertEqual({decision.direction for decision in snapshot.decisions},
                             {app.DIRECTION_ORDER[i] for i in KLNG_INDEXES})
            app.write_monthly_workbook(config, [result], [], snapshot)
            app.write_legacy_workbook(config, [result])
            paths = app.write_delivery_workbooks(config, [result], snapshot, root=folder / "delivery")
            self.assert_summary(config.output_path, KLNG_INDEXES, app.LOGICAL_DAY, row=5)
            self.assert_summary(config.legacy_output_path, KLNG_INDEXES, app.LOGICAL_DAY)
            for path, scope in ((paths.period_a_full, app.PERIOD_A), (paths.period_a_summary, app.PERIOD_A),
                                (paths.period_b_full, app.PERIOD_B), (paths.period_b_summary, app.PERIOD_B),
                                (paths.integrated_summary, app.LOGICAL_DAY)):
                self.assert_summary(path, KLNG_INDEXES, scope)
            root = radar_xml(paths.integrated_summary)
            self.assertEqual(root.find(".//c:radarStyle", CHART_NS).get("val"), "filled")
            series = root.findall(".//c:radarChart/c:ser", CHART_NS)
            fills = [
                item.find(".//c:spPr/a:solidFill/a:srgbClr", {**CHART_NS, **DRAWING_NS})
                for item in series
            ]
            self.assertEqual(
                [item.get("val") for item in fills],
                [
                    color.lstrip("#")
                    for color in app.LEGACY_RADAR_PALETTE[: len(fills)]
                ],
            )
            self.assertTrue(
                all(
                    item.find("a:alpha", DRAWING_NS).get("val") == "40000"
                    for item in fills
                )
            )
            statistics = statistics_chart_xml(paths.integrated_summary)
            bars = statistics.findall(".//c:barChart/c:ser", CHART_NS)
            line = statistics.find(".//c:lineChart/c:ser", CHART_NS)
            self.assertEqual(
                [
                    item.find(
                        ".//c:spPr/a:solidFill/a:srgbClr",
                        {**CHART_NS, **DRAWING_NS},
                    ).get("val")
                    for item in bars
                ],
                ["00B0F0", "ED7D31"],
            )
            line_shape = line.find(".//c:spPr/a:ln", {**CHART_NS, **DRAWING_NS})
            self.assertEqual(line_shape.get("w"), "28575")
            self.assertEqual(
                line_shape.find("a:solidFill/a:srgbClr", DRAWING_NS).get("val"),
                "FF0000",
            )

    def test_delivery_statistics_helpers_and_chart_cache_match_full_and_summary(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root_folder = Path(temporary)
            for port, indexes in (("HWLN", tuple(range(32))), ("KLNG", KLNG_INDEXES)):
                with self.subTest(port=port):
                    folder = root_folder / port
                    folder.mkdir()
                    config, _, result = self.prepare(folder, port=port)
                    snapshot = app.build_default_decision_snapshot(config, [result])
                    paths = app.write_delivery_workbooks(
                        config, [result], snapshot, root=folder / "delivery"
                    )
                    self.assertIn(
                        "2026-06 AIS通訊最近最遠距離(浬)",
                        chart_title(statistics_chart_xml(paths.integrated_summary)),
                    )
                    for _scope_title, full_path, summary_path in (
                        (
                            "Historical Period A 00:00–11:59",
                            paths.period_a_full,
                            paths.period_a_summary,
                        ),
                        (
                            "Historical Period B 12:00–23:59",
                            paths.period_b_full,
                            paths.period_b_summary,
                        ),
                        (
                            "Historical Period A/B 整合 MAX",
                            paths.integrated_summary,
                            paths.integrated_summary,
                        ),
                    ):
                        helper_matrices = []
                        chart_details = []
                        for path in (full_path, summary_path):
                            workbook = openpyxl.load_workbook(
                                path, read_only=True, data_only=True
                            )
                            try:
                                helper = list(
                                    workbook["總表"].iter_rows(
                                        min_row=1,
                                        max_row=len(indexes) + 1,
                                        min_col=35,
                                        max_col=38,
                                        values_only=True,
                                    )
                                )
                            finally:
                                workbook.close()
                            self.assertEqual(
                                helper[0], ("納入方位", "MAX", "MIN", "AVERAGE")
                            )
                            self.assertEqual(
                                tuple(row[0] for row in helper[1:]),
                                tuple(
                                    app.LEGACY_DIRECTION_ABBREVIATIONS[i]
                                    for i in app.legacy_chart_direction_indexes(config)
                                ),
                            )
                            for column in range(1, 4):
                                self.assertTrue(
                                    any(row[column] is not None for row in helper[1:])
                                )
                                self.assertTrue(
                                    any(row[column] is None for row in helper[1:])
                                )
                            helper_matrices.append(helper)

                            chart = statistics_chart_xml(path)
                            self.assertIn(
                                "2026-06 AIS通訊最近最遠距離(浬)",
                                chart_title(chart),
                            )
                            plot_visible_only = chart.find(".//c:plotVisOnly", CHART_NS)
                            # XlsxWriter's public show_hidden_data() encodes false
                            # by omitting plotVisOnly; its default emitted form is val=1.
                            self.assertIsNone(plot_visible_only)
                            series = [
                                *chart.findall(".//c:barChart/c:ser", CHART_NS),
                                *chart.findall(".//c:lineChart/c:ser", CHART_NS),
                            ]
                            self.assertEqual(len(series), 3)
                            formulas_and_caches = []
                            for series_index, item in enumerate(series):
                                category_formula = item.find("c:cat/*/c:f", CHART_NS)
                                value_formula = item.find("c:val/c:numRef/c:f", CHART_NS)
                                self.assertIsNotNone(category_formula)
                                self.assertIsNotNone(value_formula)
                                point_count = item.find(
                                    "c:val/c:numRef/c:numCache/c:ptCount", CHART_NS
                                )
                                self.assertIsNotNone(point_count)
                                self.assertEqual(
                                    int(point_count.get("val")), len(indexes)
                                )
                                cache = {
                                    int(point.get("idx")): float(
                                        point.find("c:v", CHART_NS).text
                                    )
                                    for point in item.findall(
                                        "c:val/c:numRef/c:numCache/c:pt", CHART_NS
                                    )
                                }
                                # Chart points must preserve the same blank vs
                                # actual-zero distinction as their source helper
                                # cells; absent observations are not radius zero.
                                expected_cache = {
                                    offset: float(row[series_index + 1])
                                    for offset, row in enumerate(helper[1:])
                                    if row[series_index + 1] is not None
                                }
                                self.assertEqual(cache, expected_cache)
                                for offset, row in enumerate(helper[1:]):
                                    if row[series_index + 1] is None:
                                        self.assertNotIn(offset, cache)
                                formulas_and_caches.append(
                                    (category_formula.text, value_formula.text, cache)
                                )

                            last_row = len(indexes) + 1
                            self.assertEqual(
                                [(category, value) for category, value, _cache in formulas_and_caches],
                                [
                                    (
                                        f"總表!$AI$2:$AI${last_row}",
                                        f"總表!${column}$2:${column}${last_row}",
                                    )
                                    for column in ("AJ", "AK", "AL")
                                ],
                            )
                            chart_details.append(formulas_and_caches)

                        self.assertEqual(helper_matrices[0], helper_matrices[1])
                        self.assertEqual(chart_details[0], chart_details[1])

    def test_future_functions_are_excel_serialized_without_changing_retained_maxima(self) -> None:
        """Check OOXML, not just XlsxWriter's precomputed values; no recalc claim."""
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            config, _, result = self.prepare(folder)
            snapshot = app.build_default_decision_snapshot(config, [result])
            app.write_monthly_workbook(config, [result], [], snapshot)
            paths = app.write_delivery_workbooks(config, [result], snapshot, root=folder / "delivery")
            for path, scope in ((paths.period_a_full, app.PERIOD_A), (paths.period_b_full, app.PERIOD_B)):
                workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
                try:
                    sheet = workbook["6月1日"]
                    details = list(sheet.iter_rows(min_row=2, max_col=3, values_only=True))
                    with zipfile.ZipFile(path) as archive:
                        daily_xml = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
                    for index, direction in enumerate(app.DIRECTION_ORDER):
                        column = openpyxl.utils.get_column_letter(index + 8)
                        formula = daily_xml.find(f".//s:c[@r='{column}2']/s:f", SHEET_NS)
                        self.assertIsNotNone(formula)
                        self.assertEqual(formula.text,
                                         f'IF(COUNTIF($C:$C,{column}$1)=0,"",'
                                         f'_xlfn.MAXIFS($A:$A,$C:$C,{column}$1))')
                        retained = [row[0] for row in details if row[2] == direction]
                        actual_max = max(retained) if retained else None
                        self.assertEqual(actual_max, self.expected(index, scope))
                        self.assertEqual(sheet[f"{column}2"].value, actual_max)
                finally:
                    workbook.close()
            # STDEV.S also requires the future-function prefix; checking its
            # cached blank would otherwise miss Excel's #NAME? on recalculation.
            for path in (config.output_path, *paths.all_files()):
                with self.subTest(path=path.name), zipfile.ZipFile(path) as archive:
                    formulas = [node.text for name in archive.namelist()
                                if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
                                for node in ET.fromstring(archive.read(name)).findall(".//s:f", SHEET_NS)]
                    stdev = [formula for formula in formulas if "STDEV.S(" in formula]
                    self.assertEqual(len(stdev), 32)
                    for formula in stdev:
                        self.assertIn("_xlfn.STDEV.S(", formula)
                        self.assertNotIn("_xlfn._xlfn.", formula)

    def test_blank_links_and_statistics_are_guarded_for_native_recalculation(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            config, _, result = self.prepare(
                folder, legacy_output_path=folder / "legacy.xlsx"
            )
            snapshot = app.build_default_decision_snapshot(config, [result])
            app.write_monthly_workbook(config, [result], [], snapshot)
            app.write_legacy_workbook(config, [result])
            paths = app.write_delivery_workbooks(
                config, [result], snapshot, root=folder / "delivery"
            )

            workbook = openpyxl.load_workbook(config.output_path, data_only=False)
            try:
                summary = workbook["總表"]
                self.assertEqual(
                    summary["AG5"].value,
                    '=IF(COUNT(\'6月1日\'!AM4)=0,"",\'6月1日\'!AM4)',
                )
                self.assertEqual(
                    summary["AG7"].value,
                    '=IF(COUNT(AG5:AG5)<1,"",MAX(AG5:AG5))',
                )
                self.assertEqual(
                    summary["AG10"].value,
                    '=IF(COUNT(AG5:AG5)<2,"",_xlfn.STDEV.S(AG5:AG5))',
                )
                self.assertEqual(
                    summary["AJ36"].value,
                    '=IF(COUNT(AG7)=0,NA(),AG7)',
                )
                # This error is deliberate and restricted to chart-only
                # helpers: native empty-string series falsely plotted zero.
                # Primary cells above retain their exact blank contracts.
                for index in range(32):
                    source = f"{openpyxl.utils.get_column_letter(index + 2)}5"
                    self.assertEqual(summary.cell(5, 41 + index).value,
                                     f'=IF(COUNT({source})=0,NA(),{source})')
                self.assertTrue(summary.column_dimensions["AO"].hidden)
            finally:
                workbook.close()
            cached = openpyxl.load_workbook(config.output_path, data_only=True)
            try:
                self.assertIsNone(cached["總表"]["AG5"].value)
                self.assertIsNone(cached["總表"]["AG7"].value)
                for address in ("AJ36", "BT5"):
                    self.assertEqual(cached["總表"][address].value, "#N/A")
                    self.assertEqual(cached["總表"][address].data_type, "e")
            finally:
                cached.close()

            workbook = openpyxl.load_workbook(
                config.legacy_output_path, data_only=False
            )
            try:
                self.assertEqual(
                    workbook["總表"]["AG2"].value,
                    '=IF(COUNT(\'6月1日\'!AM2)=0,"",\'6月1日\'!AM2)',
                )
            finally:
                workbook.close()

            workbook = openpyxl.load_workbook(paths.period_b_full, data_only=False)
            try:
                summary = workbook["總表"]
                self.assertEqual(
                    summary["R2"].value,
                    '=IF(COUNT(\'6月1日\'!X2)=0,"",\'6月1日\'!X2)',
                )
                self.assertEqual(
                    summary["R4"].value,
                    '=IF(COUNT(R2:R2)<1,"",MAX(R2:R2))',
                )
                self.assertEqual(
                    summary["R7"].value,
                    '=IF(COUNT(R2:R2)<2,"",_xlfn.STDEV.S(R2:R2))',
                )
            finally:
                workbook.close()

    def test_policy_changes_invalidate_cache_and_decision_snapshot(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            config, day_input, result = self.prepare(folder)
            app.save_cached_day(day_input, config, result)
            self.assertIsNotNone(app.load_cached_day(day_input, config))
            other = replace(config, direction_policy="klng21")
            self.assertIsNone(app.load_cached_day(day_input, other))
            self.assertNotEqual(app.source_manifest_hash(config, [result]),
                                app.source_manifest_hash(other, [result]))
            snapshot = app.build_default_decision_snapshot(config, [result])
            with self.assertRaises(ValueError):
                app.write_delivery_workbooks(other, [result], snapshot, root=folder / "wrong")
            self.assertFalse(any((folder / "wrong").glob("*.xlsx")))

    def test_pre_policy_review_cannot_silently_finalize_old_hwln_mask(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            path = folder / "v151-hwln-review.xlsx"
            workbook = openpyxl.Workbook()
            workbook.active.title = "系統資料"
            workbook.active.append(["Schema", "AIS_V15_REVIEW_2"])
            workbook.active.append(["Port", "HWLN"])
            workbook.active.append(["Period", "2026-06"])
            workbook.create_sheet("決策台帳")
            workbook.create_sheet("候選清單")
            workbook.save(path)
            workbook.close()
            original = path.read_bytes()
            for reader in (app.read_review_decisions, app.config_from_review_workbook,
                           app.finalize_review_workbook):
                with self.subTest(reader=reader.__name__), self.assertRaisesRegex(ValueError, "港別方位政策.*重新分析"):
                    reader(path)
            self.assertEqual(path.read_bytes(), original)
            self.assertEqual(list(folder.iterdir()), [path])

    def test_review_roundtrip_retains_explicit_settings(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            config, _, result = self.prepare(folder, direction_policy="klng21", radar_style="line", radar_max=150)
            snapshot = app.build_default_decision_snapshot(config, [result])
            app.write_monthly_workbook(config, [result], [], snapshot)
            loaded = app.config_from_review_workbook(config.output_path)
            self.assertEqual(loaded.direction_policy, "klng21")
            self.assertEqual(loaded.radar_style, "line")
            self.assertEqual(loaded.radar_max, 150)
            self.assertEqual(app.read_review_decisions(config.output_path), snapshot)
            self.assertEqual(app.source_manifest_hash(config, [result]), app.source_manifest_hash(loaded, [result]))

    def test_hwln_chart_uses_all32_and_automatic_axis_not_data_cutoff(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            config, _, result = self.prepare(folder, max_distance=500, legacy_output_path=folder / "legacy.xlsx")
            snapshot = app.build_default_decision_snapshot(config, [result])
            app.write_monthly_workbook(config, [result], [], snapshot)
            app.write_legacy_workbook(config, [result])
            paths = app.write_delivery_workbooks(config, [result], snapshot, root=folder / "delivery")
            self.assertIn("Modern logical-day", chart_title(radar_xml(config.output_path)))
            for path in (config.output_path, config.legacy_output_path, *paths.all_files()):
                with self.subTest(path=path.name):
                    root = radar_xml(path)
                    self.assertEqual(root.find(".//c:radarStyle", CHART_NS).get("val"), "filled")
                    self.assertIsNone(root.find(".//c:valAx/c:scaling/c:max", CHART_NS))
                    for series in root.findall(".//c:radarChart/c:ser", CHART_NS):
                        # constant_memory legacy workbooks use numRef for the
                        # formula; non-streaming summaries use strRef.
                        categories = series.find("c:cat/*/c:f", CHART_NS)
                        values = series.find("c:val/c:numRef/c:f", CHART_NS)
                        self.assertIsNotNone(categories)
                        self.assertRegex(categories.text, r"\$B\$\d+:\$AG\$\d+$")
                        if path == config.output_path:
                            self.assertEqual(values.text, "總表!$AO$5:$BT$5")
                            self.assertIsNone(root.find(".//c:plotVisOnly", CHART_NS))
                        elif path in (paths.period_a_full, paths.period_b_full, paths.integrated_summary):
                            self.assertEqual(values.text, "總表!$AO$2:$BT$2")
                            self.assertIsNone(root.find(".//c:plotVisOnly", CHART_NS))
                        else:
                            self.assertRegex(values.text, r"\$B\$\d+:\$AG\$\d+$")
                        cached = {int(point.get("idx")): float(point.find("c:v", CHART_NS).text)
                                  for point in series.findall("c:val/c:numRef/c:numCache/c:pt", CHART_NS)}
                        # A visually empty direction must not become a zero
                        # radius vertex in the chart's cached data.
                        self.assertNotIn(31, cached)
                        scope = (app.PERIOD_A if path in (paths.period_a_full, paths.period_a_summary)
                                 else app.PERIOD_B if path in (paths.period_b_full, paths.period_b_summary)
                                 else app.LOGICAL_DAY)
                        expected_cached = {i: self.expected(i, scope) for i in range(32)
                                           if self.expected(i, scope) is not None}
                        self.assertEqual(cached, expected_cached)

    def test_requested_radar_axis_is_written_but_clipping_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            config, _, result = self.prepare(folder, radar_max=150,
                                             legacy_output_path=folder / "legacy.xlsx")
            snapshot = app.build_default_decision_snapshot(config, [result])
            paths = app.write_delivery_workbooks(config, [result], snapshot, root=folder / "good")
            root = radar_xml(paths.integrated_summary)
            self.assertEqual(float(root.find(".//c:valAx/c:scaling/c:max", CHART_NS).get("val")), 150.0)
            clipped = replace(config, radar_max=80)
            clipped_snapshot = app.build_default_decision_snapshot(clipped, [result])
            with self.assertRaisesRegex(ValueError, "截掉"):
                app.write_delivery_workbooks(clipped, [result], clipped_snapshot, root=folder / "clipped")
            self.assertFalse(any((folder / "clipped").glob("*.xlsx")))
            with self.assertRaisesRegex(ValueError, "截掉"):
                app.write_monthly_workbook(clipped, [result], [], clipped_snapshot)
            self.assertFalse(clipped.output_path.exists())
            with self.assertRaisesRegex(ValueError, "截掉"):
                app.write_legacy_workbook(clipped, [result])
            self.assertFalse(clipped.legacy_output_path.exists())

    def test_entire_empty_period_exports_blank_chart_without_zero_data(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            source = folder / "D&TMOK HWLN_20260601_any.xlsx"
            write_station_fixture(source, morning_only=True)
            config = self.config(folder)
            result = app.process_day_input(app.build_processing_job(config).days[0], config)
            results = [result, app.empty_day_result(dt.date(2026, 6, 2)),
                       app.empty_day_result(dt.date(2026, 6, 3))]
            snapshot = app.build_default_decision_snapshot(config, results)
            paths = app.write_delivery_workbooks(config, results, snapshot, root=folder / "delivery")
            self.assertEqual(len(paths.all_files()), 5)
            for path in paths.all_files():
                self.assertTrue(path.is_file())
                with zipfile.ZipFile(path) as archive:
                    self.assertIsNone(archive.testzip())
            for path in (paths.period_b_full, paths.period_b_summary):
                with self.subTest(path=path.name):
                    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
                    try:
                        rows = list(workbook["總表"].iter_rows(min_row=2, max_row=4,
                                                            min_col=2, max_col=33, values_only=True))
                        self.assertEqual(rows, [(None,) * 32] * 3)
                    finally:
                        workbook.close()
                    root = radar_xml(path)
                    series = root.findall(".//c:radarChart/c:ser", CHART_NS)
                    self.assertEqual(len(series), 1)
                    self.assertEqual(series[0].find("c:tx/c:v", CHART_NS).text, "本時段無有效資料")
                    formula = series[0].find("c:val/c:numRef/c:f", CHART_NS).text
                    self.assertEqual(formula, "總表!$AO$2:$BT$2" if path in
                                     (paths.period_a_full, paths.period_b_full, paths.integrated_summary)
                                     else "總表!$B$2:$AG$2")
                    self.assertEqual(series[0].findall("c:val/c:numRef/c:numCache/c:pt", CHART_NS), [])
            # Occupied periods retain their actual date series; the two missing
            # dates remain table rows without drawing empty filled polygons.
            for path in (paths.period_a_full, paths.period_a_summary, paths.integrated_summary):
                root = radar_xml(path)
                series = root.findall(".//c:radarChart/c:ser", CHART_NS)
                self.assertEqual(len(series), 1)
                self.assertEqual(series[0].find("c:tx/c:v", CHART_NS).text, "6月1日")


if __name__ == "__main__":
    unittest.main()
