from datetime import date, datetime
import json
from pathlib import Path
import tempfile
import unittest
from unittest import mock
import zipfile

import openpyxl
from openpyxl.utils.cell import get_column_letter
import xlsxwriter

from scripts import audit_default_workbook_content as audit


class DefaultContentTests(unittest.TestCase):
    def setUp(self):
        # Existing adversarial helper fixtures intentionally model a miniature
        # one-day calendar, not a real partial-month output. Genuine 28/29/30/31
        # day/source-gap contracts are exercised without mocking in
        # test_calendar_source_validation.py using the production writer.
        calendar_fixture = mock.patch.object(audit.calendar,'monthrange',return_value=(3,1))
        calendar_fixture.start()
        self.addCleanup(calendar_fixture.stop)
        self.temporary = tempfile.TemporaryDirectory(prefix="ais-default-content-")
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.modern = self.root / "modern.xlsx"
        self.delivery = self.root / "delivery"
        self.delivery.mkdir()
        self.cache = self.root / ".month_cache"
        self.cache.mkdir()
        self.day = date(2026, 1, 1)
        self.scopes = {scope: {label: {"selected": 0.0 if index == 0 else None,
                                      "candidates": ([{"rank": 1, "candidate_id": "north-id", "distance": 0.,
                                          "bearing": 0., "timestamp": "2026-01-01T00:00:00", "mmsi": None,
                                          "channel": "A", "fragment": "source.xlsx", "sheet": "AIS",
                                          "source_row": 2, "msg_type": 1}] if index == 0 else []),
                                      "selected_candidate_id": "north-id" if index == 0 else None}
                               for index, label in enumerate(audit.LABELS)}
                       for scope in ("logical_day", "period_a", "period_b")}
        payload = {"signature": {"port": "HWLN"}, "result": {
            "day": self.day.isoformat(), "directions": self.scopes["logical_day"],
            "period_directions": {scope: self.scopes[scope] for scope in ("period_a", "period_b")}}}
        self.cache_file = self.cache / "20260101.json"
        self.cache_file.write_text(json.dumps(payload), encoding="utf-8")
        self.job = {"token": "HWLN_202601", "modern": str(self.modern),
                    "delivery": str(self.delivery), "logical_days": [self.day.isoformat()], "port": "HWLN",
                    "year":2026,"month":1}

    def write_book(self, path, modern=False, corrupt=None):
        workbook = xlsxwriter.Workbook(path)
        sheet = workbook.add_worksheet("總表")
        header_row = 3 if modern else 0
        sheet.write_row(header_row, 1, audit.LABELS if modern else audit.ABBREVIATIONS)
        format_date = workbook.add_format({"num_format": "yyyy-mm-dd"})
        sheet.write_datetime(header_row + 1, 0, datetime(2026, 1, 1), format_date)
        sheet.write(header_row + 1, 1, 99 if corrupt == "summary" else 0)
        stat_row = 6 if modern else 3
        for offset, label in enumerate(("MAX", "MIN", "AVERAGE", "STDEV.S")):
            sheet.write(stat_row + offset, 0, label)
        for offset in range(3):
            sheet.write(stat_row + offset, 1, 0)
        if modern:
            candidates = workbook.add_worksheet("候選清單")
            candidates.write_row(0, 0, audit.CANDIDATE_HEADERS)
            for position, scope in enumerate(("logical_day", "period_a", "period_b"), 1):
                candidates.write_row(position, 0, [scope, None, audit.LABELS[0], 1, "north-id", 0, 0,
                                                   None, None, "A", "source.xlsx", "AIS",
                                                   999 if corrupt == "candidate" else 2, 1])
                candidates.write_datetime(position, 1, datetime(2026, 1, 1), format_date)
                candidates.write_datetime(position, 7, datetime(2026, 1, 1), format_date)
            ledger = workbook.add_worksheet("決策台帳")
            ledger.write(0, 0, "Scope")
            row = 1
            for scope in ("logical_day", "period_a", "period_b"):
                for index, label in enumerate(audit.LABELS):
                    selected = self.scopes[scope][label]
                    ledger.write(row, 0, scope)
                    ledger.write_datetime(row, 1, datetime(2026, 1, 1), format_date)
                    ledger.write(row, 2, label)
                    ledger.write(row, 3, selected["selected_candidate_id"])
                    ledger.write(row, 4, selected["selected"])
                    ledger.write(row, 5, "wrong-id" if corrupt == "ledger" and row == 1 else selected["selected_candidate_id"])
                    number = row + 1
                    lookup = (f"IFERROR(INDEX('候選清單'!$F$2:$F$4,"
                              f"MATCH(F{number},'候選清單'!$E$2:$E$4,0)),\"#INVALID_CANDIDATE\")")
                    formula = (f'=IF(ISNUMBER(G{number}),G{number},IF(F{number}="","",{lookup}))'
                               if scope == "logical_day" else
                               f'=IF(G{number}<>"","#INVALID_PERIOD_NUMERIC",IF(F{number}="","",{lookup}))')
                    if corrupt == "ledger_formula" and row == 1:
                        formula = formula.replace("MATCH(F2,", "MATCH(F3,")
                    ledger.write_formula(row, 7, formula, None, selected["selected"] if selected["selected"] is not None else "")
                    row += 1
            daily = workbook.add_worksheet("1月1日")
            daily.write(1, 7, 0)
            for direction in range(32):
                cached = 0 if direction == 0 else ""
                daily.write_formula(3, 7 + direction, f"='決策台帳'!H{direction + 2}", None,
                                    77 if corrupt == "daily" and direction == 0 else cached)
                source = f"'1月1日'!{get_column_letter(direction + 8)}4"
                sheet.write_formula(4, direction + 1, f'=IF(COUNT({source})=0,"",{source})', None,
                                    99 if corrupt == "summary" and direction == 0 else cached)
                column = get_column_letter(direction + 2)
                for stat, function in enumerate(("MAX", "MIN", "AVERAGE", "_xlfn.STDEV.S")):
                    source = f"{column}5:{column}5"
                    sheet.write_formula(6 + stat, direction + 1,
                                        f'=IF(COUNT({source})<{2 if stat == 3 else 1},"",{function}({source}))',
                                        None, cached if stat < 3 else "")
                sheet.write(direction + 4, 34, audit.LABELS[direction])
                for stat in range(3):
                    source = f"{column}{7 + stat}"
                    formula = f'=IF(COUNT({source})=0,NA(),{source})'
                    helper_cached = 0 if direction == 0 else "#N/A"
                    if corrupt == "numeric_source_na" and direction == 0 and stat == 0:
                        helper_cached = "#N/A"
                    if corrupt == "missing_as_blank" and direction == 1 and stat == 0:
                        helper_cached = ""
                    if corrupt == "shifted_helper_source" and direction == 1 and stat == 0:
                        formula = '=IF(COUNT(D7)=0,NA(),D7)'
                    if corrupt == "primary_source_na" and direction == 1 and stat == 0:
                        sheet.write_formula(6, 2, '=IF(COUNT(C5:C5)<1,"",MAX(C5:C5))', None, "#N/A")
                    if corrupt == "helper_error_as_text" and direction == 1 and stat == 0:
                        sheet.write_string(direction + 4, 35 + stat, "#N/A")
                    else:
                        sheet.write_formula(direction + 4, 35 + stat, formula, None, helper_cached)
            if corrupt == "unexpected_na_outside_helper":
                daily.write_formula(50, 0, "=NA()", None, "#N/A")
            if corrupt == "unexpected_other_error":
                sheet.write_formula(50, 0, "=1/0", None, "#DIV/0!")
            sheet.write_row(3, 40, audit.LABELS)
            sheet.set_column(40, 71, 10, None, {"hidden": True})
            for direction in range(32):
                source = f"{get_column_letter(direction + 2)}5"
                formula = f'=IF(COUNT({source})=0,NA(),{source})'
                value = 0 if direction == 0 else "#N/A"
                if corrupt == "radar_numeric_source_na" and direction == 0:
                    value = "#N/A"
                if corrupt == "radar_missing_as_zero" and direction == 1:
                    value = 0
                if corrupt == "radar_shifted_ref" and direction == 1:
                    formula = '=IF(COUNT(D5)=0,NA(),D5)'
                sheet.write_formula(4, 40 + direction, formula, None, value)
        if not modern:
            integrated = path.stem.endswith("每日較大值整合總表")
            full = not integrated and not path.stem.endswith(" 總表")
            if full:
                daily = workbook.add_worksheet("1月1日")
                for direction in range(32):
                    column = get_column_letter(direction + 8)
                    daily.write_formula(1, direction + 7,
                                        f'=IF(COUNTIF($C:$C,{column}$1)=0,"",_xlfn.MAXIFS($A:$A,$C:$C,{column}$1))',
                                        None, 0 if direction == 0 else "")
                    reference = f"'1月1日'!{column}2"
                    sheet.write_formula(1, direction + 1, f'=IF(COUNT({reference})=0,"",{reference})',
                                        None, 0 if direction == 0 else "")
            if integrated:
                for name in ("Period A", "Period B"):
                    period = workbook.add_worksheet(name)
                    period.write(1, 1, 0)
                for direction in range(32):
                    coordinate = get_column_letter(direction + 2) + "2"
                    sources = f"'Period A'!{coordinate},'Period B'!{coordinate}"
                    sheet.write_formula(1, direction + 1, f'=IF(COUNT({sources})=0,"",MAX({sources}))',
                                        None, 0 if direction == 0 else "")
            for direction in range(32):
                column = get_column_letter(direction + 2)
                source = f"{column}2:{column}2"
                for stat, function in enumerate(("MAX", "MIN", "AVERAGE", "_xlfn.STDEV.S")):
                    sheet.write_formula(3 + stat, direction + 1,
                                        f'=IF(COUNT({source})<{2 if stat == 3 else 1},"",{function}({source}))',
                                        None, 0 if direction == 0 and stat < 3 else "")
            for direction, name in enumerate(audit.ABBREVIATIONS):
                sheet.write(direction + 1, 34, name)
                if direction == 0:
                    sheet.write_row(direction + 1, 35, [0, 0, 0])
            if full or integrated:
                sheet.write_row(0, 40, audit.ABBREVIATIONS)
                sheet.set_column(40, 71, None, None, {"hidden": True})
                for direction in range(32):
                    source = get_column_letter(direction + 2) + "2"
                    formula = f'=IF(COUNT({source})=0,NA(),{source})'
                    cached = 0 if direction == 0 else "#N/A"
                    if corrupt == "historical_numeric_source_na" and direction == 0:
                        cached = "#N/A"
                    if corrupt == "historical_missing_as_zero" and direction == 1:
                        cached = 0
                    if corrupt == "historical_shifted_helper" and direction == 1:
                        formula = '=IF(COUNT(D2)=0,NA(),D2)'
                    if corrupt == "historical_literal_error" and direction == 1:
                        sheet.write_string(1, direction + 40, "#N/A")
                    else:
                        sheet.write_formula(1, direction + 40, formula, None, cached)
            if corrupt == "historical_unexpected_na":
                sheet.write_formula(80, 0, "=NA()", None, "#N/A")
        for index in range(2):
            if index == 1:
                chart = workbook.add_chart({"type": "column"})
                helper_start = 4 if modern else 1
                for name, column in ((("最遠", 35), ("最近", 36)) if modern else (("MAX", 35), ("MIN", 36))):
                    chart.add_series({"name": name, "categories": ["總表", helper_start, 34, helper_start + 31, 34],
                                      "values": ["總表", helper_start, column, helper_start + 31, column],
                                      "values_data": [0] + [None] * 31})
                mean = workbook.add_chart({"type": "line"})
                mean.add_series({"name": "平均" if modern else "平均值", "categories": ["總表", helper_start, 34, helper_start + 31, 34],
                                 "values": ["總表", helper_start, 37, helper_start + 31, 37], "values_data": [0] + [None] * 31})
                chart.combine(mean)
            else:
                chart = workbook.add_chart({"type": "radar"})
                chart.add_series({"name": "1/1" if modern else "1月1日", "categories": ["總表", header_row, 1, header_row, 32],
                                  "values": ["總表", header_row + 1, 40 if modern or full or integrated else 1, header_row + 1, 71 if modern or full or integrated else 32],
                                  "values_data": [0] + [None] * 31})
                if modern or full or integrated:
                    chart.show_hidden_data()
            sheet.insert_chart(12 + index * 16, 0, chart)
        workbook.close()

    def complete(self, corrupt=None):
        self.write_book(self.modern, modern=True, corrupt=corrupt)
        for name in ("HWLN 1月 32方位數值_2.xlsx", "HWLN 1月 32方位數值.xlsx",
                     "HWLN 1月 32方位數值_2 總表.xlsx", "HWLN 1月 32方位數值 總表.xlsx",
                     "HWLN_1月_32方位_每日較大值整合總表.xlsx"):
            self.write_book(self.delivery / name, corrupt=corrupt)

    def allowed_missing(self, path):
        values = openpyxl.load_workbook(path, read_only=True, data_only=True)
        formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            issues, _entries, permitted = audit.check_modern_chart_helpers(
                formulas, values, [self.day], list(range(32)), [[0] + [None] * 31])
            self.assertEqual(issues, [])
            radar_issues, _radar_entries, radar_permitted = audit.check_modern_radar_helpers(
                formulas, values, [self.day], list(range(32)), [[0] + [None] * 31])
            self.assertEqual(radar_issues, [])
            return permitted | radar_permitted
        finally:
            values.close()
            formulas.close()

    def test_all_six_books_and_default_decisions_pass(self):
        self.complete()
        result = audit.audit_job(self.job)
        self.assertTrue(result["pass"], result["issues"])
        self.assertEqual(result["workbooks"], 6)
        self.assertEqual(result["chart_numeric_points"], 768)

    def test_five_historical_native_contracts_bind_all_matrices_current_hashes_and_formulas(self):
        self.complete()
        result = audit.audit_job(self.job)
        self.assertTrue(result["pass"], result["issues"])
        contracts = result["historical_chart_contracts"]
        self.assertEqual({item["role"] for item in contracts}, {"full_a", "full_b", "summary_a", "summary_b", "integrated"})
        self.assertEqual(len(contracts), 5)
        for item in contracts:
            self.assertEqual(item["schema"], 1)
            self.assertEqual(item["state"], "PASS_SERIALIZATION_PREREQUISITES")
            self.assertEqual(item["workbook_sha256"], audit.sha256(item["workbook"]))
            self.assertEqual(item["cache_sources"][0]["sha256"], audit.sha256(self.cache_file))
            self.assertEqual(item["expected_matrix"], [[0] + [None] * 31])
            self.assertEqual(item["statistic_order"], list(range(32)))
            self.assertEqual(item["statistic_matrix"], [[0, 0, 0]] + [[None, None, None]] * 31)
            self.assertEqual(item["direction_labels"], list(audit.ABBREVIATIONS))
            self.assertTrue(all(cell["verified"] for cell in item["source_cells"]))
            self.assertEqual(item["verified_source_cell_count"], len(item["source_cells"]))
            formula_role = item["role"] in ("full_a", "full_b", "integrated")
            self.assertEqual(item["radar_series"][0]["values"], "總表!AO2:BT2" if formula_role else "總表!B2:AG2")
            self.assertEqual(len(item["radar_helpers"]), 32 if formula_role else 0)
            self.assertEqual(item["radar_verified_intentional_missing_count"], 31 if formula_role else 0)
            if formula_role:
                self.assertEqual(item["radar_helpers"][0]["expected_value"], 0)
                self.assertFalse(item["radar_helpers"][0]["expected_missing"])
                self.assertEqual(item["radar_helpers"][0]["cached_value"], 0)
            self.assertEqual([series["name"] for series in item["statistic_series"]], ["MAX", "MIN", "平均值"])

    def test_historical_correct_cache_cannot_hide_shifted_primary_statistic_or_integrated_formula(self):
        for role, old, new in (
            ("full_a", b"COUNT('1", b"COUNT('2"),
            ("summary_a", b'COUNT(B2:B2)&lt;1,"",MAX(B2:B2)', b'COUNT(C2:C2)&lt;1,"",MAX(C2:C2)'),
            ("integrated", b"'Period B'!B2", b"'Period B'!C2"),
        ):
            with self.subTest(role=role):
                self.complete()
                paths = {"full_a": self.delivery / "HWLN 1月 32方位數值_2.xlsx",
                         "summary_a": self.delivery / "HWLN 1月 32方位數值_2 總表.xlsx",
                         "integrated": self.delivery / "HWLN_1月_32方位_每日較大值整合總表.xlsx"}
                path = paths[role]
                with zipfile.ZipFile(path) as source:
                    parts = {name: source.read(name) for name in source.namelist()}
                self.assertIn(old, parts["xl/worksheets/sheet1.xml"])
                parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(old, new)
                with zipfile.ZipFile(path, "w") as target:
                    for name, value in parts.items():
                        target.writestr(name, value)
                result = audit.audit_job(self.job)
                self.assertFalse(result["pass"])
                self.assertTrue(any("prerequisite" in issue for issue in result["issues"]), result["issues"])
                self.assertTrue(all(item["state"] == "FAIL" for item in result["historical_chart_contracts"]))
                self.assertEqual(result["modern_chart_missing_contract"]["state"], "FAIL")

    def test_historical_integrated_source_corruption_not_hidden_by_correct_final_cache(self):
        self.complete()
        path = self.delivery / "HWLN_1月_32方位_每日較大值整合總表.xlsx"
        with zipfile.ZipFile(path) as source:
            parts = {name: source.read(name) for name in source.namelist()}
        old = b'<c r="B2"><v>0</v></c>'
        self.assertIn(old, parts["xl/worksheets/sheet2.xml"])
        parts["xl/worksheets/sheet2.xml"] = parts["xl/worksheets/sheet2.xml"].replace(old, b'<c r="B2"><v>42</v></c>')
        with zipfile.ZipFile(path, "w") as target:
            for name, value in parts.items():
                target.writestr(name, value)
        result = audit.audit_job(self.job)
        self.assertFalse(result["pass"])
        self.assertTrue(any("integrated Period A source" in issue for issue in result["issues"]), result["issues"])

    def test_only_exact_historical_chart_helpers_receive_missing_error_permission(self):
        for corruption in ("historical_numeric_source_na", "historical_missing_as_zero", "historical_shifted_helper",
                           "historical_literal_error", "historical_unexpected_na"):
            with self.subTest(corruption=corruption):
                self.complete(corruption)
                result = audit.audit_job(self.job)
                self.assertFalse(result["pass"])
                self.assertTrue(all(item["state"] == "FAIL" for item in result["historical_chart_contracts"]))

    def test_correct_historical_helpers_do_not_permit_missing_chart_cache_as_zero(self):
        self.complete()
        path = self.delivery / "HWLN 1月 32方位數值_2.xlsx"
        with zipfile.ZipFile(path) as source:
            parts = {name: source.read(name) for name in source.namelist()}
        parts["xl/charts/chart1.xml"] = parts["xl/charts/chart1.xml"].replace(
            b'<c:pt idx="0"><c:v>0</c:v></c:pt>',
            b'<c:pt idx="0"><c:v>0</c:v></c:pt><c:pt idx="1"><c:v>0</c:v></c:pt>')
        with zipfile.ZipFile(path, "w") as target:
            for name, value in parts.items():
                target.writestr(name, value)
        result = audit.audit_job(self.job)
        self.assertFalse(result["pass"])
        self.assertTrue(any("cache/cell" in issue for issue in result["issues"]), result["issues"])

    def test_modern_summary_daily_and_ledger_mutations_fail(self):
        for field in ("summary", "daily", "ledger", "candidate", "ledger_formula"):
            with self.subTest(field=field):
                self.complete(field)
                self.assertFalse(audit.audit_job(self.job)["pass"])

    def test_five_matching_value_books_do_not_hide_missing_output_role(self):
        self.complete()
        old = self.delivery / "HWLN 1月 32方位數值_2.xlsx"
        old.rename(self.delivery / "HWLN 1月 duplicated-period-b.xlsx")
        result = audit.audit_job(self.job)
        self.assertFalse(result["pass"])
        self.assertTrue(any("role filenames" in issue for issue in result["issues"]), result["issues"])

    def test_missing_or_extra_cache_day_fails_closed(self):
        self.complete()
        (self.cache / "20260102.json").write_bytes(b"{}")
        with self.assertRaises(ValueError):
            audit.audit_job(self.job)

    def test_wrong_cache_port_is_not_accepted(self):
        payload = json.loads(self.cache_file.read_text(encoding="utf-8"))
        payload["signature"]["port"] = "KLNG"
        self.cache_file.write_text(json.dumps(payload), encoding="utf-8")
        with self.assertRaises(ValueError):
            audit.audit_job(self.job)

    def test_real_zero_is_not_blank_and_nonfinite_is_not_valid(self):
        self.assertFalse(audit.equal(None, 0))
        self.assertFalse(audit.equal(0, None))
        self.assertTrue(audit.equal(None, ""))
        self.assertFalse(audit.equal(0, float("nan")))
        self.assertFalse(audit.equal(0, False))
        self.assertFalse(audit.equal(0, 1e-12))
        self.assertFalse(audit.equal(1.25, 1.25 + 1e-10))
        value = 123.45678901234567
        self.assertTrue(audit.equal(value, float(format(value, ".16G"))))

    def test_sample_statistics_empty_one_and_many(self):
        self.assertEqual(audit.stats([None, ""]), (None,) * 4)
        self.assertEqual(audit.stats([0, None]), (0, 0, 0, None))
        self.assertEqual(audit.stats([1, 2, 3]), (3, 1, 2, 1))

    def test_stale_cache_point_and_blank_as_zero_are_detected(self):
        self.complete()
        wb = openpyxl.load_workbook(self.modern, read_only=True, data_only=True)
        try:
            with zipfile.ZipFile(self.modern) as source:
                parts = {name: source.read(name) for name in source.namelist()}
            xml = parts["xl/charts/chart1.xml"]
            xml = xml.replace(b'<c:pt idx="0"><c:v>0</c:v></c:pt>',
                              b'<c:pt idx="0"><c:v>88</c:v></c:pt><c:pt idx="1"><c:v>0</c:v></c:pt>')
            parts["xl/charts/chart1.xml"] = xml
            changed = self.root / "changed.xlsx"
            with zipfile.ZipFile(changed, "w") as target:
                for name, value in parts.items():
                    target.writestr(name, value)
            issues, _ = audit.check_chart_caches(changed, wb, self.allowed_missing(self.modern))
            self.assertEqual(len(issues), 2, issues)
        finally:
            wb.close()

    def test_two_empty_chart_parts_or_literal_series_cannot_pass(self):
        self.complete()
        wb = openpyxl.load_workbook(self.modern, read_only=True, data_only=True)
        try:
            with zipfile.ZipFile(self.modern) as source:
                parts = {name: source.read(name) for name in source.namelist()}
            changed = self.root / "empty-charts.xlsx"
            with zipfile.ZipFile(changed, "w") as target:
                for name, value in parts.items():
                    if name.startswith("xl/charts/chart") and name.endswith(".xml"):
                        value = b'<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:chart/></c:chartSpace>'
                    target.writestr(name, value)
            issues, points = audit.check_chart_caches(changed, wb)
            self.assertEqual(points, 0)
            self.assertEqual(len(issues), 2)
        finally:
            wb.close()

    def test_self_consistent_but_wrong_chart_binding_is_rejected(self):
        self.complete()
        with zipfile.ZipFile(self.modern) as source:
            parts = {name: source.read(name) for name in source.namelist()}
        # Same numerical zero/blank vector in MAX as the daily row: cache/cell
        # equality alone passes, but it is NOT the intended daily radar source.
        parts["xl/charts/chart1.xml"] = parts["xl/charts/chart1.xml"].replace(b"$AO$5:$BT$5", b"$B$7:$AG$7")
        changed = self.root / "wrong-source.xlsx"
        with zipfile.ZipFile(changed, "w") as target:
            for name, content in parts.items():
                target.writestr(name, content)
        values = openpyxl.load_workbook(changed, read_only=True, data_only=True)
        formulas = openpyxl.load_workbook(changed, read_only=True, data_only=False)
        try:
            self.assertEqual(audit.check_chart_caches(changed, values, self.allowed_missing(changed))[0], [])
            issues = audit.check_modern_bindings(changed, formulas, [self.day], list(range(32)), [[0] + [None] * 31])
            self.assertTrue(any("intended" in issue for issue in issues), issues)
        finally:
            values.close()
            formulas.close()

    def test_correct_cached_value_does_not_hide_wrong_final_formula(self):
        self.complete()
        wb = openpyxl.load_workbook(self.modern)
        wb["1月1日"]["H4"] = "='決策台帳'!H3"
        changed = self.root / "wrong-final.xlsx"
        wb.save(changed)
        wb.close()
        formulas = openpyxl.load_workbook(changed, read_only=True, data_only=False)
        try:
            issues = audit.check_modern_bindings(changed, formulas, [self.day], list(range(32)), [[0] + [None] * 31])
            self.assertTrue(any("final/summary" in issue for issue in issues), issues)
        finally:
            formulas.close()

    def test_intentional_chart_missing_contract_is_hash_bound_and_preserves_real_zero(self):
        self.complete()
        result = audit.audit_job(self.job)
        self.assertTrue(result["pass"], result["issues"])
        contract = result["modern_chart_missing_contract"]
        self.assertEqual(contract["schema"], 1)
        self.assertEqual(contract["state"], "PASS_SERIALIZATION_PREREQUISITES")
        self.assertEqual(contract["role"], "Modern/default")
        self.assertEqual(contract["workbook_sha256"], audit.sha256(self.modern))
        self.assertEqual(contract["cache_sources"][0]["sha256"], audit.sha256(self.cache_file))
        self.assertEqual(contract["expected_matrix"], [[0] + [None] * 31])
        self.assertEqual(contract["verified_intentional_missing_count"], 93)
        self.assertEqual(contract["radar_verified_intentional_missing_count"], 31)
        self.assertEqual(len(contract["radar_helpers"]), 32)
        self.assertEqual(contract["radar_helpers"][0]["cell"], "AO5")
        self.assertEqual(contract["radar_helpers"][0]["direction"], "北")
        self.assertEqual(contract["radar_helpers"][0]["source"], "B5")
        self.assertEqual(contract["radar_helpers"][0]["formula"], '=IF(COUNT(B5)=0,NA(),B5)')
        self.assertEqual(len(contract["helpers"]), 96)
        for helper in contract["helpers"]:
            self.assertTrue(helper["verified"])
            if helper["direction_index"] == 0:
                self.assertEqual(helper["expected_value"], 0)
                self.assertFalse(helper["expected_missing"])
                self.assertEqual(helper["cached_type"], "n")
                self.assertEqual(helper["cached_value"], 0)
            else:
                self.assertTrue(helper["expected_missing"])
                self.assertIsNone(helper["expected_value"])
                self.assertEqual(helper["cached_type"], "e")
                self.assertEqual(helper["cached_value"], "#N/A")

    def test_only_exact_independently_missing_helpers_can_use_na(self):
        for corruption in ("numeric_source_na", "missing_as_blank", "shifted_helper_source",
                           "primary_source_na", "helper_error_as_text", "unexpected_na_outside_helper",
                           "unexpected_other_error", "radar_numeric_source_na", "radar_missing_as_zero", "radar_shifted_ref"):
            with self.subTest(corruption=corruption):
                self.complete(corruption)
                result = audit.audit_job(self.job)
                self.assertFalse(result["pass"], result["issues"])
                self.assertEqual(result["modern_chart_missing_contract"]["state"], "FAIL")

    def test_chart_cache_missing_as_zero_fails_even_with_correct_intentional_na(self):
        self.complete()
        permitted = self.allowed_missing(self.modern)
        with zipfile.ZipFile(self.modern) as source:
            parts = {name: source.read(name) for name in source.namelist()}
        parts["xl/charts/chart2.xml"] = parts["xl/charts/chart2.xml"].replace(
            b'<c:pt idx="0"><c:v>0</c:v></c:pt>',
            b'<c:pt idx="0"><c:v>0</c:v></c:pt><c:pt idx="1"><c:v>0</c:v></c:pt>')
        changed = self.root / "na-cache-as-zero.xlsx"
        with zipfile.ZipFile(changed, "w") as target:
            for name, content in parts.items():
                target.writestr(name, content)
        values = openpyxl.load_workbook(changed, read_only=True, data_only=True)
        try:
            issues, _points = audit.check_chart_caches(changed, values, permitted)
            self.assertEqual(sum("cache/cell" in issue for issue in issues), 3, issues)
        finally:
            values.close()

    def test_no_implicit_error_permission_from_formula_count_or_blank(self):
        self.complete()
        values = openpyxl.load_workbook(self.modern, read_only=True, data_only=True)
        try:
            issues, _points = audit.check_chart_caches(self.modern, values)
            self.assertEqual(len(issues), 124)
            errors, intentional = audit.check_saved_errors(self.modern)
            self.assertEqual(len(errors), 124)
            self.assertEqual(intentional, 0)
        finally:
            values.close()

    def test_wrong_grid_or_statistic_count_guards_are_not_na_prerequisites(self):
        self.complete()
        with zipfile.ZipFile(self.modern) as source:
            parts = {name: source.read(name) for name in source.namelist()}
        original = parts["xl/worksheets/sheet1.xml"]
        for old, new in ((b'COUNT(C5:C5)&lt;1,"",MAX(C5:C5)', b'COUNT(D5:D5)&lt;1,"",MAX(D5:D5)'),
                         (b"COUNT('1", b"COUNT('2")):
            with self.subTest(old=old):
                self.assertIn(old, original)
                changed = self.root / "wrong-prerequisite.xlsx"
                with zipfile.ZipFile(changed, "w") as target:
                    for name, content in parts.items():
                        target.writestr(name, content.replace(old, new) if name == "xl/worksheets/sheet1.xml" else content)
                formulas = openpyxl.load_workbook(changed, read_only=True, data_only=False)
                values = openpyxl.load_workbook(changed, read_only=True, data_only=True)
                try:
                    stat_issues, _entries, _permitted = audit.check_modern_chart_helpers(
                        formulas, values, [self.day], list(range(32)), [[0] + [None] * 31])
                    radar_issues, _entries, _permitted = audit.check_modern_radar_helpers(
                        formulas, values, [self.day], list(range(32)), [[0] + [None] * 31])
                    self.assertTrue(stat_issues or radar_issues)
                finally:
                    formulas.close()
                    values.close()

    def test_hidden_radar_source_explicitly_disabled_is_rejected(self):
        self.complete()
        with zipfile.ZipFile(self.modern) as source:
            parts = {name: source.read(name) for name in source.namelist()}
        parts["xl/charts/chart1.xml"] = parts["xl/charts/chart1.xml"].replace(
            b"</c:chart>", b'<c:plotVisOnly val="1"/></c:chart>')
        changed = self.root / "hidden-radar-disabled.xlsx"
        with zipfile.ZipFile(changed, "w") as target:
            for name, content in parts.items():
                target.writestr(name, content)
        formulas = openpyxl.load_workbook(changed, read_only=True, data_only=False)
        try:
            issues = audit.check_modern_bindings(changed, formulas, [self.day], list(range(32)), [[0] + [None] * 31])
            self.assertTrue(any("plotVisOnly" in issue for issue in issues))
        finally:
            formulas.close()

    def test_historical_missing_nonempty_day_and_wrong_source_are_rejected(self):
        self.complete()
        path = self.delivery / "HWLN 1月 32方位數值 總表.xlsx"
        values = openpyxl.load_workbook(path, read_only=True, data_only=True)
        row = [0] + [None] * 31
        try:
            # An entirely missing second day legitimately has no filled radar
            # series. A second observed day MUST have its own series/source.
            days = [self.day, date(2026, 1, 2)]
            self.assertEqual(audit.check_historical_bindings(path, values, days, list(range(32)), [row, [None] * 32]), [])
            issues = audit.check_historical_bindings(path, values, days, list(range(32)), [row, row])
            self.assertTrue(any("coverage" in issue for issue in issues), issues)
            with zipfile.ZipFile(path) as source:
                parts = {name: source.read(name) for name in source.namelist()}
            parts["xl/charts/chart1.xml"] = parts["xl/charts/chart1.xml"].replace(b"$B$2:$AG$2", b"$B$4:$AG$4")
            changed = self.root / "historical-wrong-source.xlsx"
            with zipfile.ZipFile(changed, "w") as target:
                for name, content in parts.items():
                    target.writestr(name, content)
            self.assertEqual(audit.check_chart_caches(changed, values)[0], [])
            issues = audit.check_historical_bindings(changed, values, [self.day], list(range(32)), [row])
            self.assertTrue(any("intended" in issue for issue in issues), issues)
        finally:
            values.close()


if __name__ == "__main__":
    unittest.main()
