"""Diagnostic actual writer replay of ALL March cache days; not a new raw-data regression."""
from __future__ import annotations

import argparse
from datetime import date
import hashlib
import json
from pathlib import Path

import openpyxl
from openpyxl.utils.cell import get_column_letter
import ais_monthly_app as app
from scripts import audit_default_workbook_content as audit
from scripts.independent_source_audit import save_report


def sha(path):
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def compare_cells(old, new, helper_plan=(), radar_plan=()):
    issues, checked, allowed = [], 0, []
    intended_helpers = {(entry["sheet"], entry["cell"]): entry for entry in helper_plan if entry["verified"]}
    intended_radar = {(entry["sheet"], entry["cell"]): entry for entry in radar_plan if entry["verified"]}
    for data_only in (False, True):
        left = openpyxl.load_workbook(old, read_only=True, data_only=data_only)
        right = openpyxl.load_workbook(new, read_only=True, data_only=data_only)
        try:
            if left.sheetnames != right.sheetnames:
                issues.append("worksheet names/order differ")
            for name in left.sheetnames:
                a, b = left[name], right[name]
                if (a.max_row != b.max_row or (a.max_column != b.max_column
                        and not (name == "總表" and a.max_column <= 40 and b.max_column == 72))):
                    issues.append(f"sheet extent differs: {name}")
                bounds = {"max_row": max(a.max_row, b.max_row), "max_col": max(a.max_column, b.max_column)}
                for row_number, (row_left, row_right) in enumerate(zip(a.iter_rows(**bounds), b.iter_rows(**bounds)), 1):
                    for column_number, (cell_left, cell_right) in enumerate(zip(row_left, row_right), 1):
                        checked += 1
                        before, after = cell_left.value, cell_right.value
                        if before == after:
                            continue
                        coordinate = f"{get_column_letter(column_number)}{row_number}"
                        location = name, coordinate
                        if location in (("操作說明", "B7"), ("系統資料", "B11")) and before == str(old) and after == str(new):
                            allowed.append({"data_only": data_only, "sheet": name, "cell": coordinate,
                                            "reason": "actual output workbook self-path only"})
                        elif location in intended_helpers:
                            plan = intended_helpers[location]
                            old_formula = f'=IF(COUNT({plan["source"]})=0,"",{plan["source"]})'
                            permitted = ((not data_only and before == old_formula and after == plan["formula"])
                                         or (data_only and plan["expected_missing"] and audit.blank(before)
                                             and after == "#N/A" and cell_right.data_type == "e"))
                            if permitted:
                                allowed.append({"data_only": data_only, "sheet": name, "cell": coordinate,
                                                "reason": "independently verified chart-only helper missing semantics"})
                            else:
                                issues.append(f"unintended helper value/formula differs: {location}, data_only={data_only}")
                        elif location in intended_radar:
                            plan = intended_radar[location]
                            permitted = (audit.blank(before) and
                                         ((not data_only and after == plan["formula"] and cell_right.data_type == "f")
                                          or (data_only and plan["expected_missing"] and after == "#N/A" and cell_right.data_type == "e")
                                          or (data_only and not plan["expected_missing"] and cell_right.data_type == "n"
                                              and audit.equal(plan["expected_value"], after))))
                            if permitted:
                                allowed.append({"data_only": data_only, "sheet": name, "cell": coordinate,
                                                "reason": "independently verified added chart-only radar helper"})
                            else:
                                issues.append(f"unintended radar helper differs: {location}, data_only={data_only}")
                        elif (name == "總表" and row_number == 4 and 41 <= column_number <= 72
                              and audit.blank(before) and after == audit.LABELS[column_number - 41]):
                            allowed.append({"data_only": data_only, "sheet": name, "cell": coordinate,
                                            "reason": "exact independent 32-direction added radar helper header"})
                        else:
                            issues.append(f"saved value/formula differs: {location}, data_only={data_only}")
        finally:
            left.close()
            right.close()
    return issues, checked, allowed


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("run_root", type=Path)
    parser.add_argument("output_dir", type=Path)
    parser.add_argument("--expected-source-sha", required=True)
    parser.add_argument("--generated-workbook", type=Path,
                        help="Read-only audit an existing generated candidate; never rewrite it")
    args = parser.parse_args(argv)
    production = Path(app.__file__)
    if sha(production).lower() != args.expected_source_sha.lower():
        raise ValueError("production source differs from specified diagnostic candidate")
    if args.output_dir.exists():
        raise ValueError("diagnostic output must be fresh; do not replace prior evidence")
    run = json.loads((args.run_root / "run_report.json").read_text(encoding="utf-8"))
    jobs = [job for job in run["jobs"] if job["token"] == "HWLN_202603"]
    if len(jobs) != 1 or jobs[0]["status"] != "PASS":
        raise ValueError("no complete March execution to diagnose")
    job = jobs[0]
    old = Path(job["modern"])
    days = [date.fromisoformat(value) for value in job["logical_days"]]
    caches = list(old.parent.glob("*_cache"))
    if len(caches) != 1:
        raise ValueError("exactly one source cache required")
    paths = sorted(caches[0].glob("*.json"))
    if {path.stem for path in paths} != {day.strftime("%Y%m%d") for day in days}:
        raise ValueError("complete March cache coverage differs")
    hashes = {str(path): sha(path) for path in [old, *paths]}
    payloads = [json.loads(path.read_text(encoding="utf-8")) for path in paths]
    results = [app._day_result_from_dict(payload["result"]) for payload in payloads]
    scopes = {day: {"logical_day": payload["result"]["directions"], **payload["result"]["period_directions"]}
              for day, payload in zip(days, payloads)}
    new = args.generated_workbook or (args.output_dir / old.name)
    if args.generated_workbook is None:
        config = app.AppConfig(input_dir=Path(job["folder"]), output_path=new, port="HWLN", year=2026, month=3)
        app.write_monthly_workbook(config, results, [], app.build_default_decision_snapshot(config, results))
    if not new.is_file() or new.resolve() == old.resolve():
        raise ValueError("generated diagnostic workbook must exist separately from source")
    hashes[str(new.resolve())] = sha(new)
    values = openpyxl.load_workbook(new, read_only=True, data_only=True)
    formulas = openpyxl.load_workbook(new, read_only=True, data_only=False)
    try:
        expected = [[scopes[day]["logical_day"][label]["selected"] for label in audit.LABELS] for day in days]
        helper_issues, helper_entries, permitted_missing = audit.check_modern_chart_helpers(
            formulas, values, days, list(range(32)), expected)
        radar_issues, radar_entries, radar_missing = audit.check_modern_radar_helpers(
            formulas, values, days, list(range(32)), expected)
        issues, checked, allowed = compare_cells(old, new, helper_entries, radar_entries)
        issues += helper_issues
        issues += radar_issues
        error_issues, total_intentional_count = audit.check_saved_errors(new, permitted_missing | radar_missing)
        issues += error_issues
        if total_intentional_count != len(permitted_missing | radar_missing):
            issues.append("saved intentional errors differ from exact verified helper role coverage")
        old_values = openpyxl.load_workbook(old, read_only=True, data_only=True)
        try:
            old_cache_issues, old_points = audit.check_chart_caches(old, old_values)
        finally:
            old_values.close()
        new_cache_issues, new_points = audit.check_chart_caches(new, values, permitted_missing | radar_missing)
        issues += new_cache_issues
        issues += audit.check_modern_bindings(new, formulas, days, list(range(32)), expected)
        candidate_issues, candidate_count = audit.check_candidate_serialization(values, scopes, days, list(range(32)))
        issues += candidate_issues
        issues += audit.check_ledger_final_formulas(formulas, days, list(range(32)), candidate_count)
        missing = {cell: {"value": values["總表"][cell].value, "formula": formulas["總表"][cell].value}
                   for cell in ("AJ36", "AK36", "AL36", "AG37", "AG38", "AG39")}
        if (any(missing[cell]["value"] is not None for cell in ("AG37", "AG38", "AG39"))
                or any(missing[cell]["value"] != "#N/A" for cell in ("AJ36", "AK36", "AL36"))):
            issues.append("expected missing March NbW primary blanks/chart-only NA differ")
    finally:
        values.close()
        formulas.close()
    unchanged = hashes == {path: sha(Path(path)) for path in hashes}
    if not unchanged or sha(production).lower() != args.expected_source_sha.lower():
        issues.append("source workbook/cache/production changed during diagnostic replay")
    bound_hashes = {str(path.resolve()): hashes[str(path.resolve())] if str(path.resolve()) in hashes else hashes[str(path)]
                    for path in [new, *paths]}
    contract = audit.build_modern_missing_contract(job, new, days, list(range(32)), expected, paths,
                                                   bound_hashes, helper_entries, len(permitted_missing),
                                                   radar_entries, len(radar_missing), issues)
    report = {"source_app_sha256": args.expected_source_sha, "original": str(old), "original_sha256": hashes[str(old)],
              "generated": str(new), "generated_sha256": sha(new), "days": len(days), "candidate_rows_checked": candidate_count,
              "saved_cells_compared_formula_and_value": checked, "allowed_output_selfpath_changes": allowed,
              "old_chart_cache_issues": old_cache_issues, "old_points": old_points, "new_points": new_points,
              "missing_cells": missing, "inputs_unchanged": unchanged, "issues": issues, "pass": not issues,
              "jobs": [{"token": job["token"], "pass": not issues, "modern_chart_missing_contract": contract}],
              "boundary": "Full 31-day cached-data actual writer causal replay only. NOT a full raw-input EXE rerun or native visual acceptance."}
    save_report(args.output_dir / "diagnostic_report.json", report,
                [args.run_root, Path(job["folder"])], [old, *paths, production])
    print(json.dumps({"pass": report["pass"], "days": len(days), "cells": checked,
                      "candidate_rows": candidate_count, "chart_points": new_points,
                      "intended_changes": len(allowed), "issues": issues,
                      "report": str(args.output_dir / "diagnostic_report.json")}, ensure_ascii=False))
    return 0 if report["pass"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
