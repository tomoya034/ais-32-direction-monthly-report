from __future__ import annotations

import datetime as dt
import math
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import openpyxl

import ais_monthly_app as app


# Fixed index metadata for reading pre-1.5.1 external baselines, never a runtime alias.
V150_DIRECTION_ORDER = (
    "北", "北微東", "北北東", "東北微北", "東北", "東北微東", "東北東", "東微北",
    "東", "東微南", "東南東", "東南微南", "東南", "東南微西", "南南東", "南微東",
    "南", "南微西", "南南西", "西南微南", "西南", "西南微西", "西南西", "西微南",
    "西", "西微北", "西北西", "西北微西", "西北", "西北微北", "北北西", "北微西",
)

DIRECTION_CONTRACT = (
    (0, "N", "北"), (1, "NbE", "北微東"), (2, "NNE", "東北偏北"), (3, "NEbN", "東北微北"),
    (4, "NE", "東北"), (5, "NEbE", "東北微東"), (6, "ENE", "東北偏東"), (7, "EbN", "東微北"),
    (8, "E", "東"), (9, "EbS", "東微南"), (10, "ESE", "東南偏東"), (11, "SEbE", "東南微東"),
    (12, "SE", "東南"), (13, "SEbS", "東南微南"), (14, "SSE", "東南偏南"), (15, "SbE", "南微東"),
    (16, "S", "南"), (17, "SbW", "南微西"), (18, "SSW", "西南偏南"), (19, "SWbS", "西南微南"),
    (20, "SW", "西南"), (21, "SWbW", "西南微西"), (22, "WSW", "西南偏西"), (23, "WbS", "西微南"),
    (24, "W", "西"), (25, "WbN", "西微北"), (26, "WNW", "西北偏西"), (27, "NWbW", "西北微西"),
    (28, "NW", "西北"), (29, "NWbN", "西北微北"), (30, "NNW", "西北偏北"), (31, "NbW", "北微西"),
)


def write_direction_fixture(path: Path) -> None:
    workbook = openpyxl.Workbook(write_only=True)
    sheet = workbook.create_sheet("AIS")
    sheet.append(["Year", "Month", "Day", "Hour", "Minute", "Second", "channel", "msg_type",
                  "mmsi", "LONGITUDE_DESC", "bearing", "distance in nautical miles"])
    for hour in (1, 13):
        for index in range(32):
            for offset in range(3):
                second = index * 3 + offset
                sheet.append([2026, 6, 1, hour, second // 60, second % 60, "A", 1,
                              800000000 + index, "East", index * 11.25 + 1,
                              50.0 + index - offset])
    workbook.save(path)


class DirectionMappingTests(unittest.TestCase):
    def test_complete_index_abbreviation_chinese_contract(self) -> None:
        self.assertEqual(len(app.DIRECTION_ORDER), 32)
        self.assertEqual(len(app.LEGACY_DIRECTION_ABBREVIATIONS), 32)
        self.assertEqual(len(set(app.DIRECTION_ORDER)), 32)
        actual = tuple((i, abbreviation, label) for i, (abbreviation, label) in enumerate(
            zip(app.LEGACY_DIRECTION_ABBREVIATIONS, app.DIRECTION_ORDER)))
        self.assertEqual(actual, DIRECTION_CONTRACT)
        self.assertEqual(actual[11], (11, "SEbE", "東南微東"))
        self.assertEqual(actual[13], (13, "SEbS", "東南微南"))
        self.assertEqual([i for i in range(32) if V150_DIRECTION_ORDER[i] != app.DIRECTION_ORDER[i]],
                         [2, 6, 10, 11, 13, 14, 18, 22, 26, 30])

    def test_degree_index_retains_v150_floor_bin_behavior(self) -> None:
        for index in range(32):
            start = index * 11.25
            for degree in (start, start + 5.625, math.nextafter(start + 11.25, -math.inf)):
                self.assertEqual(app.degree_to_direction_index(degree), (index, degree))
        for degree in range(-72000, 72001):
            value = degree / 100.0
            normalized = value % 360.0
            self.assertEqual(app.degree_to_direction_index(value),
                             (min(int(normalized // 11.25), 31), normalized))
        for invalid in (None, "bad", math.nan, math.inf, -math.inf):
            self.assertIsNone(app.degree_to_direction_index(invalid))

    def test_open_sea_and_wsw_coastal_review_contract(self) -> None:
        self.assertEqual(app.OPEN_SEA_INDEXES, tuple(range(11)) + tuple(range(22, 32)))
        self.assertEqual(app.COASTAL_REVIEW_DIRECTIONS, {"西南偏西", "西微南", "西"})
        self.assertEqual(app.DIRECTION_ORDER[22], "西南偏西")
        config = app.AppConfig(Path("."), Path("test.xlsx"), "KLNG", 2026, 6)
        for maximum, status in ((12.0, "待複核"), (10.0, "自動採用")):
            rows = [(maximum - delta, row, 247.5) for row, delta in enumerate((0, .1, .2), 2)]
            selected = app.select_cluster_from_sorted_rows(
                "西南偏西", rows, config, over_cap_count=0, over_cap_max=None)
            self.assertEqual(selected.status, status)
            self.assertEqual((selected.selected, selected.selected_rank, selected.cluster_count),
                             (maximum, 1, 3))


class DirectionPersistenceTests(unittest.TestCase):
    def test_old_review_schema_is_rejected_at_all_entry_points(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = root / "old-review.xlsx"
            workbook = openpyxl.Workbook()
            workbook.active.title = "系統資料"
            workbook.active.append(["Schema", "AIS_V15_REVIEW_1"])
            workbook.create_sheet("決策台帳")
            workbook.create_sheet("候選清單")
            workbook.save(path)
            workbook.close()
            before = path.read_bytes()
            for reader in (app.read_review_decisions, app.config_from_review_workbook,
                           app.finalize_review_workbook):
                with self.subTest(reader=reader.__name__):
                    with self.assertRaisesRegex(ValueError, "32 方位中文名稱.*原始月份資料重新分析.*decision direction"):
                        reader(path)
            self.assertEqual(path.read_bytes(), before)
            self.assertEqual([p.name for p in root.iterdir()], [path.name])
            for schema in (None, "unknown"):
                with self.assertRaisesRegex(ValueError, "schema 版本不相容"):
                    app.validate_review_workbook_schema(schema)

    def test_cache_v4_invalidates_without_changing_numeric_spool(self) -> None:
        self.assertEqual(app.CACHE_VERSION, 5)
        self.assertEqual(app.NORMALIZED_SPOOL_VERSION, 1)
        self.assertEqual(app.REVIEW_WORKBOOK_SCHEMA, "AIS_V15_REVIEW_2")
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "D&TMOK KLNG_20260601_any.xlsx"
            write_direction_fixture(source)
            config = app.AppConfig(root, root / "analysis.xlsx", "KLNG", 2026, 6)
            day_input = app.build_processing_job(config).days[0]
            result = app.process_day_input(day_input, config)
            spool = app._normalized_spool_path(config, day_input.day)
            original = spool.read_bytes()
            with mock.patch.object(app, "CACHE_VERSION", 4):
                app.save_cached_day(day_input, config, result)
            self.assertIsNone(app.load_cached_day(day_input, config))
            self.assertEqual(spool.read_bytes(), original)
            app.save_cached_day(day_input, config, result)
            self.assertIsNotNone(app.load_cached_day(day_input, config))

    def test_all_workbook_direction_displays_follow_index_contract(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "D&TMOK KLNG_20260601_any.xlsx"
            write_direction_fixture(source)
            config = app.AppConfig(root, root / "analysis.xlsx", "KLNG", 2026, 6)
            result = app.process_day_input(app.build_processing_job(config).days[0], config)
            snapshot = app.build_default_decision_snapshot(config, [result])
            app.write_monthly_workbook(config, [result], [], snapshot)
            self.assertEqual(app.read_review_decisions(config.output_path), snapshot)
            paths = app.write_delivery_workbooks(config, [result], snapshot, root=root / "delivery")
            modern = openpyxl.load_workbook(config.output_path)
            try:
                self.assertEqual([cell.value for cell in modern["總表"][4][1:33]], app.DIRECTION_ORDER)
                self.assertEqual([cell.value for cell in modern["6月1日"][1][7:39]], app.DIRECTION_ORDER)
                for name in ("候選清單", "決策台帳"):
                    self.assertEqual({row[2] for row in modern[name].iter_rows(min_row=2, values_only=True)},
                                     set(app.OPEN_SEA_DIRECTIONS))
                review_labels = {cell.value for row in modern["待複核"].iter_rows() for cell in row}
                self.assertIn("西南偏西", review_labels)
            finally:
                modern.close()
            for path in paths.all_files():
                workbook = openpyxl.load_workbook(path, data_only=True)
                try:
                    self.assertEqual([cell.value for cell in workbook["總表"][1][1:33]],
                                     app.LEGACY_DIRECTION_ABBREVIATIONS)
                    if path in (paths.period_a_full, paths.period_b_full):
                        self.assertEqual([cell.value for cell in workbook["6月1日"][1][7:39]],
                                         app.DIRECTION_ORDER)
                        for distance, bearing, label in workbook["6月1日"].iter_rows(
                                min_row=2, max_col=3, values_only=True):
                            if distance is not None:
                                self.assertEqual(label, app.DIRECTION_ORDER[app.degree_to_direction_index(bearing)[0]])
                        self.assertEqual(workbook["工作"].sheet_state, "hidden")
                        for degree in range(361):
                            self.assertEqual(workbook["工作"].cell(degree + 2, 11).value,
                                             app.DIRECTION_ORDER[app.degree_to_direction_index(degree)[0]])
                finally:
                    workbook.close()


if __name__ == "__main__":
    unittest.main()
