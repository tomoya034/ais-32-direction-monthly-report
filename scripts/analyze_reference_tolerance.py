from __future__ import annotations

import argparse
import json
import math
from datetime import date
from pathlib import Path

import openpyxl

from ais_monthly_app import (
    DIRECTION_ORDER,
    HISTORICAL_MESSAGE_TYPES,
    PERIOD_A,
    PERIOD_B,
    read_normalized_spool,
)


def reference_grid(path: Path, days: int) -> list[list[float | None]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["總表"]
        return [
            [
                None if sheet.cell(row, column).value in (None, "")
                else float(sheet.cell(row, column).value)
                for column in range(2, 34)
            ]
            for row in range(2, days + 2)
        ]
    finally:
        workbook.close()


def select(distances: list[float], tolerance: float, cluster_size: int = 3) -> float | None:
    values = sorted((value for value in distances if value <= 500.0), reverse=True)
    if not values:
        return None
    for start in range(0, len(values) - cluster_size + 1):
        floor = values[start] * (1.0 - tolerance)
        if values[start + cluster_size - 1] + 1e-12 >= floor:
            return values[start]
    return values[0]


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("spool_directory", type=Path)
    parser.add_argument("period_a_reference", type=Path)
    parser.add_argument("period_b_reference", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    arguments = parser.parse_args()
    references = {
        PERIOD_A: reference_grid(arguments.period_a_reference, 30),
        PERIOD_B: reference_grid(arguments.period_b_reference, 30),
    }
    records: dict[tuple[int, str, int], list[float]] = {}
    for day_number in range(1, 31):
        _counts, iterator = read_normalized_spool(
            arguments.spool_directory / f"202606{day_number:02d}.bin"
        )
        for direction_index, record in iterator:
            if record.msg_type not in HISTORICAL_MESSAGE_TYPES:
                continue
            records.setdefault((day_number, record.period, direction_index), []).append(
                record.distance
            )

    results = []
    included = tuple(range(11)) + tuple(range(22, 32))
    for percent in [value / 1000 for value in range(5, 151, 5)]:
        matches = 0
        compared = 0
        blanks = 0
        examples = []
        for day_number in range(1, 31):
            for period in (PERIOD_A, PERIOD_B):
                for direction_index in included:
                    expected = references[period][day_number - 1][direction_index]
                    actual = select(
                        records.get((day_number, period, direction_index), []), percent
                    )
                    compared += 1
                    if expected is None and actual is None:
                        matches += 1
                        blanks += 1
                    elif expected is not None and actual is not None and math.isclose(
                        expected, actual, rel_tol=0.0, abs_tol=1e-9
                    ):
                        matches += 1
                    elif len(examples) < 10:
                        examples.append(
                            {
                                "day": day_number,
                                "period": period,
                                "direction": DIRECTION_ORDER[direction_index],
                                "reference": expected,
                                "actual": actual,
                            }
                        )
        results.append(
            {
                "tolerance": percent,
                "matches": matches,
                "compared": compared,
                "match_rate": matches / compared,
                "blank_matches": blanks,
                "examples": examples,
            }
        )
    results.sort(key=lambda item: (-item["matches"], item["tolerance"]))
    payload = {"best": results[:5], "all": results}
    arguments.output.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(payload["best"], ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
